"""Recompute fixture answers from the saved XLSX, without calling an LLM."""
import csv
import json
from pathlib import Path
import unittest
import openpyxl
from scripts.test_fixture_oracle import evaluate, format_answer
from src.workbook_index import load_regions, retrieve

ROOT = Path(__file__).resolve().parents[1]


class ExtendedFixturesTest(unittest.TestCase):
    def test_every_reference_matches_saved_workbook(self):
        seen = set()
        total = 0
        for folder in sorted((ROOT / "Test").iterdir()):
            if not (folder / "answer_checks.json").exists():
                continue
            with self.subTest(suite=folder.name):
                queries = json.loads((folder / "queries.json").read_text(encoding="utf-8"))
                checks = json.loads((folder / "answer_checks.json").read_text(encoding="utf-8"))
                with (folder / "expected_submission.csv").open(encoding="utf-8", newline="") as stream:
                    reader = csv.DictReader(stream)
                    self.assertEqual(reader.fieldnames, ["id", "answer"])
                    expected = list(reader)
                self.assertEqual([q["id"] for q in queries], [r["id"] for r in expected])
                self.assertEqual([q["id"] for q in queries], [c["id"] for c in checks])
                self.assertEqual(len({q["query"] for q in queries}), len(queries))
                book = openpyxl.load_workbook(folder / "xlsx_files/data.xlsx", data_only=True)
                try:
                    source = {f"{ws.title}!{cell.coordinate}": cell.value
                              for ws in book for row in ws for cell in row}
                    for q, check, reference in zip(queries, checks, expected):
                        self.assertNotIn(q["id"], seen)
                        seen.add(q["id"])
                        self.assertTrue(q["query"].strip())
                        self.assertEqual(q["file_name"], "data.xlsx")
                        actual = format_answer(evaluate(check["expression"], source), check["decimals"])
                        self.assertEqual(actual, check["answer"], q["id"])
                        self.assertEqual(actual, reference["answer"], q["id"])
                finally:
                    book.close()
                total += len(queries)
        self.assertEqual(total, 200)

    def test_independent_edge_and_date_controls(self):
        folder = ROOT / "Test/11_missing_ties_unicode"
        book = openpyxl.load_workbook(folder / "xlsx_files/data.xlsx", data_only=True)
        try:
            ws = book["Records"]
            self.assertEqual(ws["A2"].value, "001")
            self.assertEqual(ws["A2"].data_type, "s")
            values = [ws.cell(r, 4).value for r in range(2, 14)]
            self.assertEqual(sum(v is None for v in values), 3)
            self.assertEqual(sum(v == 0 for v in values), 2)
            self.assertEqual(sum(v for v in values if v is not None), 90)
        finally:
            book.close()
        with (folder / "expected_submission.csv").open(encoding="utf-8") as stream:
            answers = {r["id"]: r["answer"] for r in csv.DictReader(stream)}
        self.assertEqual(answers["11_q003"], 'Công ty "Ánh Dương"')
        self.assertEqual(answers["11_q016"], "005")
        self.assertEqual(answers["11_q017"], "005, 006, 001")
        dates = ROOT / "Test/10_dates_boundaries/expected_submission.csv"
        with dates.open(encoding="utf-8") as stream:
            answers = {r["id"]: r["answer"] for r in csv.DictReader(stream)}
        self.assertEqual([answers[f"10_q{i:03d}"] for i in range(1,7)], ["23", "22", "25", "24", "23", "24"])
        self.assertEqual(answers["10_q007"], "2")
        self.assertEqual(answers["10_q023"], "2025-01-01")

    def test_large_file_keeps_tail_wide_column_and_full_totals(self):
        path = ROOT / "Test/12_large_full_scan/xlsx_files/data.xlsx"
        regions = load_regions(path)
        selected, _ = retrieve(regions, "Secret Value in column AI for ID 2500")
        self.assertTrue(any(r.cells.get((2501,35)) == 32507 for r in selected))
        book = openpyxl.load_workbook(path, data_only=True)
        try:
            ws = book["Big"]
            self.assertEqual((ws.max_row, ws.max_column), (2501,35))
            self.assertEqual(sum(ws.cell(r,3).value for r in range(2,2502)), 1000250)
            self.assertEqual(sum(ws.cell(r,35).value for r in range(2,2502)), 40658750)
            self.assertEqual([ws.cell(r,1).value for r in range(2,2502) if ws.cell(r,4).value == "Rare"],
                             [499,998,1497,1996,2495])
        finally:
            book.close()

    def test_old_smoke_age_reference_uses_actual_admission_age(self):
        folder = ROOT / "Test/01_smoke_replica"
        book = openpyxl.load_workbook(folder / "xlsx_files/mini_people.xlsx", data_only=True)
        try:
            ws = book["People"]
            ages = sorted(((ws.cell(r,6).value-ws.cell(r,5).value).days, ws.cell(r,4).value)
                          for r in range(3,8))
            self.assertEqual(ages[1][1], "Dang")
        finally:
            book.close()
        with (folder / "expected_submission.csv").open(encoding="utf-8") as stream:
            self.assertEqual(next(csv.DictReader(stream))["answer"], "Dang")


if __name__ == "__main__":
    unittest.main()
