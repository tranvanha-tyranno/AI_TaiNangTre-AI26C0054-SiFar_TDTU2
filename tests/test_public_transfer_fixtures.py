"""Independent row-based controls for all 300 new answers, without the oracle DSL."""
import csv
from decimal import Decimal, ROUND_HALF_UP
import json
from pathlib import Path
import random
import unittest

import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def num(v):
    if v is None or str(v).strip() in {"..", "—", "(-8-)", ""}: return None
    text=str(v).replace("$", "").replace(",", "")
    return -Decimal(text[1:-1]) if text.startswith("(") else Decimal(text)


def fmt(v, digits=None):
    if isinstance(v, (tuple, list)): return ", ".join(fmt(x) for x in v)
    if isinstance(v, bool): return "Yes" if v else "No"
    if isinstance(v, str): return v
    v=Decimal(str(v))
    if digits is not None: return f"{v.quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP):.{digits}f}"
    return str(int(v)) if v==int(v) else format(v.normalize(), "f")


def independent_answers(index, book):
    out=[]
    def add(v, d=None): out.append(fmt(v,d))
    def rows(sheet, first, last): return [list(r) for r in book[sheet].iter_rows(min_row=first,max_row=last,values_only=True)]
    def v(sheet, coord): return book[sheet][coord].value
    def n(sheet, coord): return num(v(sheet,coord))
    def mean(values): return sum(values)/len(values)
    def rank(data, column, label=0): return sorted(data,key=lambda r:(-num(r[column]),r[label]))
    if index==13:
        data=rows("People",3,66)
        ages=lambda r:(r[5]-r[4]).days
        completed=lambda r:r[5].year-r[4].year-((r[5].month,r[5].day)<(r[4].month,r[4].day))
        youngest=sorted(data,key=lambda r:(ages(r),r[0]))
        for k in [2,3,5,9]: add(youngest[k-1][3])
        for i in [1,2,19,64]: add(completed(data[i-1]))
        for count in [5,10,15]: add(max(r[6] for r in youngest[:count]))
        for campus in ["Hue","Can Tho","Da Lat"]:
            sub=[r for r in data if r[7]==campus]
            add(mean([num(r[6]) for r in sub]),2)
            add(sum(r[6]>=75 and completed(r)<30 for r in sub))
            add(rank(sub,6)[0][0])
        for threshold in [60,75,90]: add(sum(r[6]>threshold for r in data))
        add(sum(r[5].year==2026 for r in data))
        add(sorted(data,key=lambda r:(-ages(r),r[0]))[1][6])
        add([r[0] for r in rank(data,6)[:3]])
        add(sum(r[6]>100 for r in data)); add(ages(data[-1])); add(len({r[3] for r in data}))
        add(completed(data[1])-completed(data[0]))
    elif index==14:
        w=book["Trade"]
        for j in range(6):
            r=4+3*j
            add(w.cell(r,5).value>w.cell(r,4).value)
            add(w.cell(r,5).value-w.cell(r+1,5).value)
            add(w.cell(26+j,5).value-w.cell(r,5).value)
            add(100*(num(w.cell(r,3).value)-num(w.cell(r,2).value))/num(w.cell(r,2).value),2)
        for c in range(2,6): add(sum(w.cell(r,c).value for r in range(4,22,3)))
        add(sum(w.cell(r,5).value<0 for r in range(6,22,3)))
        add(abs(w["F10"].value-w["F19"].value))
    elif index==15:
        data=rows("Prices",3,22)
        for label in ["Housing","Food and beverages","Services","Commodities","Transport"]:
            r=next(r for r in data if r[1]==label)
            add(num(r[5])-num(r[6]),1); add(num(r[6])-num(r[7]),1); add(num(r[4])-num(r[3]),3)
        sub=[r for r in data if r[0]>0]
        for t in [0,1,2,3]: add(sum(r[6]>t for r in sub))
        start=next(i for i,r in enumerate(data) if r[1]=="Housing")
        housing=[]
        for r in data[start+1:]:
            if r[0]<=1: break
            if r[0]==2: housing.append(r)
        add([r[1] for r in rank(housing,5,1)]); add(mean([num(r[5]) for r in housing]),2)
        for key in [lambda r:abs(num(r[6])), lambda r:abs(num(r[6])-num(r[7])),lambda r:num(r[2]),lambda r:num(r[5])]:
            add(sorted(sub,key=lambda r:(-key(r),r[1]))[0][1])
        for level in [1,2,3]: add(sum(r[0]==level for r in data))
        by={r[1]:r for r in data}
        add(num(by["Housing"][2])-num(by["Food and beverages"][2]),3)
        add(by["Services"][5]>by["Commodities"][5])
    elif index==16:
        w=book["Demographics"]; starts={"Men":4,"Women":11}
        for sex in ["Women","Men"]:
            r=starts[sex]
            for c in range(2,8): add(w.cell(r,c).value-w.cell(r+1,c).value)
            add(mean([num(w.cell(r+4,c).value) for c in range(2,8)]),2)
            add(100*sum(num(w.cell(r+3,c).value) for c in range(2,8))/sum(num(w.cell(r+1,c).value) for c in range(2,8)),2)
            for c in [2,4,7]: add(100*num(w.cell(r+1,c).value)/num(w.cell(r,c).value),2)
        for c in range(2,8): add(w.cell(6,c).value-w.cell(13,c).value)
        add(w["G15"].value>w["G8"].value)
        add(sum(w.cell(r,c).value for r in [6,13] for c in range(2,8)))
    elif index==17:
        data=rows("Employment",3,10); occ=rows("Occupations",2,9)
        for r in data:
            add(r[2]-r[1]); add(100*num(r[4])/(num(r[4])+num(r[5])),2)
        for r in occ[:4]: add(r[1]+r[2])
        for key in [lambda r:r[1],lambda r:r[5],lambda r:r[6],lambda r:r[2]-r[1]]:
            add(sorted(data,key=lambda r:(-key(r),r[0]))[0][0])
        add([r[0] for r in rank(occ[:3],3)])
        add(sum(r[4] for r in data)); add([sum(r[4] for r in data),sum(r[5] for r in data)])
        add(100*num(occ[1][4])/sum(num(x) for x in data[1][1:4]),2)
        add(data[1][2]+data[2][2]); add(data[3][3]>data[4][3])
    elif index==18:
        w=book["Time use"]
        starts=[6,14,22]; activities=["Housework","Childcare","Sleeping","Leisure","Travel"]
        for start in starts:
            for offset in range(4): add(num(w.cell(start+offset,5).value)-num(w.cell(start+offset,2).value),2)
            for offset in range(2): add(num(w.cell(start+offset,7).value)-num(w.cell(start+offset,6).value),1)
            left=num(w.cell(start,2).value); right=num(w.cell(start,5).value)
            add(["not-employed fathers" if right>left else "full-time employed mothers",fmt(abs(left-right),2)])
        for offset,_ in enumerate(activities):
            candidates=[(w.cell(start-2,1).value,num(w.cell(start+offset,2).value)) for start in starts]
            add(sorted(candidates,key=lambda r:(-r[1],r[0]))[0][0])
        for start in starts: add(num(w.cell(start,4).value)+num(w.cell(start+1,4).value),2)
        add(num(w["B10"].value)*60,2)
    elif index==19:
        data=rows("Incidence",3,1102)
        for t in [0,2.5,5,7.5,10]: add(sum(len(r[1])==6 and num(r[2]) is not None and num(r[2])>t for r in data))
        for i in [1,499,550,999,1099,1100]: add(data[i-1][3],2)
        for low,high in [(1,300),(301,700),(701,1100)]: add(rank([r for r in data[low-1:high] if num(r[2]) is not None],2)[0][0])
        quart=book["Quartiles"]; means=[]
        for first in [3,9,15,21]:
            sizes=rows("Quartiles",first+1,first+4)
            chosen=sorted(sizes,key=lambda r:(-(num(r[4])-num(r[3])),r[0]))[0]
            add([chosen[0],fmt(num(chosen[4])-num(chosen[3]),1)])
            add(quart.cell(first+2,5).value>5); add(quart.cell(first+4,4).value,1)
            means.append((quart.cell(first-1,1).value,num(quart.cell(first+4,2).value)))
        add(sum(m is None for _,m in means)); add(sorted([r for r in means if r[1] is not None],key=lambda r:(-r[1],r[0]))[0][0])
        add(sum(len(r[1])==6 and num(r[2])==5 for r in data)); add(sum(num(r[2]) is None for r in data))
    elif index==20:
        data=rows("Productivity",3,102); annual=[r for r in data if r[1]=="Annual"]; quarter=[r for r in data if r[1]=="Quarter"]
        by={r[0]:r for r in annual}
        for c in range(2,6):
            add(rank(annual,c)[0][0]); add(mean([num(r[c]) for r in annual]),2); add(rank(quarter,c)[0][0])
            for year in [2017,2027]: add(num(by[year][c])-num(by[year-1][c]),1)
        add(sum(r[2]>r[3] for r in annual))
        add(sorted(annual,key=lambda r:(-(num(r[2])-num(r[3])),r[0]))[0][0])
        for t in [90,100,110]: add(sum(r[4]>=t for r in annual))
        drops=[(r[0],num(prev[2])-num(r[2])) for prev,r in zip(annual,annual[1:])]
        chosen=sorted(drops,key=lambda r:(-r[1],r[0]))[0]; add([chosen[0],fmt(chosen[1],1)])
        add(sum(2015<=r[0]<=2022 for r in annual))
        for c in [3,5]: add(mean([num(r[c]) for r in annual if 2020<=r[0]<=2027]),2)
        add(v("Productivity","D105"))
    elif index==21:
        segs=["Cloud","Devices","Services"]
        for c in ["B","C","D"]:
            add(sum(n(seg,c+"4") for seg in segs))
            chosen=sorted([(seg,n(seg,c+"3")) for seg in segs],key=lambda r:(-r[1],r[0]))[0]; add(chosen)
            add(sum(n(seg,c+"6") for seg in segs))
        for seg in segs:
            add(100*(n(seg,"B4")-n(seg,"C4"))/n(seg,"C4"),2)
            add(100*n(seg,"B5")/n(seg,"B4"),2)
        for r in range(3,8): add(n("Fair value",f"E{r}")-n("Fair value",f"I{r}"))
        add(n("Fair value","E8")-n("Fair value","I8"))
        for c in ["B","C","D","E"]: add(n("Funding",c+"5")/n("Funding",c+"6"),3)
        add(n("Funding","B5")-n("Funding","D5")); add(n("Funding","B7"))
        add(100*n("Fair value","C8")/n("Fair value","E8"),2)
        add(n("Cloud","B4")>n("Devices","B4")); add(v("Funding","C7"))
    elif index==22:
        care=rows("Care",3,7)
        for label in ["information","medication","counselling"]:
            r=next(r for r in care if r[0]==label); add(num(r[1])+num(r[2]),1)
        add(max(num(r[1])+num(r[2]) for r in care if r[0]!="any" and num(r[1]) is not None),1)
        add(sum(num(r[1]) is None for r in care))
        inc=rows("Thu nhập",3,8)
        for r in inc: add(100*num(r[2])/num(r[1]),1)
        add(max(abs(num(r[1])-num(r[2])) for r in inc)); add(rank(inc,2)[0][0])
        justice=rows("Justice",3,12)
        for a,b in [(0,9),(2,7),(3,8)]: add(100*(num(justice[a][3])-num(justice[b][3]))/num(justice[a][3]),2)
        add(rank(justice,3)[0][0]); add(sum(num(r[1]) for r in justice))
        for r in [3,6,9]:
            add(n("劳动",f"C{r+1}")-n("劳动",f"C{r}"),1)
            add(n("劳动",f"B{r}")-n("劳动",f"C{r}"),1)
        chronic=rows("Santé",2,4)
        moods=[num(r[4]) for r in chronic]
        add(max(moods)-min(moods),1); add(rank(chronic,3)[0][0]); add(mean(moods),2)
        add(sum(num(r[1])==0 for r in care)); add(num(justice[-1][1])/num(justice[-1][3]),3)
        add(num(inc[-2][2])+num(inc[-1][2]))
    return out


class PublicTransferFixturesTest(unittest.TestCase):
    def test_all_300_answers_with_independent_row_calculations(self):
        manifest=json.loads((ROOT/"Test/public_transfer_manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(len(manifest["suites"]),10)
        self.assertEqual(len(manifest["source_inventory"]),25)
        for suite in manifest["suites"]:
            folder=ROOT/"Test"/suite["name"]
            with self.subTest(suite=suite["name"]):
                with (folder/"expected_submission.csv").open(encoding="utf-8",newline="") as stream:
                    expected=list(csv.DictReader(stream))
                q=json.loads((folder/"queries.json").read_text(encoding="utf-8"))
                self.assertEqual(len(q),30)
                self.assertEqual([r["id"] for r in expected],[r["id"] for r in q])
                book=openpyxl.load_workbook(folder/"xlsx_files/data.xlsx",data_only=True)
                try: actual=independent_answers(int(suite["name"][:2]),book)
                finally: book.close()
                self.assertEqual(len(actual),30)
                for query,reference,answer in zip(q,expected,actual):
                    self.assertEqual(answer,reference["answer"],query["id"]+": "+query["query"])

    def test_fresh_values_with_alternate_seed_and_no_public_query_copies(self):
        from scripts.gen_public_transfer_suites import DEFINITIONS, TransferSuite
        public={q["query"] for name in ["Public_data_1","Public_data_2"] for q in json.loads((ROOT/name/"queries.json").read_text(encoding="utf-8"))}
        for name,fn,description,sources in DEFINITIONS:
            with self.subTest(suite=name):
                first=TransferSuite(name,description,sources,random.Random(261016+int(name[:2])),261016)
                second=TransferSuite(name,description,sources,random.Random(271017+int(name[:2])),271017)
                repeat=TransferSuite(name,description,sources,random.Random(261016+int(name[:2])),261016)
                for s in [first,second,repeat]: fn(s)
                self.assertEqual(first.sheets,repeat.sheets)
                self.assertEqual(first.checks,repeat.checks)
                self.assertEqual(len(second.queries),30)
                self.assertNotEqual(first.cells,second.cells)
                self.assertGreater(sum(a["answer"]!=b["answer"] for a,b in zip(first.checks,second.checks)),10)
                self.assertTrue(all(q["query"] not in public for q in first.queries))


if __name__=="__main__": unittest.main()
