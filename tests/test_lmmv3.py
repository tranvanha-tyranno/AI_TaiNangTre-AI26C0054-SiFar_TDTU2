"""Synthetic unit/integration checks only: no reference datasets or paid API."""
import copy
import csv
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import openpyxl
from src.lmmv3.workbook import Workbook, coverage_check, tokens
from src.lmmv3.contracts import check_spec, schema, validate
from src.lmmv3.tools import run, format_answer
from src.lmmv3.pipeline import answer
from src.tools.executor import PlanError

QUESTION='求金额总和，保留两位小数。'


def specification():
    return {'language':'zh','requirements':[{'id':'metric','kind':'metric','description':'Sum amount','quote':'金额总和'}],
            'operation':'sum','output_count':1,'fields':[{'name':'sum','type':'number','decimals':2,'quote':'两位小数'}],
            'separator':', ','ambiguities':[]}


def computation():
    return {'steps':[{'id':'data','op':'read','table':'t1','start_row':2,'end_row':4},
                     {'id':'total','op':'aggregate','input':'data','metric':'sum','expr':{'col':'B'}}],
            'answer_spec':{'answer_type':'number','answer_field':'金额','operation':'sum'},'output':{'step':'total'}}


def response(plan=None):
    return {'action':'plan','coverage':[{'requirement':'metric','table':'t1','range':'B1:B4','rationale':'Amount heading and rows'}],
            'plan':plan or computation()}


class FakeClient:
    prompt_cap=.01;completion_cap=.01
    def __init__(self,replies): self.replies=list(replies);self.purposes=[]
    def complete(self,messages,ledger,purpose):
        self.purposes.append(purpose)
        call=ledger.reserve(messages,100)
        call.update(status='completed',prompt_tokens=10,completion_tokens=10,total_tokens=20,
                    actual_cost_usd=.000001,estimated_cost_usd=.000001,budget_charge_usd=.000001)
        if not self.replies: raise AssertionError('Unexpected additional generation')
        return json.dumps(self.replies.pop(0),ensure_ascii=False)


class V3Test(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.path=Path(self.temp.name)/'arbitrary.xlsx'
        book=openpyxl.Workbook();sheet=book.active;sheet.title='金额'
        sheet.append(['名称','金额','权重']);sheet.append(['甲',10,1]);sheet.append(['乙',20,2]);sheet.append(['丙',30,3]);book.save(self.path);book.close()

    def test_grounded_multilingual_pipeline_has_eight_phases(self):
        client=FakeClient([specification(),response()]);record={}
        value=answer(self.path,QUESTION,client,record,{})
        self.assertEqual(value,'60.00');self.assertTrue(record['validated'])
        self.assertEqual({p['phase'] for p in record['phases']},set(range(1,9)))
        self.assertEqual(client.purposes,['v3_spec','v3_plan'])
        self.assertEqual(record['usage']['api_calls'],2)
        self.assertIn('金额!B2',record['evidence'])

    def test_missing_evidence_triggers_tool_then_plan(self):
        first=response();first['coverage'][0]['table']='unknown'
        inspect={'action':'inspect','requests':[{'op':'read_cell','table':'t1','cell':'B2'}],'missing_requirements':['metric']}
        client=FakeClient([specification(),first,inspect,response()]);record={}
        self.assertEqual(answer(self.path,QUESTION,client,record,{}),'60.00')
        self.assertTrue(record['validation_events'])
        self.assertTrue(any(i['action']=='inspect' for i in record['interactions']))

    def test_unseen_coverage_is_not_claimed_complete(self):
        book=Workbook(self.path);book.remember('t1','B1')
        with self.assertRaisesRegex(PlanError,'Missing requirement evidence'):
            coverage_check(specification(),response()['coverage'],book)

    def test_fabricated_question_quote_rejected(self):
        spec=specification();spec['fields'][0]['quote']='three decimal places'
        with self.assertRaisesRegex(PlanError,'quote'): check_spec(spec,QUESTION)

    def test_inspection_preserves_blank_merged_and_exact_ranges(self):
        book=openpyxl.load_workbook(self.path);sheet=book.active;sheet.merge_cells('A6:C6');sheet['A6']='标题';book.save(self.path);book.close()
        workspace=Workbook(self.path);region=next(r for r in workspace.regions if r.whole_sheet)
        result=workspace.inspect([{'op':'read_range','table':region.id,'range':'A5:C6'}],QUESTION)[0]
        self.assertEqual(len(result['cells']),6);self.assertFalse(result['truncated'])
        merged=next(c for c in result['cells'] if c['cell'].endswith('B6'))
        self.assertEqual(merged['raw'],'');self.assertEqual(merged['value'],'标题');self.assertEqual(merged['merged_anchor'],'A6')
        self.assertEqual(result['cells'][0]['type'],'blank')

    def test_weighted_mean_and_median_preserve_sources(self):
        workspace=Workbook(self.path)
        for operation,expected in [('weighted_mean','23.33333333333333333333333333'),('median','20')]:
            plan=computation()
            plan['steps'][-1]=({'id':'total','op':'weighted_mean','input':'data','value':{'col':'B'},'weight':{'col':'C'}}
                if operation=='weighted_mean' else {'id':'total','op':'aggregate','input':'data','metric':'median','expr':{'col':'B'}})
            result,_=run(plan,workspace.regions);self.assertEqual(result['answer'],expected)
            self.assertIn('金额!B3',result['evidence'])

    def test_percent_change_zero_denominator_and_forward_ref(self):
        plan={'steps':[{'id':'old','op':'cell','table':'t1','row':2,'column':'B'},
                       {'id':'new','op':'cell','table':'t1','row':3,'column':'B'},
                       {'id':'change','op':'percent_change','old':{'ref':'old'},'new':{'ref':'new'}}],
              'answer_spec':{'answer_type':'number','answer_field':'amount','operation':'ratio'},'output':{'step':'change'}}
        regions=Workbook(self.path).regions
        self.assertEqual(run(plan,regions)[0]['answer'],'100')
        plan['steps'][-1]['old']={'value':0}
        with self.assertRaises(PlanError): run(plan,regions)
        plan['steps'][0]={'id':'old','op':'calculate','expr':{'ref':'new'}}
        with self.assertRaisesRegex(PlanError,'preceding scalar'): run(plan,regions)

    def test_mixed_fields_format_numbers_without_changing_labels(self):
        spec=specification();spec['output_count']=2
        spec['fields']=[{'name':'code','type':'label','decimals':None,'quote':'金额'},spec['fields'][0]]
        self.assertEqual(format_answer({'values':['001','14.201118428968']},spec),'001, 14.20')

    def test_reviewer_rejection_never_becomes_submission(self):
        plan=computation();plan['steps']=[{'id':'cell','op':'cell','table':'t1','row':2,'column':'A'}];plan['output']={'step':'cell'}
        spec=specification();spec['fields'][0].update(type='label',decimals=None)
        client=FakeClient([spec,response(plan),{'verdict':'revise','issues':['Wrong scope; inspect more rows']}]);record={}
        with self.assertRaises(PlanError): answer(self.path,QUESTION,client,record,{'v3_max_repairs':0})
        self.assertFalse(record['validated']);self.assertNotIn('answer',record)

    def test_risky_review_accepts_once(self):
        spec=specification();spec['fields'][0].update(type='label',decimals=None)
        plan=computation();plan['steps']=[{'id':'cell','op':'cell','table':'t1','row':2,'column':'A'}];plan['output']={'step':'cell'}
        client=FakeClient([spec,response(plan),{'verdict':'accept','issues':[]}]);record={}
        self.assertEqual(answer(self.path,QUESTION,client,record,{}),'甲')
        self.assertEqual(client.purposes.count('v3_review'),1)

    def test_isolation_and_all_provider_schemas(self):
        from src.planner import tool_schema
        legacy=tool_schema();self.assertNotIn('v3_spec',json.dumps(legacy))
        validate(specification(),tool_schema('v3_spec'));validate(response(),tool_schema('v3_plan'))
        self.assertTrue(tokens('原住民就业率') & tokens('就业率'))

    def test_shared_budget_includes_spec_inspection_and_review(self):
        client=FakeClient([specification()]);record={}
        with self.assertRaises(Exception): answer(self.path,QUESTION,client,record,{'v3_max_api_calls':1})
        self.assertEqual(record['usage']['api_calls'],1);self.assertFalse(record['validated'])

    def test_explicit_locale_and_date_parsing_do_not_guess(self):
        from src.tools.executor import expression, Result
        from decimal import Decimal
        from datetime import date
        self.assertEqual(expression({'parse_number':[{'value':'1.234,50'},{'value':'decimal_comma'}]}, {}, {}).value,Decimal('1234.50'))
        with self.assertRaises(PlanError): expression({'parse_number':[{'value':'1.23,50'},{'value':'decimal_comma'}]}, {}, {})
        self.assertEqual(expression({'parse_date':[{'value':'03/04/2020'},{'value':'%m/%d/%Y'}]}, {}, {}).value,date(2020,3,4))
        self.assertEqual(expression({'parse_date':[{'value':'03/04/2020'},{'value':'%d/%m/%Y'}]}, {}, {}).value,date(2020,4,3))

    def test_usage_keeps_reasoning_separate_from_completion(self):
        from src.usage import UsageLedger
        ledger=UsageLedger();call=ledger.reserve([],1)
        call.update(prompt_tokens=10,completion_tokens=8,total_tokens=18,reasoning_tokens=3)
        self.assertEqual(ledger.summary()['input_plus_reasoning_tokens'],13)
        self.assertEqual(ledger.summary()['total_tokens'],18)
        call['reasoning_tokens']=None
        self.assertIsNone(ledger.summary()['input_plus_reasoning_tokens'])

    def test_cli_routes_v3_and_writes_numeric_ids_in_order(self):
        import main
        queries=Path(self.temp.name)/'queries.json';output=Path(self.temp.name)/'submission.csv';logs=Path(self.temp.name)/'logs'
        queries.write_text(json.dumps([{'id':i,'file_name':self.path.name,'query':QUESTION} for i in [10,2,1]]),encoding='utf-8')
        client=FakeClient([specification(),response()]*3)
        args=['main.py','--mode','llm-v3','--xlsx_dir',self.temp.name,'--queries',str(queries),
              '--output',str(output),'--log-dir',str(logs),'--config',str(Path(self.temp.name)/'unused.yaml')]
        with patch('sys.argv',args),patch.object(main,'get_client',return_value=object()), \
             patch.object(main,'PlannerClient',return_value=client),patch('sys.stdout',new=io.StringIO()):
            main.main()
        with output.open(encoding='utf-8',newline='') as stream: rows=list(csv.DictReader(stream))
        self.assertEqual([r['id'] for r in rows],['1','2','10'])
        self.assertEqual([r['answer'] for r in rows],['60.00']*3)
        self.assertEqual(json.loads((logs/'summary.json').read_text())['pipeline_version'],'lmmv3-8phase-v1')

    def test_dashboard_accepts_v3_without_launching_a_worker(self):
        import dashboard
        queries=[{'id':1,'file_name':self.path.name,'source_file':self.path.name,'query':QUESTION,'dataset':'synthetic'}]
        with patch.object(dashboard,'OUTPUTS',Path(self.temp.name)/'outputs'), \
             patch.object(dashboard,'datasets_for',return_value=[{'id':'synthetic','xlsx_dir':str(self.path.parent)}]), \
             patch.object(dashboard,'queries',return_value=queries),patch.object(dashboard,'reference',return_value={}), \
             patch.object(dashboard.threading,'Thread'),patch.dict(dashboard.JOB,running=False):
            run_id=dashboard.start_run('llm-v3',None,'synthetic')
            metadata=json.loads((dashboard.OUTPUTS/'runs'/run_id/'metadata.json').read_text(encoding='utf-8'))
            self.assertEqual(metadata['mode'],'llm-v3')

    def test_contextual_ambiguity_can_be_resolved_before_planning(self):
        ambiguous=specification();ambiguous['ambiguities']=['Need to inspect amount heading']
        revise={'action':'revise_spec','reason':'Heading confirms requested metric','spec':specification()}
        client=FakeClient([ambiguous,revise,response()]);record={}
        self.assertEqual(answer(self.path,QUESTION,client,record,{}),'60.00')
        self.assertEqual(len(record['specification_revisions']),1)


if __name__=='__main__': unittest.main()
