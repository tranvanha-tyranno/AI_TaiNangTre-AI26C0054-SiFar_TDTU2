"""Workbook-derived coverage hints and conservative checks for unbounded counts."""
import re
from openpyxl.utils import get_column_letter
from .answer_contract import ancestors
from .tools.executor import PlanError, number


def year_value(value):
    if isinstance(value, bool): return None
    try: parsed=number(value)
    except PlanError: return None
    return int(parsed) if parsed==int(parsed) and 1000<=parsed<=2999 else None


def profiles(region):
    years=[]; identifiers=[]
    for col in range(region.start_col,region.end_col+1):
        column=get_column_letter(col)
        entries=[(r,year_value(region.cells.get((r,col)))) for r in range(region.start_row,region.end_row+1)]
        blocks=[]
        for row,year in entries:
            if year is None: continue
            if not blocks or row-blocks[-1][-1][0]>2 or year<=blocks[-1][-1][1]: blocks.append([])
            blocks[-1].append((row,year))
        for block in blocks:
            if col>region.start_col+1 or len(block)<4: continue
            if sum(b[1]-a[1]==1 for a,b in zip(block,block[1:]))<len(block)-2: continue
            years.append({'column':column,'start_row':block[0][0],'end_row':block[-1][0],
                          'first_year':block[0][1],'last_year':block[-1][1],'observations':len(block)})
        # Code profiles require a code-like heading, so rates/years are not
        # silently interpreted as identifiers.
        headings=' '.join(str(region.cells.get((r,col)) or '') for r in range(region.start_row,min(region.end_row,region.start_row+12)+1))
        if not re.search(r'\b(?:code|identifier|id)\b|mã',headings,re.I): continue
        buckets={}
        for row in range(region.start_row,region.end_row+1):
            raw=str(region.cells.get((row,col)) or '').strip()
            if raw.isascii() and raw.isdigit(): buckets.setdefault(len(raw),[]).append(row)
        for digits,rows in sorted(buckets.items()):
            if len(rows)<2: continue
            identifiers.append({'column':column,'digits':digits,'start_row':min(rows),'end_row':max(rows),'observations':len(rows)})
    return {'year_axis_profiles':years,'identifier_profiles':identifiers}


def validate_read_scope(question,plan,regions):
    """Reject partial source scans, never supply an answer or widen a plan."""
    q=question.casefold()
    if not re.search(r'how many|bao nhiêu',q): return
    if re.search(r'\b(?:first|last|top)\s+\d+\s+rows|\brows?\s+\d+\s*(?:to|through|[-:])\s*\d+',q): return
    active=ancestors(plan)
    if not any(s.get('op')=='aggregate' and s.get('metric')=='count' for s in active): return
    digit=re.search(r'\b(\d+)[ -]digit\b',q)
    annual=bool(re.search(r'annual rows|rows labelled with a year|year-labelled rows',q))
    # Explicit time bounds legitimately select only part of an annual block.
    if annual and re.search(r'\b(?:1|2)\d{3}\b',q): annual=False
    if not digit and not annual: return
    index={r.id:r for r in regions}
    reads=[s for s in active if s.get('op')=='read' and s.get('table') in index]
    for read in reads:
        region=index[read['table']]; hints=profiles(region)
        bounds=[(s.get('start_row',index[s['table']].start_row),s.get('end_row',index[s['table']].end_row))
                for s in reads if index[s['table']].sheet==region.sheet]
        def covered(row): return any(start<=row<=end for start,end in bounds)
        for hint in hints['identifier_profiles'] if digit else []:
            if hint['digits']!=int(digit.group(1)): continue
            col=next(c for c in range(region.start_col,region.end_col+1) if get_column_letter(c)==hint['column'])
            missing=[r for r in range(hint['start_row'],hint['end_row']+1)
                     if not covered(r) and re.fullmatch(r'\d{'+str(hint['digits'])+r'}',str(region.cells.get((r,col)) or '').strip(),re.ASCII)]
            if missing:
                raise PlanError(f'Incomplete read scope: omitted identifier rows in {region.sheet}!{hint["column"]}, '
                                f'including {missing[:5]}; matching codes extend through row {hint["end_row"]}. '
                                'Read the full relevant region, then apply the question filters. Preview endpoints are not data endpoints.')
        for hint in hints['year_axis_profiles'] if annual else []:
            start,end=read.get('start_row',region.start_row),read.get('end_row',region.end_row)
            if end<hint['start_row'] or start>hint['end_row']: continue
            col=next(c for c in range(region.start_col,region.end_col+1) if get_column_letter(c)==hint['column'])
            missing=[r for r in range(hint['start_row'],hint['end_row']+1) if not covered(r) and year_value(region.cells.get((r,col))) is not None]
            if missing:
                raise PlanError(f'Incomplete read scope: omitted annual observations in {region.sheet}!{hint["column"]}, '
                                f'including rows {missing[:5]}; annual block spans {hint["start_row"]}:{hint["end_row"]}. '
                                'Include the full annual block and filter actual year labels; do not infer bounds from preview samples.')


def validate_group_identity(question,plan,result,regions):
    """Explicit group-name answers must match observed groups named in question."""
    if not re.search(r'only the group name|group name only',question,re.I): return
    def norm(text): return ' '.join(str(text).casefold().split())
    q=norm(question)
    sections={norm(s['label']) for region in regions for s in region.section_labels()
              if len(norm(s['label']))>=3 and norm(s['label']) in q}
    if len(sections)<2: return  # Evidence is too weak to infer a candidate set.
    aliases=plan.get('output',{}).get('aliases',{})
    allowed=sections|{norm(target) for source,target in aliases.items() if norm(source) in sections and norm(target) in q}
    if any(norm(value) not in allowed for value in result.get('values',[])):
        raise PlanError('Requested group-name identity differs from the repeated detail/occupation label. '
                        'Use the selected row\'s verified _section or its actual section-heading cell, '
                        'not the repeated metric row label. Keep numeric ranking unchanged.')


def validate_numeric_label_filters(plan,regions):
    """Diagnose label-vs-metric axis confusion; do not rewrite the computation."""
    index={r.id:r for r in regions};steps={s['id']:s for s in plan.get('steps',[])}
    def root_read(ident):
        seen=set()
        while ident in steps and ident not in seen:
            seen.add(ident);step=steps[ident]
            if step.get('op')=='read': return step
            ident=step.get('input')
        return None
    def comparisons(expr):
        if not isinstance(expr,dict): return
        for op,args in expr.items():
            if op=='eq' and isinstance(args,list) and len(args)==2:
                yield args
            elif isinstance(args,list):
                for a in args: yield from comparisons(a)
    for step in plan.get('steps',[]):
        if step.get('op')!='filter': continue
        read=root_read(step.get('input'))
        region=index.get(read.get('table')) if read else None
        if region is None: continue
        for args in comparisons(step.get('where')):
            for left,right in (args,args[::-1]):
                if not isinstance(left,dict) or set(left)!={'col'} or not isinstance(right,dict) or set(right)!={'value'}: continue
                label=right['value'];col=left['col']
                if not isinstance(label,str) or len(label.strip())<4: continue
                try: number(label);continue
                except PlanError: pass
                c=next((c for c in range(region.start_col,region.end_col+1) if get_column_letter(c)==col),None)
                if c is None: continue
                vals=[region.cells.get((r,c)) for r in range(read.get('start_row',region.start_row),read.get('end_row',region.end_row)+1)]
                nonblank=[v for v in vals if v not in (None,'')]
                if len(nonblank)<4: continue
                numeric=0
                for v in nonblank:
                    try:number(v);numeric+=1
                    except PlanError:pass
                if numeric!=len(nonblank): continue
                raise PlanError(f'Header label used as row filter: {col} contains numeric metrics in the read range, '
                                f'but eq compares it with nonnumeric label {label!r}. Inspect its header: '
                                'a sector/metric heading selects a COLUMN, not a row label. For a highest-year request, '
                                'filter annual year rows and numeric metrics, sort that metric, take one and output the year axis. '
                                'Do not repeat this impossible text-equals-number filter.')
