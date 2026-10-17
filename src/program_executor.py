"""Restricted, bounded computations over read-only in-memory spreadsheet data.

Generated code has no import, filesystem, network, reflection or dynamic-code
capabilities. Workbook access is through table IDs and is recorded for review.
"""
from __future__ import annotations

import ast
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, localcontext
import itertools
import math
import operator
import re
import statistics
import sys
import time
from types import SimpleNamespace

from openpyxl.utils import get_column_letter
from .tools.executor import PlanError

MAX_ITEMS = 200000


class InvalidDateError(PlanError):
    def __init__(self,value):
        self.value=value
        super().__init__(f'Invalid datetime {str(value)[:80]!r}; if this is a header label, use data(table) before parsing dates')


def num(value):
    """Parse explicit international numeric notation; missing remains missing."""
    if value is None or isinstance(value, bool):
        return None
    raw = str(value).strip().replace('\u2212', '-').replace('\u00a0', '').replace('\u202f', '')
    if raw in {'', '-', '--', '—', '…', 'NA', 'N/A', 'ND', '(-8-)'}:
        return None
    raw = re.sub(r'^[€$£¥]\s*', '', raw).replace(' ', '')
    if raw.startswith('(') and raw.endswith(')'):
        raw = '-'+raw[1:-1]
    raw = raw.rstrip('%')
    if ',' in raw and '.' in raw:
        if raw.rfind(',') > raw.rfind('.'):
            raw = raw.replace('.', '').replace(',', '.')
        else:
            raw = raw.replace(',', '')
    elif ',' in raw:
        raw = raw.replace(',', '') if re.fullmatch(r'[+-]?\d{1,3}(?:,\d{3})+', raw) else raw.replace(',', '.')
    try:
        result = Decimal(raw)
    except InvalidOperation:
        return None
    if not result.is_finite() or abs(result.adjusted()) > 50:
        raise PlanError('Numeric magnitude is outside the computation limit')
    return result


def dt(value):
    if isinstance(value, datetime):
        result = value
    elif isinstance(value, date):
        result = datetime.combine(value, datetime.min.time())
    else:
        raw = str(value).strip()
        try:
            result = datetime.fromisoformat(raw.replace('Z', '+00:00'))
        except ValueError:
            result = None
            for fmt in ('%d/%m/%Y', '%d/%m/%Y %H:%M', '%Y-%m-%d %H:%M:%S', '%Y/%m/%d'):
                try:
                    result = datetime.strptime(raw, fmt)
                    break
                except ValueError:
                    continue
            if result is None:
                raise InvalidDateError(raw)
    if result.tzinfo:
        result = result.astimezone(timezone.utc).replace(tzinfo=None)
    return result


def bounded_items(iterable):
    result = list(itertools.islice(iterable, MAX_ITEMS+1))
    if len(result) > MAX_ITEMS:
        raise PlanError('Computation collection exceeds the item limit')
    return result


def safe_range(*args):
    result = range(*args)
    if len(result) > MAX_ITEMS:
        raise PlanError('range exceeds the item limit')
    return result


def depends_on_data(tree):
    """Reject disconnected literal answers; semantic correctness needs review."""
    graph = {}
    source = '$spreadsheet'

    def dependencies(node):
        names = {n.id for n in ast.walk(node) if isinstance(n, ast.Name) and isinstance(n.ctx, ast.Load)}
        if any(isinstance(n, ast.Call) and isinstance(n.func, ast.Name)
               and n.func.id in {'read', 'data', 'cell', 'headers'} for n in ast.walk(node)):
            names.add(source)
        return names

    def targets(node):
        if isinstance(node, (ast.Subscript, ast.Attribute)):
            base = node.value
            while isinstance(base, (ast.Subscript, ast.Attribute)):
                base = base.value
            return {base.id} if isinstance(base, ast.Name) else set()
        return {n.id for n in ast.walk(node) if isinstance(n, ast.Name) and isinstance(n.ctx, ast.Store)}

    def walk(statements, control=set()):
        for node in statements:
            if isinstance(node, (ast.Assign, ast.AnnAssign, ast.AugAssign, ast.NamedExpr)):
                names = set().union(*(targets(t) for t in node.targets)) if isinstance(node, ast.Assign) else targets(node.target)
                deps = dependencies(node.value) | control
                for name in names:
                    graph.setdefault(name, set()).update(deps)
            elif isinstance(node, (ast.For, ast.While, ast.If)):
                deps = dependencies(node.iter if isinstance(node, ast.For) else node.test) | control
                if isinstance(node, ast.For):
                    for name in targets(node.target):
                        graph.setdefault(name, set()).update(deps)
                walk(node.body, deps)
                walk(node.orelse, deps)
            elif isinstance(node, ast.FunctionDef):
                graph.setdefault(node.name, set()).update(dependencies(node) | control)
                walk(node.body, control)
            elif isinstance(node, ast.Expr) and isinstance(node.value, ast.Call):
                call = node.value
                if isinstance(call.func, ast.Attribute):
                    for name in targets(call.func.value):
                        graph.setdefault(name, set()).update(dependencies(call) | control)
                    if isinstance(call.func.value, ast.Name):
                        graph.setdefault(call.func.value.id, set()).update(dependencies(call) | control)
            elif isinstance(node, ast.Try):
                walk(node.body, control)
                walk(node.orelse, control)
                walk(node.finalbody, control)
                for handler in node.handlers:
                    walk(handler.body, control)
    walk(tree.body)
    reached = {'result'}
    while True:
        expanded = reached | set().union(*(graph.get(name, set()) for name in reached))
        if expanded == reached:
            return source in reached
        reached = expanded


def safe_binary(op, left, right):
    if op == 'Mult':
        for sequence, factor in ((left, right), (right, left)):
            if isinstance(sequence, (str, list, tuple)) and isinstance(factor, int):
                if len(sequence)*max(0, factor) > MAX_ITEMS:
                    raise PlanError('Repeated sequence exceeds the item limit')
    if op == 'Pow' and (abs(right) > 20 or abs(left) > 10**12):
        raise PlanError('Exponent exceeds the computation limit')
    funcs = {'Add': operator.add, 'Sub': operator.sub, 'Mult': operator.mul,
             'Div': operator.truediv, 'FloorDiv': operator.floordiv,
             'Mod': operator.mod, 'Pow': operator.pow}
    result = funcs[op](left, right)
    if isinstance(result, (str, list, tuple, dict, set)) and len(result) > MAX_ITEMS:
        raise PlanError('Computation collection exceeds the item limit')
    return result


def decimal_round(value, ndigits=None):
    if isinstance(value,bool) or not isinstance(value,(int,float,Decimal)):
        raise PlanError('round requires a numeric value')
    if ndigits is not None and (type(ndigits) is not int or not -20 <= ndigits <= 20):
        raise PlanError('round precision must be an integer between -20 and 20')
    rounded = Decimal(str(value)).quantize(Decimal(1).scaleb(-(ndigits or 0)),rounding=ROUND_HALF_UP)
    return int(rounded) if ndigits is None else rounded


def percentile(values, p, method):
    numbers=sorted(num(v) for v in bounded_items(values) if num(v) is not None)
    p=num(p)
    if not numbers or p is None or not 0 <= p <= 100:
        raise PlanError('percentile requires numeric observations and p from 0 to 100')
    if method=='nearest_rank':
        return numbers[max(0,math.ceil(p*len(numbers)/100)-1)]
    if method=='linear':
        pos=p*(len(numbers)-1)/100
        lo=math.floor(pos);hi=math.ceil(pos)
        return numbers[lo]+(numbers[hi]-numbers[lo])*(pos-lo)
    raise PlanError('Specify percentile method nearest_rank or linear as requested')


ATTRIBUTES = {
    'get', 'items', 'values', 'keys', 'append', 'extend', 'add', 'update',
    'setdefault', 'pop', 'copy', 'index', 'count', 'sort', 'reverse',
    'strip', 'lstrip', 'rstrip', 'casefold', 'lower', 'upper', 'replace',
    'startswith', 'endswith', 'split', 'rsplit', 'join', 'find', 'isdigit',
    'zfill', 'partition', 'rpartition', 'removeprefix', 'removesuffix', 'strftime',
    'isdecimal', 'isnumeric', 'quantize', 'to_integral_value', 'to_integral', 'isoformat', 'total_seconds',
    'date', 'year', 'month', 'day', 'hour', 'minute', 'second',
    'ceil', 'floor', 'median', 'mean',
}
FORBIDDEN = (ast.ClassDef, ast.With, ast.AsyncWith,
             ast.AsyncFunctionDef, ast.Await, ast.Yield, ast.YieldFrom,
             ast.Global, ast.Nonlocal, ast.Delete)


class CheckedProgram(ast.NodeTransformer):
    def visit(self, node):
        if isinstance(node, FORBIDDEN):
            raise PlanError(f'Unsupported program construct: {type(node).__name__}')
        return super().visit(node)

    def visit_Name(self, node):
        if node.id.startswith('_') and node.id != '_':
            raise PlanError('Private names are not available in computations')
        return node

    def visit_ImportFrom(self, node):
        # These names already exist in the closed computation environment.
        # Rewrite the declaration; never invoke Python's import machinery.
        available = {'decimal': {'Decimal', 'ROUND_HALF_UP'},
                     'datetime': {'datetime', 'date', 'timedelta'},
                     'collections': {'defaultdict', 'Counter'},
                     'math': {'ceil','floor'}, 'statistics':{'median','mean'}}
        if node.level or node.module not in available or any(a.name not in available[node.module] for a in node.names):
            raise PlanError('Only preprovided computation helpers can be declared')
        replacements = []
        for alias in node.names:
            if alias.asname:
                if alias.asname.startswith('_'):
                    raise PlanError('Private import aliases are not available')
                replacements.append(ast.Assign(targets=[ast.Name(id=alias.asname,ctx=ast.Store())],
                                               value=ast.Name(id=alias.name,ctx=ast.Load())))
        return replacements or ast.copy_location(ast.Pass(), node)

    def visit_Import(self,node):
        statements=[]
        for alias in node.names:
            if alias.name not in {'math','statistics'} or (alias.asname or '').startswith('_'):
                raise PlanError('Only preprovided computation namespaces can be declared')
            if alias.asname:
                statements.append(ast.Assign(targets=[ast.Name(id=alias.asname,ctx=ast.Store())],
                                             value=ast.Name(id=alias.name,ctx=ast.Load())))
        return statements or ast.copy_location(ast.Pass(),node)

    def visit_Attribute(self, node):
        if node.attr not in ATTRIBUTES:
            raise PlanError(f'Attribute {node.attr!r} is not available in computations')
        return self.generic_visit(node)

    def visit_Constant(self, node):
        if isinstance(node.value, (str, bytes)) and len(node.value) > 20000:
            raise PlanError('Program literal exceeds the size limit')
        return node

    def visit_Call(self, node):
        self.generic_visit(node)
        if isinstance(node.func,ast.Attribute) and node.func.attr == 'get':
            if not 1 <= len(node.args) <= 2 or node.keywords:
                raise PlanError('Dictionary get expects key and optional default')
            label=ast.unparse(node.func.value)[:160]
            return ast.copy_location(ast.Call(func=ast.Name(id='_lookup',ctx=ast.Load()),
                args=[node.func.value,ast.Constant(label),*node.args],keywords=[]),node)
        if isinstance(node.func,ast.Attribute) and node.func.attr == 'strftime':
            if len(node.args) != 1 or node.keywords:
                raise PlanError('strftime expects one date-format argument')
            return ast.copy_location(ast.Call(func=ast.Name(id='_format_date',ctx=ast.Load()),
                args=[node.func.value,node.args[0]],keywords=[]),node)
        return node

    def visit_BinOp(self, node):
        self.generic_visit(node)
        name = type(node.op).__name__
        if name not in {'Add', 'Sub', 'Mult', 'Div', 'FloorDiv', 'Mod', 'Pow'}:
            raise PlanError(f'Unsupported binary operator: {name}')
        return ast.copy_location(ast.Call(func=ast.Name(id='_binary', ctx=ast.Load()),
            args=[ast.Constant(name), node.left, node.right], keywords=[]), node)


def execute_program(code, regions, seconds=8):
    if not isinstance(code, str) or not code.strip() or len(code) > 20000:
        raise PlanError('Program must contain 1..20000 characters')
    tree = ast.parse(code, mode='exec')
    if sum(1 for _ in ast.walk(tree)) > 6000:
        raise PlanError('Program exceeds the syntax size limit')
    dependent = depends_on_data(tree)
    tree = ast.fix_missing_locations(CheckedProgram().visit(tree))
    known = {r.id: r for r in regions}
    reads = []
    lookup_audit = {}

    class SpreadsheetRow(dict):
        def __init__(self,*args,aliases=None,**kwargs):
            super().__init__(*args,**kwargs); self._aliases=aliases or {}
        def _key(self,key):
            if key in self:return key
            return self._aliases.get(str(key).strip(),key)
        def __getitem__(self,key):
            key=self._key(key)
            if key not in self: raise KeyError(key)
            return super().__getitem__(key)
        def get(self,key,default=None):
            key=self._key(key)
            if key not in self:
                raise PlanError(f'Unknown row field {key!r}; available fields are Excel letters and header aliases: {list(self)}')
            return super().get(key,default)

    def table(table_id):
        if table_id not in known:
            raise PlanError(f'Unknown table ID {table_id}')
        return known[table_id]

    def read(table_id, start_row=None, end_row=None):
        region = table(table_id)
        lo = region.start_row if start_row is None else int(start_row)
        hi = region.end_row if end_row is None else int(end_row)
        if not region.start_row <= lo <= hi <= region.end_row:
            raise PlanError('read bounds must be inside the observed region')
        result = []
        for row in range(lo, hi+1):
            aliases={}
            for col,heads in region.column_headers().items():
                for head in heads:
                    if head and str(head).strip(): aliases.setdefault(str(head).strip(),col)
            result.append(SpreadsheetRow({**{get_column_letter(col): region.cells.get((row, col))
                             for col in range(region.start_col, region.end_col+1)},
                           '_row': row, '_sheet': region.sheet,
                           'row': row, 'sheet': region.sheet},aliases=aliases))
        reads.append({'table': table_id, 'sheet': region.sheet, 'start_row': lo,
                      'end_row': hi, 'rows': len(result)})
        return result

    def headers(table_id):
        region = table(table_id)
        reads.append({'table':table_id,'sheet':region.sheet,'metadata':'column_headers'})
        return region.column_headers()

    def data(table_id):
        region = table(table_id)
        rows = read(table_id)
        # Header detection over text-only lookup tables is advisory. Without
        # merged heading geometry, exclude only the first observed header.
        heading_rows = set(region.header_rows if region.merged else region.header_rows[:1])
        templates = [{get_column_letter(c):region.cells.get((r,c))
                      for c in range(region.start_col,region.end_col+1)} for r in heading_rows]
        result = [r for r in rows if r['_row'] not in heading_rows
                  and not any(all(r[c] == value for c,value in h.items()) for h in templates)]
        kept_rows = {r['_row'] for r in result}
        reads[-1]['data_rows'] = len(result)
        reads[-1]['excluded_header_rows'] = [r['_row'] for r in rows if r['_row'] not in kept_rows]
        return result

    def cell(table_id, row, column):
        region = table(table_id)
        col = next((c for c in range(region.start_col, region.end_col+1)
                    if get_column_letter(c) == column), None)
        if col is None or not region.start_row <= row <= region.end_row:
            raise PlanError('cell coordinate is outside the observed region')
        reads.append({'table': table_id, 'sheet': region.sheet, 'cell': f'{column}{row}'})
        return region.cells.get((row, col))

    def safe_dict(*args, **kwargs):
        if args and not isinstance(args[0], dict):
            args = (bounded_items(args[0]),)
        return dict(*args, **kwargs)

    def format_date(value, fmt):
        if not isinstance(value,date) or not isinstance(fmt,str) or len(fmt)>1000:
            raise PlanError('strftime requires a date and bounded format string')
        return value.strftime(fmt)

    def lookup(mapping,label,key,default=None):
        if not isinstance(mapping,dict):
            raise PlanError('get requires a dictionary or spreadsheet row')
        if not isinstance(mapping,SpreadsheetRow) and ((id(mapping),label) in lookup_audit or len(lookup_audit)<30):
            audit_key=(id(mapping),label)
            if audit_key not in lookup_audit:
                lookup_audit[audit_key]=(mapping,{'mapping':label,'lookups':0,'matched':0,'missing':0,'missing_key_samples':[]})
            entry=lookup_audit[audit_key][1]
            entry['lookups']+=1
            entry['matched']+=key in mapping
            entry['missing']+=key not in mapping
            if key not in mapping and len(entry['missing_key_samples'])<3:
                entry['missing_key_samples'].append(str(key)[:200])
        return mapping.get(key,default)

    builtin = {'abs': abs, 'all': all, 'any': any, 'bool': bool, 'dict': safe_dict,
               'enumerate': enumerate, 'float': float, 'int': int, 'next': next,
               'iter': iter, 'filter': filter, 'map': map,
               'isinstance': isinstance, 'len': len, 'list': lambda xs=(): bounded_items(xs),
               'max': max, 'min': min, 'range': safe_range, 'round': decimal_round,
               'set': lambda xs=(): set(bounded_items(xs)),
               'sorted': lambda xs, **kw: sorted(bounded_items(xs), **kw),
               'str': str, 'sum': sum, 'tuple': lambda xs=(): tuple(bounded_items(xs)),
               'zip': zip, 'ValueError': ValueError, 'TypeError': TypeError,
               'KeyError': KeyError, 'ZeroDivisionError': ZeroDivisionError}
    env = {'__builtins__': builtin, '_binary': safe_binary, '_format_date':format_date,
           '_lookup':lookup,
           'read': read, 'data': data, 'cell': cell, 'headers': headers,
           'num': num, 'dt': dt, 'Decimal': Decimal, 'datetime': datetime,
           'date': date, 'timedelta': timedelta, 'defaultdict': defaultdict,
           'Counter': Counter, 'search': re.search, 'fullmatch': re.fullmatch,
           'sub': re.sub, 'ROUND_HALF_UP': ROUND_HALF_UP,
           'ceil':math.ceil,'floor':math.floor,'median':statistics.median,
           'mean':statistics.mean,'percentile':percentile,
           'math':SimpleNamespace(ceil=math.ceil,floor=math.floor),
           'statistics':SimpleNamespace(median=statistics.median,mean=statistics.mean)}
    deadline = time.monotonic()+max(.01, seconds)
    ticks = 0

    def trace(frame, event, arg):
        nonlocal ticks
        if frame.f_code.co_filename != '<spreadsheet-program>':
            return None
        ticks += 1
        if ticks > 2000000 or (ticks % 100 == 0 and time.monotonic() > deadline):
            raise PlanError('Computation execution limit exceeded')
        return trace

    previous = sys.gettrace()
    initial_names = set(env)
    try:
        with localcontext() as decimal_context:
            decimal_context.rounding = ROUND_HALF_UP
            sys.settrace(trace)
            exec(compile(tree, '<spreadsheet-program>', 'exec'), env, env)
    finally:
        sys.settrace(previous)
    if time.monotonic() > deadline:
        raise PlanError('Computation execution time exceeded')
    if not reads:
        raise PlanError('Computation must read observed spreadsheet data')
    if 'result' not in env:
        raise PlanError('Assign the computed answer to result')
    value = env['result']
    if value not in ('', None) and not dependent:
        raise PlanError('Final result is disconnected from spreadsheet reads')
    if isinstance(value, dict):
        value = list(value.values())
    if isinstance(value, (list, tuple)):
        if len(value) > 100:
            raise PlanError('Answer exceeds 100 output fields')
        values = list(value)
    else:
        values = [value]
    if any(not isinstance(v, (str, int, float, Decimal, bool, date)) and v is not None for v in values):
        raise PlanError('Answer fields must be scalars, not nested collections')
    if any(isinstance(v, float) and not math.isfinite(v) for v in values):
        raise PlanError('Answer cannot contain non-finite numbers')
    def sample(value,depth=0):
        if depth >= 3:
            return str(value)[:160]
        if isinstance(value,dict):
            return {str(k):sample(v,depth+1) for k,v in list(value.items())[:30]}
        if isinstance(value,(list,tuple)):
            return [sample(v,depth+1) for v in value[:3]]
        if isinstance(value,set):
            return [sample(v,depth+1) for v in sorted(value,key=str)[:3]]
        return value if isinstance(value,(str,int,float,bool,type(None))) else str(value)
    variables=[]
    size=0
    import json
    for name,value in reversed(list(env.items())):
        if name in initial_names or name.startswith('_') or callable(value):
            continue
        entry={'name':name,'sample':sample(value)}
        if isinstance(value,(list,tuple,dict,set)):
            entry['count']=len(value)
        encoded=json.dumps(entry,ensure_ascii=False,default=str)
        if size+len(encoded)>14000:
            continue
        variables.append(entry);size+=len(encoded)
    audits=[]
    for mapping,entry in lookup_audit.values():
        audits.append({**entry,'mapping_size':len(mapping),
                       'mapping_key_samples':[str(k)[:200] for k in list(mapping)[:3]]})
    return {'values': values, 'reads': reads, 'execution_events': ticks,
            'variables':variables,'lookup_audit':audits}


def format_result(execution, decimals=None, separator=', '):
    if decimals is not None and (type(decimals) is not int or not 0 <= decimals <= 10):
        raise PlanError('decimals must be null or an integer from 0 to 10')
    fields = []
    for value in execution['values']:
        if value is None:
            fields.append('')
        elif isinstance(value, bool):
            fields.append('Yes' if value else 'No')
        elif isinstance(value, (int, float, Decimal)):
            numeric = Decimal(str(value))
            if not numeric.is_finite():
                raise PlanError('Answer cannot contain non-finite numbers')
            if decimals is not None:
                fields.append(format(numeric.quantize(Decimal(1).scaleb(-decimals), rounding=ROUND_HALF_UP), f'.{decimals}f'))
            else:
                fields.append(format(numeric, 'f').rstrip('0').rstrip('.') if '.' in format(numeric, 'f') else format(numeric, 'f'))
        elif isinstance(value, date):
            fields.append(value.isoformat())
        else:
            fields.append(value)
    return separator.join(fields)
