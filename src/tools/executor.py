"""Execute declarative plans over complete regions, with cell provenance."""
from __future__ import annotations
from dataclasses import dataclass
from datetime import datetime, date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import math
import re
import time
import unicodedata
from collections import Counter
from contextvars import ContextVar
from openpyxl.utils import get_column_letter


class PlanError(ValueError):
    pass


_TYPED_COMPARISONS = ContextVar('spreadsheet_typed_comparisons', default=False)


@dataclass
class Result:
    value: object
    sources: set


def number(value):
    if isinstance(value, bool) or value is None or isinstance(value,(date,datetime)):
        raise PlanError(f"Expected a number, got {type(value).__name__}")
    try:
        raw = unicodedata.normalize('NFKC', str(value)).strip().replace('\u2212','-')
        # Financial exports frequently store formatted amounts as text.
        # Recognize explicit currency/accounting syntax, never strip arbitrary notes.
        raw = re.sub(r'^[\$\u20ac\u00a3\u00a5]\s*', '', raw)
        if raw.startswith('(') and raw.endswith(')'):
            raw = '-' + re.sub(r'^[\$\u20ac\u00a3\u00a5]\s*', '', raw[1:-1].strip())
        raw = raw.replace('\u00a0',' ').replace('\u202f',' ')
        if re.fullmatch(r'-?\d{1,3}(?: \d{3})+(?:\.\d+)?%?', raw): raw = raw.replace(' ','')
        if ',' in raw and not re.fullmatch(r'[+-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?%?',raw):
            raise PlanError('Ambiguous numeric separators')
        result = Decimal(raw.replace(',','').rstrip('%'))
    except InvalidOperation:
        raise PlanError("Non-numeric or missing cell; filter the correct data rows")
    if not result.is_finite():
        raise PlanError("Non-finite numeric value")
    if abs(result.adjusted())>50 or len(result.as_tuple().digits)>100:
        raise PlanError('Numeric magnitude exceeds execution limit')
    return result


def numeric(result):
    try:
        return number(result.value)
    except PlanError as error:
        cells=', '.join(sorted(result.sources)[:3])
        raise PlanError(f'{error}; at {cells}. Filter is_number on this column before numeric operations when rows include blanks/notes.') from None


def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value,date):
        return value
    for fmt in ('%Y-%m-%d','%Y-%m-%d %H:%M:%S','%d/%m/%Y','%m/%d/%Y'):
        try:
            return datetime.strptime(str(value),fmt).date()
        except ValueError:
            pass
    raise PlanError("Invalid date; select date cells")


def expression(expr, row, results, depth=0):
    if depth > 15:
        raise PlanError("Expression nesting too deep")
    if not isinstance(expr,dict) or len(expr) != 1:
        raise PlanError("Expression must have one key: col, value, ref, or an operator")
    op, arg = next(iter(expr.items()))
    if op == 'col':
        if arg not in row:
            available=[key for key in row if not key.startswith('__')]
            raise PlanError(f"Unknown column {arg}; available columns are {available}. "
                            "After unpivot, label is the column heading and _row is the original row number. "
                            "Filter the source row before unpivot, or use keep_columns to preserve its label.")
        return Result(row[arg], set(row.get('__sources',{}).get(arg,[])))
    if op == 'value':
        if isinstance(arg,(dict,list)):
            raise PlanError("Literal must be a scalar")
        return Result(arg,set())
    if op == 'ref':
        if arg not in results or isinstance(results[arg].value,list):
            scalars=[key for key,result in results.items() if not isinstance(result.value,list)]
            raise PlanError(f'Scalar reference does not exist: {arg}. ref only names preceding scalar steps: {scalars}. '
                            'To read the current row, use col (including col:_row), never ref. For group labels, use actual section-heading cells or derive section before filtering details.')
        return results[arg]
    if not isinstance(arg,list):
        raise PlanError(f"{op} expects an array of expressions")
    if op == 'if':
        if len(arg)!=3: raise PlanError('if requires [condition, true expression, false expression]')
        condition=expression(arg[0],row,results,depth+1)
        if not isinstance(condition.value,bool): raise PlanError('if condition must be boolean')
        branch=expression(arg[1 if condition.value else 2],row,results,depth+1)
        return Result(branch.value,condition.sources|branch.sources)
    if op in {'and','or'}:
        if len(arg)<2: raise PlanError(f'{op} needs at least two conditions')
        sources=set()
        for condition in arg:
            result=expression(condition,row,results,depth+1)
            if not isinstance(result.value,bool): raise PlanError('Logical conditions must be boolean')
            sources.update(result.sources)
            if op=='and' and not result.value: return Result(False,sources)
            if op=='or' and result.value: return Result(True,sources)
        return Result(op=='and',sources)
    args = [expression(a,row,results,depth+1) for a in arg]
    values = [a.value for a in args]
    sources = set().union(*(a.sources for a in args))
    unary = {'abs','number','date','lower','trim','length','is_null','is_number','not'}
    binary = {'sub','div','round','date_diff_days','age_years','eq','ne','gt','ge','lt','le','contains','starts_with'}
    if (op in unary and len(values)!=1) or (op in binary and len(values)!=2):
        raise PlanError(f"Wrong argument count for {op}")
    if op in {'add','sub','mul','div'}:
        if op in {'add','mul'}:
            if len(values)<2: raise PlanError(f'{op} needs at least two values')
            numbers=list(map(numeric,args))
            return Result(sum(numbers,Decimal(0)) if op=='add' else math.prod(numbers),sources)
        a,b = map(numeric,args)
        if op == 'div' and b == 0:
            raise PlanError("Division by zero")
        value = {'add':lambda:a+b,'sub':lambda:a-b,'mul':lambda:a*b,'div':lambda:a/b}[op]()
    elif op == 'parse_number':
        if len(values)!=2 or values[1] not in {'decimal_comma','decimal_dot'}:
            raise PlanError('parse_number requires [text, decimal_comma|decimal_dot]')
        raw=unicodedata.normalize('NFKC',str(values[0])).strip().replace('\u2212','-')
        raw=raw.replace('\u00a0',' ').replace('\u202f',' ')
        decimal,group=(',', '.') if values[1]=='decimal_comma' else ('.', ',')
        pattern=rf'[+-]?(?:\d+|\d{{1,3}}(?:{re.escape(group)}\d{{3}})+|\d{{1,3}}(?: \d{{3}})+)(?:{re.escape(decimal)}\d+)?'
        if not re.fullmatch(pattern,raw): raise PlanError('Number does not match declared locale; no implicit cleanup')
        value=number(raw.replace(group,'').replace(' ','').replace(decimal,'.'))
    elif op == 'parse_date':
        if len(values)!=2 or not isinstance(values[1],str) or len(values[1])>30:
            raise PlanError('parse_date requires [text, explicit strptime format]')
        try: value=datetime.strptime(str(values[0]),values[1]).date()
        except ValueError: raise PlanError('Date does not match declared format') from None
    elif op == 'abs': value = abs(numeric(args[0]))
    elif op == 'number': value = numeric(args[0])
    elif op == 'round':
        digits=number(values[1])
        if digits!=int(digits) or not 0<=digits<=10: raise PlanError('round digits must be an integer 0..10')
        value=numeric(args[0]).quantize(Decimal(1).scaleb(-int(digits)),rounding=ROUND_HALF_UP)
    elif op == 'date': value = as_date(values[0])
    elif op in {'date_diff_days','age_years'}:
        end,start = map(as_date,values)
        value = (end-start).days if op=='date_diff_days' else end.year-start.year-((end.month,end.day)<(start.month,start.day))
    elif op in {'eq','ne','gt','ge','lt','le'}:
        a,b=values
        if isinstance(a,(date,datetime)) or isinstance(b,(date,datetime)):
            try: a,b=as_date(a),as_date(b)
            except PlanError: pass
        if op in {'eq','ne'} and (a is None or b is None):
            return Result((a==b) if op=='eq' else (a!=b),sources)
        if _TYPED_COMPARISONS.get():
            if isinstance(a,bool) or isinstance(b,bool):
                if op in {'eq','ne'} and isinstance(a,bool) and isinstance(b,bool):
                    return Result(a==b if op=='eq' else a!=b,sources)
                raise PlanError('Boolean comparisons require two booleans and eq/ne; booleans are not numbers.')
            if op in {'gt','ge','lt','le'} and isinstance(a,str) and isinstance(b,str):
                raw=[unicodedata.normalize('NFKC',v).strip().replace('\u2212','-') for v in (a,b)]
                numeric_text=all(re.fullmatch(r'[+-]?\d[\d,. ]*',v) for v in raw)
                identifiers=any(re.match(r'[+-]?0\d',v) for v in raw)
                if numeric_text and not identifiers:
                    try:a,b=map(number,raw)
                    except PlanError:
                        raise PlanError('Ambiguous numeric text comparison; declare locale explicitly with parse_number.') from None
        if isinstance(a,(int,float,Decimal)) or isinstance(b,(int,float,Decimal)):
            a,b=number(a),number(b)
        elif op not in {'eq','ne'} and (a is None or b is None):
            raise PlanError("Missing value in ordering comparison")
        value={'eq':lambda:a==b,'ne':lambda:a!=b,'gt':lambda:a>b,'ge':lambda:a>=b,'lt':lambda:a<b,'le':lambda:a<=b}[op]()
    elif op=='contains': value=str(values[1]).casefold() in str(values[0] or '').casefold()
    elif op=='starts_with': value=str(values[0] or '').casefold().startswith(str(values[1]).casefold())
    elif op=='lower': value=str(values[0] or '').casefold()
    elif op=='trim': value=str(values[0] if values[0] is not None else '').strip()
    elif op=='length': value=len(str(values[0] if values[0] is not None else '').strip())
    elif op=='is_null': value=values[0] is None or values[0]==''
    elif op=='is_number':
        try: number(values[0]); value=True
        except PlanError: value=False
    elif op=='not': value=not values[0]
    elif op=='and': value=all(values)
    elif op=='or': value=any(values)
    else: raise PlanError(f"Operator not allowed: {op}")
    return Result(value,sources)


def rows_of(results, name):
    if name not in results or not isinstance(results[name].value,list):
        raise PlanError(f"Row dataset does not exist: {name}")
    return results[name].value


def aggregate(rows, metric, expr, results, missing='error'):
    if metric=='count':
        return Result(len(rows),set().union(*(set(row.get('__filter_sources',[])) | set().union(*(set(v) for v in row.get('__sources',{}).values())) for row in rows)))
    if metric=='distinct_count':
        values=[];sources=set()
        for row in rows:
            value=expression(expr,row,results)
            if value.value not in (None,''):
                values.append(str(value.value));sources.update(value.sources)
        return Result(len(set(values)),sources)
    if metric=='sum' and len(rows)>1:
        # Tables frequently print a grand total above/below its components.
        # Summing both is syntactically valid but double-counts the data.
        labels=[]
        for row in rows:
            label=next((str(v).strip() for col,v in row.items()
                        if not col.startswith('_') and isinstance(v,str) and v.strip()),'')
            labels.append(label)
        total_rows=[(label,row) for label,row in zip(labels,rows)
                    if re.match(r'^(?:grand total|subtotal|total|tổng cộng|tổng)(?:\b|\s|,|:)',label,re.I)]
        if total_rows and len(total_rows)<len(rows):
            locations=[]
            for label,row in total_rows[:2]:
                try: refs=sorted(expression(expr,row,results).sources)
                except PlanError: refs=[]
                locations.append({'label':label,'row':row.get('_row'),'cells':refs[:4]})
            raise PlanError(f'Sum includes a Total/subtotal row together with detail rows. Total rows: {locations}. '
                            'For the overall total, use cell on that matching printed Total alone (no new sum needed). '
                            'For detail sums, exclude Total and overlapping groups from read/filter bounds.')
    selected=[]
    for row in rows:
        val=expression(expr,row,results)
        if val.value in (None,'') and missing=='skip':
            continue
        selected.append(Result(number(val.value),val.sources))
    if not selected:
        raise PlanError("No numeric rows selected")
    values=[v.value for v in selected]
    if metric=='sum': value=sum(values,Decimal(0))
    elif metric=='mean': value=sum(values,Decimal(0))/len(values)
    elif metric=='min': value=min(values)
    elif metric=='max': value=max(values)
    elif metric=='median':
        ordered=sorted(values);middle=len(ordered)//2
        value=ordered[middle] if len(ordered)%2 else (ordered[middle-1]+ordered[middle])/Decimal(2)
    else: raise PlanError(f"Unknown aggregate {metric}")
    return Result(value,set().union(*(v.sources for v in selected)))


def format_value(value, fmt):
    decimals=fmt.get('decimals')
    if decimals is not None and not isinstance(value,(bool,date,datetime)):
        try: number(value)
        except PlanError: return format_value(value,{})
        if not isinstance(decimals,int) or not 0<=decimals<=10:
            raise PlanError("decimals must be 0..10")
        return format(number(value).quantize(Decimal(1).scaleb(-decimals),rounding=ROUND_HALF_UP),f'.{decimals}f')
    if isinstance(value,bool):
        return 'Yes' if value else 'No'
    if isinstance(value,(int,float,Decimal)):
        v=number(value)
        return str(int(v)) if v==v.to_integral_value() else format(v,'f').rstrip('0').rstrip('.')
    if isinstance(value,(datetime,date)):
        return value.date().isoformat() if isinstance(value,datetime) else value.isoformat()
    if value is None:
        raise PlanError("Final answer is missing")
    cleaned=re.sub(r'(?<=\w)-\s+(?=\w)','',str(value))
    return re.sub(r'\s+',' ',cleaned).strip()


def execute(plan, regions, *, typed_comparisons=False):
    token=_TYPED_COMPARISONS.set(typed_comparisons)
    try:
        return _execute(plan,regions)
    finally:
        _TYPED_COMPARISONS.reset(token)


def _execute(plan, regions):
    if not isinstance(plan,dict) or not isinstance(plan.get('steps'),list) or not 1<=len(plan['steps'])<=25:
        raise PlanError("Plan requires 1..25 steps")
    final=plan.get('output')
    if not isinstance(final,dict): raise PlanError('output must be an object')
    unsupported=set(final)-{'step','columns','separator','decimals','aliases'}
    if unsupported:
        raise PlanError(f'Unsupported output fields: {sorted(unsupported)}. '
                        'limit/offset belong in a take STEP after sorting, not in output. '
                        'Totals require an aggregate or sum_columns STEP; output only formats the final result.')
    index={r.id:r for r in regions}
    results,trace,row_samples={},[],{}
    materialized_cells=0
    for step in plan['steps']:
        started=time.monotonic()
        if not isinstance(step,dict): raise PlanError("Step must be an object")
        name,op=step.get('id'),step.get('op')
        allowed={
            'read':{'table','start_row','end_row'},'cell':{'table','row','column'},
            'sum_columns':{'table','row','columns'},
            'concat':{'inputs'},'collect':{'inputs'},'record':{'fields'},'unpivot':{'input','columns','label_table','label_row','keep_columns'},
            'lag':{'input','column','source','offset'},'fill_down':{'input','columns'},
            'filter':{'input','where'},'derive':{'input','column','expr'},'sort':{'input','by'},
            'take':{'input','offset','limit'},'select':{'input','columns'},
            'aggregate':{'input','metric','expr','missing'},'group':{'input','by','metric','expr','column','missing'},
            'join':{'left','right','left_on','right_on','cardinality'},'calculate':{'expr','input'}}
        if op not in allowed: raise PlanError(f'Tool not allowed: {op}')
        required={
            'read':['table'],'cell':['table','row','column'],'sum_columns':['table','row','columns'],
            'concat':['inputs'],'collect':['inputs'],'record':['fields'],'unpivot':['input','columns'],
            'lag':['input','column','source'],'fill_down':['input','columns'],
            'filter':['input','where'],'derive':['input','column','expr'],'sort':['input','by'],
            'take':['input'],'select':['input','columns'],'aggregate':['input','metric'],
            'group':['input','by','metric'],'join':['left','right','left_on','right_on'],'calculate':['expr']}
        missing=[field for field in required[op] if field not in step]
        if missing:
            datasets=[key for key,result in results.items() if isinstance(result.value,list)]
            raise PlanError(f'Step {name} ({op}) is missing required fields {missing}. '
                            f'For input, explicitly name the intended preceding row dataset: {datasets}. '
                            'Return the complete plan with those fields present.')
        unexpected=set(step)-allowed[op]-{'id','op'}
        if unexpected: raise PlanError(f'Unsupported fields for {op}: {sorted(unexpected)}. '
                                      f'Allowed fields: {sorted(allowed[op])}. Each operation is a separate step; '
                                      'put where in filter and by in sort, never on derive/aggregate.')
        if not isinstance(name,str) or not re.fullmatch(r'[a-zA-Z][a-zA-Z0-9_]{0,30}',name) or name in results:
            raise PlanError("Invalid/duplicate step id")
        if op=='read':
            region=index.get(step.get('table'))
            if region is None: raise PlanError("Unknown retrieved table ID")
            start,end=step.get('start_row',region.start_row),step.get('end_row',region.end_row)
            if not isinstance(start,int) or not isinstance(end,int) or not region.start_row<=start<=end<=region.end_row:
                raise PlanError("Read bounds outside retrieved table")
            value=[]
            sections=iter(region.section_labels()); upcoming=next(sections,None); section=None
            for r in range(start,end+1):
                while upcoming and upcoming['row']<=r:
                    section=upcoming;upcoming=next(sections,None)
                anchor=f'{region.sheet}!{get_column_letter(region.start_col)}{r}'
                row={'_row':r,'_sheet':region.sheet,'__sources':{'_sheet':[anchor],'_row':[anchor]}}
                row.update(_section=section['label'] if section else None,
                           _section_row=section['row'] if section else None)
                section_sources=[f"{region.sheet}!{section['column']}{section['row']}"] if section else []
                row['__sources'].update(_section=section_sources,_section_row=section_sources)
                for c in range(region.start_col,region.end_col+1):
                    col=get_column_letter(c)
                    row[col]=region.cells.get((r,c))
                    row['__sources'][col]=[f"{region.sheet}!{col}{r}"]
                value.append(row)
            result=Result(value,set())
        elif op=='cell':
            region=index.get(step.get('table'))
            if region is None: raise PlanError('Unknown table for cell lookup')
            col=step.get('column'); row=step.get('row')
            columns={get_column_letter(c):c for c in range(region.start_col,region.end_col+1)}
            if col not in columns or not isinstance(row,int) or not region.start_row<=row<=region.end_row: raise PlanError('Cell outside selected table')
            result=Result(region.cells.get((row,columns[col])),{f'{region.sheet}!{col}{row}'})
        elif op=='sum_columns':
            region=index.get(step.get('table'))
            row=step.get('row'); columns=step.get('columns')
            available={get_column_letter(c):c for c in range(region.start_col,region.end_col+1)} if region else {}
            if region is None or not isinstance(row,int) or not region.start_row<=row<=region.end_row:
                raise PlanError('Row outside selected table')
            if not isinstance(columns,list) or not columns or len(set(columns))!=len(columns):
                raise PlanError('sum_columns requires distinct numeric column letters')
            values=[]; sources=set()
            for col in columns:
                if col not in available: raise PlanError('Unknown sum_columns column')
                source=f'{region.sheet}!{col}{row}'
                values.append(numeric(Result(region.cells.get((row,available[col])),{source})))
                sources.add(source)
            result=Result(sum(values,Decimal(0)),sources)
        elif op=='concat':
            inputs=step.get('inputs')
            if not isinstance(inputs,list) or not 2<=len(inputs)<=10: raise PlanError('concat needs 2..10 datasets')
            value=[]; evidence=set()
            for input_id in inputs:
                value.extend(rows_of(results,input_id)); evidence.update(results[input_id].sources)
            result=Result(value,evidence)
        elif op=='collect':
            inputs=step.get('inputs')
            if not isinstance(inputs,list) or not 1<=len(inputs)<=20: raise PlanError('collect needs scalar step IDs in answer order')
            value=[]
            for input_id in inputs:
                if input_id not in results or isinstance(results[input_id].value,list): raise PlanError('collect inputs must be scalar cells/calculations')
                scalar=results[input_id]
                value.append({'value':scalar.value,'__sources':{'value':sorted(scalar.sources)}})
            result=Result(value,set())
        elif op=='record':
            fields=step.get('fields')
            if not isinstance(fields,dict) or not fields or len(fields)>20: raise PlanError('record requires named scalar step IDs in fields')
            row={'__sources':{}}
            for column,expr in fields.items():
                if not isinstance(column,str) or column.startswith('_'): raise PlanError('Invalid record column')
                if isinstance(expr,str): expr={'ref':expr}
                scalar=expression(expr,{},results)
                if not scalar.sources and isinstance(scalar.value,str) and scalar.value.strip():
                    # A literal label is acceptable only if tools can resolve
                    # it to real cells. Numeric literals cannot become answers.
                    label=' '.join(scalar.value.split())
                    scalar.sources={f'{r.sheet}!{get_column_letter(c)}{row}' for r in regions
                                    for (row,c),v in r.cells.items()
                                    if isinstance(v,str) and ' '.join(v.split())==label}
                if not scalar.sources: raise PlanError('record fields must depend on spreadsheet scalars; labels can be cells or conditionals on cells')
                row[column]=scalar.value; row['__sources'][column]=sorted(scalar.sources)
            result=Result([row],set())
        elif op in {'lag','fill_down'}:
            rows=rows_of(results,step.get('input')); value=[]
            if op=='lag':
                column,source,offset=step.get('column'),step.get('source'),step.get('offset',1)
                if not isinstance(column,str) or column.startswith('_') or not isinstance(offset,int) or not 1<=offset<=1000:
                    raise PlanError('lag requires new column and offset 1..1000')
                for i,row in enumerate(rows):
                    if source not in row or column in row: raise PlanError('lag source missing or column already exists')
                    prior=rows[i-offset] if i>=offset else None
                    value.append({**row,column:prior[source] if prior else None,
                        '__sources':{**row['__sources'],column:prior['__sources'].get(source,[]) if prior else []}})
            else:
                columns=step.get('columns'); carried={}
                if not isinstance(columns,list) or not columns: raise PlanError('fill_down requires columns')
                for row in rows:
                    cloned={**row,'__sources':dict(row['__sources'])}
                    for column in columns:
                        if column not in row: raise PlanError('Unknown fill_down column')
                        if row[column] not in (None,''): carried[column]=(row[column],row['__sources'].get(column,[]))
                        elif column in carried: cloned[column],cloned['__sources'][column]=carried[column]
                    value.append(cloned)
            result=Result(value,set())
        elif op=='unpivot':
            rows=rows_of(results,step.get('input'))
            columns=step.get('columns')
            if not isinstance(columns,list) or not columns: raise PlanError('unpivot requires columns')
            keep_columns=step.get('keep_columns',[])
            if not isinstance(keep_columns,list) or any(c in {'column','label','value','_row'} or c.startswith('__') for c in keep_columns):
                raise PlanError('Invalid unpivot keep_columns')
            label_row=step.get('label_row'); label_region=index.get(step.get('label_table'))
            # Horizontal tables often keep category names in the nearest
            # textual header row of the same retrieved region. Infer that
            # row only when the planner omitted label_table/label_row; this
            # remains coordinate based and cannot invent a label.
            if label_region is None:
                source_region=None
                for candidate_region in regions:
                    if not rows: continue
                    sample=rows[0]
                    valid=True
                    for col in columns:
                        col_index=next((c for c in range(candidate_region.start_col,candidate_region.end_col+1)
                                    if get_column_letter(c)==col),-1)
                        if (sample.get('_row'),col_index) not in candidate_region.cells:
                            valid=False;break
                    if valid:
                        source_region=candidate_region;break
                if source_region is not None:
                    row_numbers=sorted({r for r,c in source_region.cells if r < min((row.get('_row',source_region.end_row) for row in rows),default=source_region.end_row)})
                    candidates=[]
                    for candidate in row_numbers:
                        values=[source_region.cells.get((candidate,next((c for c in range(source_region.start_col,source_region.end_col+1) if get_column_letter(c)==col),-1))) for col in columns]
                        if values and all(isinstance(v,str) and v.strip() for v in values): candidates.append(candidate)
                    if candidates:
                        label_row=max(candidates); label_region=source_region
            value=[]
            for row in rows:
                if any(c not in row for c in keep_columns): raise PlanError('Unknown unpivot keep column')
                for col in columns:
                    if col not in row: raise PlanError('Unknown unpivot column')
                    label=col; label_sources=[]
                    if label_row is not None:
                        if label_region is None: raise PlanError('unpivot label_table missing')
                        c=next((c for c in range(label_region.start_col,label_region.end_col+1) if get_column_letter(c)==col),None)
                        if not isinstance(label_row,int) or c is None: raise PlanError('Label cell outside table')
                        source=next((r for r in index.values() if r.sheet==label_region.sheet and (label_row,c) in r.cells),None)
                        if source is None: raise PlanError('Label cell not available in retrieved sheet')
                        label=source.cells.get((label_row,c))
                        label_sources=[f'{label_region.sheet}!{col}{label_row}']
                    value.append({**{c:row[c] for c in keep_columns},
                                  'column':col,'label':label,'value':row[col],'_row':row.get('_row'),
                                  '__sources':{**{c:row['__sources'].get(c,[]) for c in keep_columns},
                                               'label':label_sources or row['__sources'].get(col,[]),'value':row['__sources'].get(col,[])}})
            result=Result(value,set())
        elif op in {'filter','derive','sort','take','select','aggregate','group'}:
            rows=rows_of(results,step.get('input'))
            if op=='filter':
                value=[]
                checked_sources=set()
                for row in rows:
                    predicate=expression(step.get('where'),row,results)
                    checked_sources.update(predicate.sources)
                    if not isinstance(predicate.value,bool): raise PlanError("Filter condition must be boolean")
                    if predicate.value:
                        value.append({**row,'__filter_sources':sorted(set(row.get('__filter_sources',[]))|predicate.sources)})
                result=Result(value,checked_sources)
            elif op=='derive':
                column=step.get('column')
                if not isinstance(column,str) or column.startswith('__') or column=='_row': raise PlanError("Invalid derived column")
                value=[]
                for row in rows:
                    if column in row: raise PlanError("Derived column must not overwrite existing column")
                    derived=expression(step.get('expr'),row,results)
                    if not derived.sources: raise PlanError('Derived columns must depend on spreadsheet cells, not invented literals')
                    value.append({**row,column:derived.value,'__sources':{**row['__sources'],column:sorted(derived.sources)}})
                result=Result(value,set())
            elif op=='sort':
                by=step.get('by')
                if not isinstance(by,list) or not by or len(by)>5: raise PlanError("sort.by requires 1..5 keys")
                value=list(rows)
                for key in reversed(by):
                    col=key.get('column')
                    if any(col not in row or row[col] is None for row in rows): raise PlanError("Sort column missing or null")
                    value.sort(key=lambda row:row[col],reverse=bool(key.get('descending',False)))
                evidence=set().union(*(set(row['__sources'].get(key['column'],[])) for row in rows for key in by))
                result=Result(value,evidence)
            elif op=='take':
                offset,limit=step.get('offset',0),step.get('limit',1)
                if not isinstance(offset,int) or not isinstance(limit,int) or offset<0 or not 1<=limit<=1000: raise PlanError("Invalid take offset/limit")
                if offset>=len(rows): raise PlanError("Requested rank outside selected rows")
                result=Result(rows[offset:offset+limit],set())
            elif op=='select':
                columns=step.get('columns')
                if not isinstance(columns,list) or not columns: raise PlanError("select requires columns")
                value=[]
                for row in rows:
                    if any(c not in row for c in columns): raise PlanError("Unknown selected column")
                    value.append({**{c:row[c] for c in columns},'__sources':{c:row['__sources'].get(c,[]) for c in columns}})
                result=Result(value,set())
            elif op=='aggregate':
                result=aggregate(rows,step.get('metric'),step.get('expr'),results,step.get('missing','error'))
            else:
                by=step.get('by'); metric=step.get('metric'); output=step.get('column','value')
                if not isinstance(by,list) or not by: raise PlanError("Group requires key columns")
                buckets={}
                for row in rows:
                    if any(c not in row for c in by): raise PlanError("Unknown group column")
                    buckets.setdefault(tuple(row[c] for c in by),[]).append(row)
                value=[]
                for keys,bucket in buckets.items():
                    val=aggregate(bucket,metric,step.get('expr'),results,step.get('missing','error'))
                    value.append({**dict(zip(by,keys)),output:val.value,'__sources':{**{c:bucket[0]['__sources'].get(c,[]) for c in by},output:sorted(val.sources)}})
                result=Result(value,set())
        elif op=='join':
            left,right=rows_of(results,step.get('left')),rows_of(results,step.get('right'))
            lk,rk=step.get('left_on'),step.get('right_on')
            if len(left)*len(right)>1000000: raise PlanError("Join exceeds execution limit")
            if any(lk not in row for row in left) or any(rk not in row for row in right):
                raise PlanError('Unknown join key')
            cardinality=step.get('cardinality','many_to_one')
            if cardinality not in {'many_to_one','one_to_one','one_to_many','many_to_many'}:
                raise PlanError('Invalid join cardinality')
            left_keys=Counter(row[lk] for row in left if row[lk] is not None)
            right_keys=Counter(row[rk] for row in right if row[rk] is not None)
            audit={'cardinality':cardinality,'left_rows':len(left),'right_rows':len(right),
                   'left_duplicate_keys':sum(n>1 for n in left_keys.values()),
                   'right_duplicate_keys':sum(n>1 for n in right_keys.values()),
                   'unmatched_left_rows':sum(row[lk] is None or row[lk] not in right_keys for row in left),
                   'unmatched_right_rows':sum(row[rk] is None or row[rk] not in left_keys for row in right)}
            if cardinality in {'many_to_one','one_to_one'} and audit['right_duplicate_keys']:
                raise PlanError(f'Join cardinality {cardinality}: duplicate right lookup keys would multiply rows. '
                                f'Audit: {audit}. Aggregate/filter the lookup table, or explicitly choose the true relationship.')
            if cardinality in {'one_to_many','one_to_one'} and audit['left_duplicate_keys']:
                raise PlanError(f'Join cardinality {cardinality}: duplicate left keys. Audit: {audit}')
            value=[]
            for a in left:
                if lk not in a: raise PlanError("Unknown left join key")
                for b in right:
                    if rk not in b: raise PlanError("Unknown right join key")
                    if a[lk] is not None and a[lk]==b[rk]:
                        row={'__sources':{}}
                        for prefix,record in [('left.',a),('right.',b)]:
                            for col,val in record.items():
                                if col.startswith('__'): continue
                                row[prefix+col]=val
                                row['__sources'][prefix+col]=record['__sources'].get(col,[])
                        value.append(row)
                        if len(value)>10000: raise PlanError('Join result exceeds 10000 rows')
            result=Result(value,set(results[step['left']].sources)|set(results[step['right']].sources))
            audit.update(output_rows=len(value),expansion_factor=round(len(value)/len(left),4) if left else 0)
        elif op=='calculate':
            result=expression(step.get('expr'),{},results)
            if not result.sources: raise PlanError("Calculation must depend on spreadsheet data")
        else: raise PlanError(f"Tool not allowed: {op}")
        if step.get('input') in results:
            result.sources.update(results[step['input']].sources)
        if isinstance(result.value,list):
            materialized_cells += sum(len(row.get('__sources',{})) for row in result.value)
            if materialized_cells>500000: raise PlanError('Plan exceeds intermediate memory limit')
        results[name]=result
        if isinstance(result.value,list):
            def sample_value(value):
                if isinstance(value,(datetime,date)): return value.isoformat()
                if isinstance(value,Decimal): return str(value)
                if isinstance(value,str): return value[:160]
                return value
            row_samples[name]=[{k:sample_value(v) for k,v in row.items() if not k.startswith('__')}
                               for row in result.value[:3]]
        coverage={}
        if _TYPED_COMPARISONS.get() and isinstance(result.value,list) and op in {'filter','take','group'}:
            sections={str(row['_section']) for row in result.value if row.get('_section') not in (None,'')}
            source_rows=[row['_row'] for row in result.value if isinstance(row.get('_row'),int)]
            coverage={'observed_coverage':{'section_labels':sorted(sections)[:20],
                'section_count':len(sections),'section_labels_truncated':len(sections)>20,
                'source_row_min':min(source_rows,default=None),'source_row_max':max(source_rows,default=None),
                'note':'Computed from all result rows. Does not establish completeness against the question.'}}
        trace.append({'id':name,'op':op,'row_count':len(result.value) if isinstance(result.value,list) else None,
                      'scalar':format_value(result.value,{}) if not isinstance(result.value,list) else None,
                      'elapsed_seconds':time.monotonic()-started,**coverage,**({'join_audit':audit} if op=='join' else {})})
    final=plan.get('output',{})
    result=results.get(final.get('step'))
    if result is None: raise PlanError("Final step does not exist")
    sources=set(result.sources)
    if isinstance(result.value,list):
        columns=final.get('columns')
        if not isinstance(columns,list) or not columns: raise PlanError("Row output requires columns")
        if not result.value: raise PlanError("Final selection is empty")
        values=[]; value_sources=[]
        for row in result.value:
            for col in columns:
                if col not in row: raise PlanError(f"Unknown output column {col}")
                values.append(format_value(row[col],final))
                value_sources.append(sorted(row['__sources'].get(col,[])))
                sources.update(row['__sources'].get(col,[]))
                sources.update(row.get('__filter_sources',[]))
    else:
        values=[format_value(result.value,final)]
        value_sources=[sorted(result.sources)]
    aliases=final.get('aliases',{})
    if not isinstance(aliases,dict): raise PlanError('Output aliases must be a mapping')
    if any(not isinstance(key,str) or not isinstance(value,str) for key,value in aliases.items()): raise PlanError('Aliases must map strings to strings')
    values=[aliases.get(value,value) for value in values]
    separator=final.get('separator',', ')
    if not isinstance(separator,str): raise PlanError('separator must be a string')
    answer=separator.join(values)
    if not sources or not answer or len(answer)>4000: raise PlanError("Final output must be nonempty and grounded in Excel")
    full_evidence=sorted(sources)
    return {'answer':answer,'trace':trace,'evidence':full_evidence[:100],'evidence_full':full_evidence,
            'evidence_count':len(sources),'value_count':len(values),'values':values,'validated':True,
            'value_sources':value_sources,'row_samples':row_samples}
