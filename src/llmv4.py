"""Incremental v4 features; no reference answers, extra LLM stages or new tools."""
import copy
import json
import re
import unicodedata
from openpyxl.utils import get_column_letter, range_boundaries
from .answer_contract import infer_answer_contract, validate_answer_contract, validate_answer_field
from .workbook_index import Region, text
from .tools.executor import PlanError

VERSION = 'llmv4-contract-context-v1'


def infer_contract(question):
    q = unicodedata.normalize('NFKC', question).casefold()
    contract = infer_answer_contract(question)
    contract['contract_version'] = VERSION
    # Explicit return fields take precedence over metric names inside filters.
    year = bool(re.search(r'chỉ (?:trả lời|ghi) (?:mỗi )?năm\b|\bnăm nào\b', q))
    label = year or bool(re.search(r'(?:nhóm|ngành|khu vực|loại|tỉnh|thành phố|sheet) nào\b|chỉ trả lời tên\b', q))
    if label:
        contract.update(reduction=None, answer_type='label', answer_field='year' if year else None)
    pairs = [('Yes', 'No'), ('Có', 'Không'), ('Oui', 'Non'), ('是', '否'), ('はい', 'いいえ'), ('네', '아니요'), ('예', '아니요')]
    boolean_pair = None
    for yes, no in pairs:
        if re.search(re.escape(yes.casefold()) + r'\s*(?:/|hoặc|hay|or|ou|或|还是|か|または|또는)\s*' + re.escape(no.casefold()), q):
            boolean_pair = (yes, no)
            break
    if boolean_pair is None and re.search(r'\bcó\b.+\bkhông\s*\?', q):
        boolean_pair = ('Có', 'Không')
    if boolean_pair:
        contract.update(yes_no=True, kind='scalar', count=1, reduction=None,
                        answer_type='boolean', answer_field=None, decimals=None, integer=False,
                        boolean_labels=list(boolean_pair))
    template = re.search(r'(?:trả lời (?:dạng|theo dạng)|answer (?:as|in (?:the )?format))\s*[`\"“]([^`\"”]+)[`\"”]', q)
    if template and ',' in template.group(1):
        fields = [s.strip() for s in template.group(1).split(',')]
        label_words = r'nhóm|tên|ngành|sheet|quy mô|group|name|class|category|industry|sector|label|code|\bid\b'
        numeric = [i for i, field in enumerate(fields) if not re.search(label_words, field)]
        contract.update(kind='tuple', count=len(fields), fields=fields, numeric_positions=numeric,
                        answer_type=None, answer_field=None, reduction=None, yes_no=False)
        contract['component_decimals'] = {i: (0 if re.search(r'\byear\b|\bnăm\b', fields[i]) else contract.get('decimals')) for i in numeric}
    return contract


def validate_contract(contract, plan, result, question):
    check = dict(contract)
    if contract.get('kind') == 'tuple':
        check.update(decimals=None, integer=False)
    validate_answer_contract(check, plan, result, question)
    if contract.get('kind') == 'tuple':
        for i in contract['numeric_positions']:
            digits = contract.get('component_decimals', {}).get(i)
            value = result['values'][i]
            if digits is not None:
                pattern = rf'-?\d+\.\d{{{digits}}}' if digits else r'-?\d+'
                if not re.fullmatch(pattern, value):
                    raise PlanError(f'Output component {i+1} requires {digits} decimal places, not {value}.')
            elif contract.get('integer') and not re.fullmatch(r'-?\d+', value):
                raise PlanError(f'Output component {i+1} requires an integer.')


def validate_field(contract, plan, result, regions):
    checked = copy.deepcopy(plan)
    spec = checked.get('answer_spec') or {}
    field = str(spec.get('answer_field', '')).casefold().strip()
    if contract.get('answer_field') == 'year' and field in {'year', 'năm', 'année', '年', '년도'}:
        spec['answer_field'] = 'year'
        # Years may be declared numeric without becoming a summed metric.
        if spec.get('answer_type') == 'number':
            spec['answer_type'] = 'label'
    elif spec.get('answer_type') == 'label' and re.search(r'^(?:mã\b|số thứ tự\b|no\.?$)', field):
        spec['answer_field'] = 'code'
    validate_answer_field(contract, checked, result, regions)


def format_answer(answer, contract, record):
    labels = contract.get('boolean_labels')
    if labels and answer in ('Yes', 'No'):
        rendered = labels[answer == 'No']
        record['answer_formatting'] = {'kind': 'boolean_language', 'source': answer, 'output': rendered}
        return rendered
    return answer


class RegionV4(Region):
    def standalone_labels(self):
        # Expanded merged cells are copies of one anchor, not multiple labels.
        copies = set()
        for merged in self.merged:
            c1, r1, c2, r2 = range_boundaries(merged)
            copies.update((r, c) for r in range(r1, r2+1) for c in range(c1, c2+1) if (r, c) != (r1, c1))
        sections = []
        for row in range(self.start_row, self.end_row+1):
            values = [(c, self.cells.get((row,c))) for c in range(self.start_col,self.end_col+1)
                      if (row,c) not in copies and self.cells.get((row,c)) not in (None,'')]
            if len(values) == 1 and values[0][0] <= self.start_col+1 and isinstance(values[0][1], str) and len(values[0][1]) <= 240:
                sections.append({'row':row, 'column':get_column_letter(values[0][0]), 'label':values[0][1].strip()})
        return sections

    def section_labels(self):
        return [s for s in self.standalone_labels() if s['row'] > max(self.header_rows, default=self.start_row-1)]

    def preview(self, query, max_rows=45):
        preview = super().preview(query, max_rows)
        # Preserve titles omitted by column_headers(), including explicit units.
        candidates = [r for r in self.header_rows if r < max(self.header_rows, default=self.start_row)]
        preview['document_context'] = [{'row':r, 'cells':{get_column_letter(c):text(self.cells[(r,c)])[:500]
            for c in range(self.start_col,self.end_col+1) if (r,c) in self.cells}} for r in candidates[:8]]
        return preview


def regions_v4(regions):
    # Do not mutate baseline regions or their values/provenance.
    result = []
    for r in regions:
        region = RegionV4(**r.__dict__)
        labels = {s['row'] for s in region.standalone_labels()}
        headers = list(region.header_rows)
        # Only trim trailing standalone labels when a multi-cell header and
        # further standalone groups provide structural evidence for blocks.
        if any(row not in labels for row in headers) and any(row > max(headers, default=0) for row in labels):
            while headers and headers[-1] in labels:
                headers.pop()
        region.header_rows = headers
        result.append(region)
    return result


def prompt(request, contract):
    request = copy.deepcopy(request)
    request.insert(0, {'role':'system', 'content':
        'LLMv4: use this output contract, including every requested component: ' + json.dumps(contract, ensure_ascii=False) +
        '\nReturn internal booleans as Yes/No; the program applies explicit output-language labels. '
        'A metric named Total/Mean inside a filter is not a request to sum/average. '
        'Use _section only when supplied section_labels supports it; it carries the actual group anchor. '
        'Read document_context for units; convert only when source and target units are explicit. '
        'For conditional counts filter first, then count. date_diff_days and age_years take [end, start]. '
        'A tuple needs all components in output; use scalar collect/record or selected row columns. '
        'Repair the named error; do not repeat an unchanged failed plan.'})
    return request
