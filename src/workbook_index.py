"""Coordinate-preserving workbook regions and lexical table retrieval."""
from __future__ import annotations
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
import re
import unicodedata
import zipfile
import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries


def text(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value) if value is not None else ""


def words(value):
    value = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode().lower()
    return set(re.findall(r"[a-z0-9]+", value)) - {"the", "a", "of", "in", "and", "only", "answer", "what", "with", "for"}


@dataclass
class Region:
    id: str
    sheet: str
    start_row: int
    end_row: int
    start_col: int
    end_col: int
    cells: dict
    context: list
    merged: list
    header_rows: list
    whole_sheet: bool = False
    formats: dict = None

    @property
    def bounds(self):
        return f"{get_column_letter(self.start_col)}{self.start_row}:{get_column_letter(self.end_col)}{self.end_row}"

    def column_headers(self):
        full_titles=[(r1,r2) for c1,r1,c2,r2 in map(range_boundaries,self.merged)
                     if c1<=self.start_col and c2>=self.end_col and c2>c1 and r1==r2]
        return {get_column_letter(c):list(dict.fromkeys(text(self.cells[(r,c)]) for r in self.header_rows
                if (r,c) in self.cells and self.cells[(r,c)] is not None and not any(r1<=r<=r2 for r1,r2 in full_titles)))
                for c in range(self.start_col,self.end_col+1)}

    def row_dimensions(self):
        dimensions=[]
        paths=self.column_headers()
        for c in range(self.start_col,self.end_col+1):
            headers=paths[get_column_letter(c)]
            if not any(re.fullmatch(r'(?:sectors?|industry|industries|category|categories|regions?|departments?|products?|groups?)',label.strip(),re.I) for label in headers): continue
            rows=[];totals=[]
            for r in range(self.start_row,self.end_row+1):
                label=self.cells.get((r,c))
                if r in self.header_rows or not isinstance(label,str) or not label.strip(): continue
                if not any(isinstance(self.cells.get((r,col)),(int,float)) and not isinstance(self.cells.get((r,col)),bool)
                           for col in range(self.start_col,self.end_col+1) if col!=c): continue
                if re.match(r'^(?:total|subtotal|grand total)\b',label,re.I): totals.append(r)
                else: rows.append(r)
            dimensions.append({'column':get_column_letter(c),'column_index':c,'headers':headers,'data_rows':rows,'total_rows':totals})
        return dimensions

    def preview(self, query, max_rows=45):
        all_rows = list(range(self.start_row, self.end_row + 1))
        tokens = words(query)
        ranked = sorted(all_rows, key=lambda r: -len(tokens & words(' '.join(text(self.cells.get((r,c))) for c in range(self.start_col,self.end_col+1)))))
        if len(all_rows) <= max_rows:
            included = all_rows
        else:
            included = sorted(set(all_rows[:10] + all_rows[-3:] + ranked[:max_rows-13]))
        # Retain merged ranges explicitly while sending the anchor text once.
        # Expanded values remain available to the executor.
        merged_bounds=[range_boundaries(region) for region in self.merged]
        def preview_cells(row):
            return {get_column_letter(c):text(self.cells[(row,c)])[:500]
                    for c in range(self.start_col,self.end_col+1) if (row,c) in self.cells
                    and not any(c1<=c<=c2 and r1<=row<=r2 and (row,c)!=(r1,c1)
                                for c1,r1,c2,r2 in merged_bounds)}
        return {"table_id": self.id, "sheet": self.sheet, "range": self.bounds,
                "context": self.context, "header_rows": self.header_rows,
                "column_hierarchy":self.column_headers(),
                "row_dimensions":[{'column':d['column'],'headers':d['headers'],'row_count':len(d['data_rows']),
                                   'data_ranges':[[span[0],span[-1]] for span in spans(d['data_rows'])],
                                   'distinct_labels':list(dict.fromkeys(text(self.cells[(r,d['column_index'])]) for r in d['data_rows']))[:40]}
                                  for d in self.row_dimensions()],
                "whole_sheet": self.whole_sheet,
                "row_labels": [{'row':r,'label':' | '.join(dict.fromkeys(text(self.cells.get((r,c)))[:300] for c in range(self.start_col,min(self.start_col+2,self.end_col+1)) if self.cells.get((r,c)) is not None))} for r in all_rows[:500]],
                "percent_cells": [f'{get_column_letter(c)}{r}' for (r,c),fmt in (self.formats or {}).items() if '%' in fmt][:100],
                "merged_ranges": self.merged, "total_rows": len(all_rows),
                "preview_is_complete": len(included) == len(all_rows),
                "columns": [get_column_letter(c) for c in range(self.start_col,self.end_col+1)],
                "rows": [{"row":r,"cells":preview_cells(r)} for r in included]}


def spans(indices):
    result = []
    for item in sorted(indices):
        if not result or item > result[-1][-1] + 1:
            result.append([item])
        else:
            result[-1].append(item)
    return result


def load_regions(path: Path, max_cells=200000):
    # Bound input expansion before openpyxl creates a workbook object.
    with zipfile.ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
            raise ValueError("Workbook decompressed size exceeds parser safety limit")
    book = openpyxl.load_workbook(path, data_only=True, keep_links=False)
    regions, occupied_total = [], 0
    try:
        for sheet in book.worksheets:
            if sheet.max_row * sheet.max_column > max_cells:
                raise ValueError("Worksheet dimensions exceed configured cell limit")
            cells = {(cell.row,cell.column):cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None}
            if not cells:
                continue
            formats={(cell.row,cell.column):cell.number_format for row in sheet.iter_rows() for cell in row if cell.value is not None and '%' in cell.number_format}
            merges = list(sheet.merged_cells.ranges)
            for merge in merges:
                anchor = cells.get((merge.min_row, merge.min_col))
                if anchor is not None:
                    for r in range(merge.min_row,merge.max_row+1):
                        for c in range(merge.min_col,merge.max_col+1):
                            cells.setdefault((r,c),anchor)
            occupied_total += len(cells)
            if occupied_total > max_cells:
                raise ValueError("Workbook exceeds configured occupied cell limit")
            row_indices={r for r,c in cells}; col_indices={c for r,c in cells}
            row_values={}
            for (r,c),value in cells.items(): row_values.setdefault(r,[]).append((c,value))
            contexts=[]
            for r,vals in sorted(row_values.items()):
                unique=list(dict.fromkeys(text(value)[:500] for c,value in sorted(vals)))
                if r<=8 or (len(unique)<=2 and not any(isinstance(value,(int,float,date,datetime)) for c,value in vals)):
                    contexts.append((r,f'row {r}: '+ ' | '.join(unique)))
            full_headers=[]
            for r in sorted(row_indices)[:8]:
                if any(isinstance(v,(int,float,date,datetime)) for (rr,c),v in cells.items() if rr==r): break
                full_headers.append(r)
            if len(spans(row_indices))>1:
                regions.append(Region(f't{len(regions)+1}',sheet.title,min(row_indices),max(row_indices),min(col_indices),max(col_indices),cells,[],[str(m) for m in merges],full_headers,True,formats))
            # Blank rows/columns split candidate regions. Keep preceding sheet
            # headings as context, including section labels outside each region.
            for row_span in spans({r for r,c in cells}):
                if len(regions)>64: raise ValueError('Too many candidate regions for current parser limits')
                columns = {c for r,c in cells if row_span[0] <= r <= row_span[-1]}
                for col_span in spans(columns):
                    selected = {(r,c):v for (r,c),v in cells.items() if row_span[0] <= r <= row_span[-1] and col_span[0] <= c <= col_span[-1]}
                    if not selected:
                        continue
                    context = [content for r,content in contexts if r<row_span[0]]
                    # Header inference is advisory only. Tools always retain all
                    # rows, and the planner supplies explicit data row bounds.
                    headers = []
                    for r in row_span[:8]:
                        vals = [selected.get((r,c)) for c in col_span]
                        if any(isinstance(v,(int,float,date,datetime)) for v in vals if v is not None):
                            break
                        headers.append(r)
                    merged = [str(m) for m in merges if m.min_row <= row_span[-1] and m.max_row >= row_span[0] and m.min_col <= col_span[-1] and m.max_col >= col_span[0]]
                    regions.append(Region(f"t{len(regions)+1}",sheet.title,row_span[0],row_span[-1],col_span[0],col_span[-1],selected,context[-20:],merged,headers,False,formats))
    finally:
        book.close()
    return regions


def retrieve(regions, query, top_k=4):
    tokens = words(query)
    scored = []
    for region in regions:
        headers = ' '.join(dict.fromkeys(text(v)[:500] for (r,c),v in region.cells.items() if r in region.header_rows or r <= region.start_row+5))
        content_words=set()
        for value in region.cells.values(): content_words.update(words(text(value)[:500]))
        score = 4*len(tokens & words(region.sheet+' '+headers)) + 2*len(tokens & words(' '.join(region.context))) + len(tokens & content_words)
        if region.whole_sheet: score += 100
        if not any(isinstance(v,(int,float,date,datetime)) for v in region.cells.values()): score-=100
        scored.append((score, region))
    ranked = sorted(scored,key=lambda pair:-pair[0])
    return [region for score,region in ranked[:top_k]], [{"table_id":r.id,"sheet":r.sheet,"range":r.bounds,"score":score,"selected":index<top_k} for index,(score,r) in enumerate(ranked)]
