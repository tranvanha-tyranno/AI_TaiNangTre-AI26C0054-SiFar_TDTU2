"""Read-only tools: original cells, advisory structure and bounded observations."""
from collections import Counter
from dataclasses import replace
from datetime import date, datetime
import re
import unicodedata
import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries, coordinate_to_tuple
from ..workbook_index import load_regions, text
from ..tools.executor import number, PlanError
from ..exploration import observe


def tokens(value):
    normalized=unicodedata.normalize('NFKC',text(value)).casefold()
    parts=re.findall(r'[^\W_]+',normalized,re.UNICODE)
    # Character pairs preserve retrieval signals for scripts without word spaces.
    return set(parts)|{part[i:i+2] for part in parts if any('\u3040'<=c<='\ud7af' for c in part) for i in range(len(part)-1)}


class Workbook:
    def __init__(self,path,max_cells=200000):
        self.regions=load_regions(path,max_cells)
        # Years may be numeric header cells. Keep this inference advisory and
        # confined to v3; tools retain all original rows for explicit reads.
        for region in self.regions:
            headers=[]
            for row in range(region.start_row,min(region.start_row+12,region.end_row+1)):
                values=[region.cells.get((row,c)) for c in range(region.start_col,region.end_col+1)]
                present=[v for v in values if v not in (None,'')]
                if not present: break
                years=[v for v in present if type(v) in (int,float) and float(v).is_integer() and 1800<=v<=2200]
                numeric=[v for v in present if type(v) in (int,float)]
                if any(isinstance(v,(date,datetime)) for v in present) or len(numeric)>len(years): break
                if all(isinstance(v,str) for v in present) or years: headers.append(row)
                else: break
            region.header_rows=headers
        self.index={r.id:r for r in self.regions}
        self.raw={};self.formats={};self.formula_warnings=[]
        book=openpyxl.load_workbook(path,data_only=False,keep_links=False)
        try:
            for sheet in book:
                for row in sheet:
                    for cell in row:
                        if cell.value is not None:
                            self.raw[(sheet.title,cell.row,cell.column)]=cell.value
                            if len(self.raw)>max_cells: raise ValueError('Workbook original cells exceed configured limit')
                            self.formats[(sheet.title,cell.row,cell.column)]=cell.number_format
                            if cell.data_type=='f' and len(self.formula_warnings)<100:
                                self.formula_warnings.append(f'{sheet.title}!{cell.coordinate}: formula uses cached value; cache may be absent or stale')
        finally: book.close()
        self.observed=[]

    def cell(self,region,row,col):
        key=(region.sheet,row,col)
        anchor=None
        for merged in region.merged:
            c1,r1,c2,r2=range_boundaries(merged)
            if r1<=row<=r2 and c1<=col<=c2:
                anchor=f'{get_column_letter(c1)}{r1}';break
        value=region.cells.get((row,col))
        kind='blank' if value is None else 'boolean' if isinstance(value,bool) else 'date' if isinstance(value,(date,datetime)) else 'number' if isinstance(value,(int,float)) else 'text'
        return {'cell':f'{region.sheet}!{get_column_letter(col)}{row}',
                'raw':text(self.raw.get(key)),'value':text(value),'type':kind,
                'number_format':self.formats.get(key),'merged_anchor':anchor}

    def structure(self,region):
        cols=[]
        headers=region.column_headers()
        for col in range(region.start_col,region.end_col+1):
            letter=get_column_letter(col)
            counts=Counter('blank' if region.cells.get((row,col)) is None else type(region.cells[(row,col)]).__name__
                           for row in range(region.start_row,region.end_row+1))
            cols.append({'column':letter,'header_path':headers[letter],'types':dict(counts)})
            cols[-1]['number_format_examples']=list(dict.fromkeys(self.formats.get((region.sheet,r,col),'General')
                 for r in range(region.start_row,region.end_row+1) if (region.sheet,r,col) in self.raw))[:8]
        annotations=[]
        for (row,col),value in region.cells.items():
            if isinstance(value,str) and re.search(r'percent|percentage|thousand|million|billion|unit|note|footnote|pourcentage|millier|nghìn|triệu|đơn vị|chú thích|单位|千|百万|百分|단위|백만|천',value,re.I):
                annotations.append({'cell':f'{region.sheet}!{get_column_letter(col)}{row}','text':value[:500]})
                if len(annotations)>=40: break
        return {'table_id':region.id,'sheet':region.sheet,'range':region.bounds,'columns':cols,
                'unit_note_candidates_advisory':annotations,
                'merged_ranges':region.merged,'header_rows_advisory':region.header_rows,
                'context':region.context,'section_candidates':region.section_labels()[:60],
                'warnings':['Header/section/total detection is advisory. Read original cells to confirm scope and units.']}

    def catalog(self,question):
        q=tokens(question)
        scored=[]
        for region in self.regions:
            headings=' '.join(v for path in region.column_headers().values() for v in path)
            score=4*len(q & tokens(region.sheet+' '+headings))+len(q & tokens(' '.join(text(v) for v in region.cells.values())))
            scored.append((score,region))
        return [{'table_id':r.id,'sheet':r.sheet,'range':r.bounds,'score':score,
                 'whole_sheet':r.whole_sheet} for score,r in sorted(scored,key=lambda item:-item[0])]

    def remember(self,table,bounds):
        item={'table':table,'range':bounds}
        if item not in self.observed: self.observed.append(item)

    def preview(self,region,question,limit):
        view=region.preview(question,limit)
        # Region.preview is retained for older variants; v3 ranks Unicode text.
        all_rows=list(range(region.start_row,region.end_row+1));query=tokens(question)
        ranked=sorted(all_rows,key=lambda row:-len(query & tokens(' '.join(text(region.cells.get((row,c))) for c in range(region.start_col,region.end_col+1)))))
        mandatory=list(dict.fromkeys(region.header_rows+all_rows[:5]+all_rows[-2:]))
        included=sorted(set((mandatory+ [r for r in ranked if r not in mandatory])[:max(limit,len(mandatory))]))
        view['rows']=[{'row':r,'cells':{get_column_letter(c):text(region.cells.get((r,c)))[:500]
                       for c in range(region.start_col,region.end_col+1) if (r,c) in region.cells},
                      'original_cells':{get_column_letter(c):text(self.raw.get((region.sheet,r,c)))[:500]
                       for c in range(region.start_col,region.end_col+1) if (region.sheet,r,c) in self.raw}}
                      for r in included]
        view['preview_is_complete']=len(included)==len(all_rows)
        view['cell_text_limit']=500;view['structure_is_advisory']=True
        for row in view['rows']:
            self.remember(region.id,f'{get_column_letter(region.start_col)}{row["row"]}:{get_column_letter(region.end_col)}{row["row"]}')
        return view

    def inspect(self,requests,question):
        if not isinstance(requests,list) or not 1<=len(requests)<=3: raise PlanError('Inspection requires 1..3 requests')
        result=[]
        for request in requests:
            op=request.get('op');region=self.index.get(request.get('table'))
            if op in {'table_schema','read_cell','profile'}:
                if region is None: raise PlanError('Unknown inspection table')
                if op=='table_schema':
                    result.append({'op':op,'structure':self.structure(region)})
                    for row in region.header_rows: self.remember(region.id,f'{get_column_letter(region.start_col)}{row}:{get_column_letter(region.end_col)}{row}')
                elif op=='read_cell':
                    try: row,col=coordinate_to_tuple(request['cell'])
                    except (ValueError,KeyError): raise PlanError('Invalid cell coordinate') from None
                    if not (region.start_row<=row<=region.end_row and region.start_col<=col<=region.end_col): raise PlanError('Cell outside table')
                    cells=[self.cell(region,r,c) for r in range(max(region.start_row,row-1),min(region.end_row,row+1)+1)
                           for c in range(max(region.start_col,col-1),min(region.end_col,col+1)+1)]
                    for cell in cells: self.remember(region.id,cell['cell'].rsplit('!',1)[1])
                    result.append({'op':op,'cells':cells,'truncated':False})
                else:
                    letter=request['column'];valid={get_column_letter(c):c for c in range(region.start_col,region.end_col+1)}
                    if letter not in valid: raise PlanError('Unknown profile column')
                    col=valid[letter];values=[region.cells.get((r,col)) for r in range(region.start_row,region.end_row+1)]
                    numeric=0
                    for value in values:
                        try: number(value);numeric+=1
                        except PlanError: pass
                    result.append({'op':op,'table':region.id,'column':letter,'row_count':len(values),
                                   'blank_count':sum(v is None or v=='' for v in values),'numeric_count':numeric,
                                   'includes_headers':True,'values_not_returned':True})
            else:
                observations,_=observe([request],self.regions,question,60)
                for item in observations:
                    if 'preview' in item:
                        preview=item['preview'];table=preview['table_id'];original=self.index[table]
                        if op=='read_range':
                            # Exact reads return every requested cell, including blanks and originals.
                            c1,r1,c2,r2=range_boundaries(request['range'])
                            item['cells']=[self.cell(original,r,c) for r in range(r1,r2+1) for c in range(c1,c2+1)]
                            item['truncated']=False;self.remember(table,request['range'])
                        else:
                            item['preview']=self.preview(original,question,60)
                    for hit in item.get('matches',[]): self.remember(hit['table'],f'{get_column_letter(hit["column_index"])}{hit["row"]}')
                result.extend(observations)
        return result


def coverage_check(spec,coverage,workbook,question=None):
    """Check grounding, not semantic truth: LLM still owns meaning alignment."""
    required={r['id'] for r in spec['requirements']};covered=set();issues=[]
    question_grounded={r['id'] for r in spec['requirements'] if r['kind']=='output' and question is not None and r['quote'] in question}
    for item in coverage:
        if item['requirement'] not in required:
            raise PlanError(f'Coverage uses unknown requirement {item["requirement"]!r}; valid IDs: {sorted(required)}')
        region=workbook.index.get(item['table'])
        if region is None: raise PlanError('Coverage uses unknown table')
        try: c1,r1,c2,r2=range_boundaries(item['range'])
        except (ValueError,TypeError): raise PlanError('Invalid coverage A1 range') from None
        if None in (c1,r1,c2,r2) or not(region.start_col<=c1<=c2<=region.end_col and region.start_row<=r1<=r2<=region.end_row):
            raise PlanError(f'Coverage outside table: {item["range"]!r}; {region.id} bounds are {region.bounds}. Cite an observed range inside these bounds; do not clamp or invent evidence.')
        if (c2-c1+1)*(r2-r1+1)>600:
            raise PlanError('Coverage evidence must cite <=600 observed cells, not an uninspected full data range')
        observed_bounds=[range_boundaries(o['range']) for o in workbook.observed if o['table']==region.id]
        seen=all(any(a<=col<=c and b<=row<=d for a,b,c,d in observed_bounds)
                 for row in range(r1,r2+1) for col in range(c1,c2+1))
        if not seen: issues.append(f'{item["requirement"]}: no observed cells in {region.id}!{item["range"]}')
        else: covered.add(item['requirement'])
    issues.extend('Missing requirement evidence: '+key for key in sorted(required-covered-question_grounded))
    if issues: raise PlanError('; '.join(issues))
    return {'complete':True,'requirements':len(required),'grounding_checked':True,
            'semantic_alignment':'model_assessed_not_proven','coverage':coverage,'question_grounded':sorted(question_grounded)}
