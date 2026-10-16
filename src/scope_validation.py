"""Conservative checks for scalar plans that mistake one category for all."""
import re
from .tools.executor import PlanError, number
from openpyxl.utils import get_column_letter


def validate_aggregate_scope(question, plan, regions):
    # This check applies only to explicit all-category requests and scalar
    # cell arithmetic. Row-based aggregates have a different axis/scope.
    if not re.search(r'\b(?:across|over|among|for|in) all (?:industries|sectors|categories)\b|tất cả (?:ngành|lĩnh vực)',question,re.I):
        return
    steps={step['id']:step for step in plan['steps']}
    leaves=[]; visited=set()
    def refs(expr):
        if not isinstance(expr,dict): return []
        if 'ref' in expr: return [expr['ref']]
        return [ref for value in expr.values() if isinstance(value,list) for child in value for ref in refs(child)]
    def walk(ident):
        if ident in visited: return True
        visited.add(ident)
        step=steps.get(ident,{})
        # A sum down ONE row is just a single category cell. Do not let this
        # syntactic wrapper bypass the horizontal coverage check.
        if step.get('op')=='aggregate' and step.get('metric')=='sum':
            source=steps.get(step.get('input'),{})
            expr=step.get('expr',{})
            if source.get('op')=='read' and source.get('start_row')==source.get('end_row') and isinstance(source.get('start_row'),int) and set(expr)=={'col'}:
                leaves.append({'op':'cell','table':source['table'],'row':source['start_row'],'column':expr['col']})
                return True
        if step.get('op') in {'cell','sum_columns'}:
            leaves.append(step); return True
        if step.get('op')=='calculate': return all(walk(ref) for ref in refs(step.get('expr')))
        if step.get('op')=='collect': return all(walk(ref) for ref in step.get('inputs',[]))
        return False
    if not walk(plan.get('output',{}).get('step')): return
    index={r.id:r for r in regions}
    for step in leaves:
        region=index[step['table']]; row=step['row']
        numeric_columns=[]
        for col in range(region.start_col,region.end_col+1):
            try: number(region.cells.get((row,col)))
            except PlanError: continue
            numeric_columns.append(col)
        if len(numeric_columns)<3: continue
        # Find the nearest row above the selected data with distinct text
        # labels over its numeric columns (not a merged year or title).
        headings=None
        for header in range(row-1,region.start_row-1,-1):
            values={c:region.cells.get((header,c)) for c in numeric_columns}
            if all(isinstance(v,str) and v.strip() for v in values.values()) and len(set(values.values()))>=3:
                headings=values; break
        if headings is None:
            # A request that explicitly says "across all industries" still
            # has an axis requirement even when the candidate fragment starts
            # below the header rows. Reject a single category cell/column;
            # the planner must use every numeric industry column or a printed
            # all-industry total.
            if re.search(r'\bacross\s+all\s+industries\b',question,re.I):
                used={step['column']} if step['op']=='cell' else set(step['columns'])
                if len(used) < len(numeric_columns):
                    raise PlanError(f'All-industry scope is incomplete at {region.sheet} row {row}: '
                                    f'used {sorted(used)}, available numeric industry columns {len(numeric_columns)}. '
                                    'Use sum_columns for every industry column or the explicit all-industry total.')
            continue
        used={step['column']} if step['op']=='cell' else set(step['columns'])
        # An explicit total alongside components can coexist with narrower
        # qualifiers (e.g. "in all industries ... and in agriculture"). Do
        # not infer summation across such a mixed total/detail/rate axis.
        if any(re.search(r'\b(?:total|all industries|all sectors|all categories)\b',
                         re.sub(r'\s+',' ',headings[c]),re.I) for c in numeric_columns):
            continue
        # Multiple cells on the same row may be summed by a calculate step.
        used.update(leaf['column'] for leaf in leaves if leaf['op']=='cell'
                    and leaf['row']==row and index[leaf['table']].sheet==region.sheet)
        missing=[get_column_letter(c) for c in numeric_columns if get_column_letter(c) not in used]
        if missing:
            raise PlanError(f'All-category scope is incomplete at {region.sheet} row {row}: '
                f'missing numeric category columns {missing}. Use sum_columns for every category, '
                'or an explicit Total column. Do not treat the first category as the total.')
