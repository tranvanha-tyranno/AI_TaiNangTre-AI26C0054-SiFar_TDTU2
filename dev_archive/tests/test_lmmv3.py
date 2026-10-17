"""Synthetic unit/integration checks only: no reference datasets or paid API."""
import copy
import csv
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import openpyxl
from src.lmmv3.workbook import Workbook, coverage_check, tokens
from src.lmmv3.contracts import check_spec, schema, validate
from src.lmmv3.tools import run, format_answer
from src.lmmv3.pipeline import answer
from src.tools.executor import PlanError

QUESTION='求金额总和，保留两位小数。'


def specification():
    return {'language':'zh','requirements':[{'id':'metric','kind':'metric','description':'Sum amount','quote':'金额总和'}],
            'operation':'sum','output_count':1,'fields':[{'name':'sum','type':'number','decimals':2,'quote':'两位小数'}],
            'separator':', ','ambiguities':[]}


def computation():
    return {'steps':[{'id':'data','op':'read','table':'t1','start_row':2,'end_row':4},
                     {'id':'total','op':'aggregate','input':'data','metric':'sum','expr':{'col':'B'}}],
            'answer_spec':{'answer_type':'number','answer_field':'金额','operation':'sum'},'output':{'step':'total'}}


def response(plan=None):
    return {'action':'plan','coverage':[{'requirement':'metric','table':'t1','range':'B1:B4','rationale':'Amount heading and rows'}],
            'plan':plan or computation()}


class FakeClient:
    prompt_cap=.01;completion_cap=.01
    def __init__(self,replies): self.replies=list(replies);self.purposes=[]
    def complete(self,messages,ledger,purpose):
        self.purposes.append(purpose)
        call=ledger.reserve(messages,100)
        call.update(status='completed',prompt_tokens=10,completion_tokens=10,total_tokens=20,
                    actual_cost_usd=.000001,estimated_cost_usd=.000001,budget_charge_usd=.000001)
        if not self.replies: raise AssertionError('Unexpected additional generation')
        return json.dumps(self.replies.pop(0),ensure_ascii=False)


class V3Test(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.path=Path(self.temp.name)/'arbitrary.xlsx'
        book=openpyxl.Workbook();sheet=book.active;sheet.title='金额'
        sheet.append(['名称','金额','权重']);sheet.append(['甲',10,1]);sheet.append(['乙',20,2]);sheet.append(['丙',30,3]);book.save(self.path);book.close()

    def test_grounded_multilingual_pipeline_has_eight_phases(self):
        client=FakeClient([specification(),response()]);record={}
        value=answer(self.path,QUESTION,client,record,{})
        self.assertEqual(value,'60.00');self.assertTrue(record['validated'])
        self.assertEqual({p['phase'] for p in record['phases']},set(range(1,9)))
        self.assertEqual(client.purposes,['v3_spec','v3_plan'])
        self.assertEqual(record['usage']['api_calls'],2)
        self.assertIn('金额!B2',record['evidence'])

    def test_missing_evidence_triggers_tool_then_plan(self):
        inspect={'action':'inspect','requests':[{'op':'read_cell','table':'t1','cell':'B2'}],'missing_requirements':['metric']}
        client=FakeClient([specification(),inspect,response()]);record={}
        self.assertEqual(answer(self.path,QUESTION,client,record,{}),'60.00')
        self.assertTrue(any(i['action']=='inspect' for i in record['interactions']))

    def test_unseen_coverage_is_not_claimed_complete(self):
        book=Workbook(self.path);book.remember('t1','B1')
        with self.assertRaisesRegex(PlanError,'Missing requirement evidence'):
            coverage_check(specification(),response()['coverage'],book)

    def test_fabricated_question_quote_rejected(self):
        spec=specification();spec['fields'][0]['quote']='three decimal places'
        with self.assertRaisesRegex(PlanError,'quote'): check_spec(spec,QUESTION)

    def test_inspection_preserves_blank_merged_and_exact_ranges(self):
        book=openpyxl.load_workbook(self.path);sheet=book.active;sheet.merge_cells('A6:C6');sheet['A6']='标题';book.save(self.path);book.close()
        workspace=Workbook(self.path);region=next(r for r in workspace.regions if r.whole_sheet)
        result=workspace.inspect([{'op':'read_range','table':region.id,'range':'A5:C6'}],QUESTION)[0]
        self.assertEqual(len(result['cells']),6);self.assertFalse(result['truncated'])
        merged=next(c for c in result['cells'] if c['cell'].endswith('B6'))
        self.assertEqual(merged['raw'],'');self.assertEqual(merged['value'],'标题');self.assertEqual(merged['merged_anchor'],'A6')
        self.assertEqual(result['cells'][0]['type'],'blank')

    def test_weighted_mean_and_median_preserve_sources(self):
        workspace=Workbook(self.path)
        for operation,expected in [('weighted_mean','23.33333333333333333333333333'),('median','20')]:
            plan=computation()
            plan['steps'][-1]=({'id':'total','op':'weighted_mean','input':'data','value':{'col':'B'},'weight':{'col':'C'}}
                if operation=='weighted_mean' else {'id':'total','op':'aggregate','input':'data','metric':'median','expr':{'col':'B'}})
            result,_=run(plan,workspace.regions);self.assertEqual(result['answer'],expected)
            self.assertIn('金额!B3',result['evidence'])

    def test_percent_change_zero_denominator_and_forward_ref(self):
        plan={'steps':[{'id':'old','op':'cell','table':'t1','row':2,'column':'B'},
                       {'id':'new','op':'cell','table':'t1','row':3,'column':'B'},
                       {'id':'change','op':'percent_change','old':{'ref':'old'},'new':{'ref':'new'}}],
              'answer_spec':{'answer_type':'number','answer_field':'amount','operation':'ratio'},'output':{'step':'change'}}
        regions=Workbook(self.path).regions
        self.assertEqual(run(plan,regions)[0]['answer'],'100')
        plan['steps'][-1]['old']={'value':0}
        with self.assertRaises(PlanError): run(plan,regions)
        plan['steps'][0]={'id':'old','op':'calculate','expr':{'ref':'new'}}
        with self.assertRaisesRegex(PlanError,'preceding scalar'): run(plan,regions)

    def test_mixed_fields_format_numbers_without_changing_labels(self):
        spec=specification();spec['output_count']=2
        spec['fields']=[{'name':'code','type':'label','decimals':None,'quote':'金额'},spec['fields'][0]]
        self.assertEqual(format_answer({'values':['001','14.201118428968']},spec),'001, 14.20')

    def test_reviewer_rejection_never_becomes_submission(self):
        plan=computation();plan['steps']=[{'id':'cell','op':'cell','table':'t1','row':2,'column':'A'}];plan['output']={'step':'cell'}
        spec=specification();spec['fields'][0].update(type='label',decimals=None)
        client=FakeClient([spec,response(plan),{'verdict':'revise','issues':['Wrong scope; inspect more rows']}]);record={}
        self.assertEqual(answer(self.path,QUESTION,client,record,{'v3_max_repairs':0}),'')
        self.assertFalse(record['validated']);self.assertNotIn('answer',record)
        self.assertEqual(record['status'],'unanswered')

    def test_risky_review_accepts_once(self):
        spec=specification();spec['fields'][0].update(type='label',decimals=None)
        plan=computation();plan['steps']=[{'id':'cell','op':'cell','table':'t1','row':2,'column':'A'}];plan['output']={'step':'cell'}
        client=FakeClient([spec,response(plan),{'verdict':'accept','issues':[]}]);record={}
        self.assertEqual(answer(self.path,QUESTION,client,record,{}),'甲')
        self.assertEqual(client.purposes.count('v3_review'),1)

    def test_isolation_and_all_provider_schemas(self):
        from src.planner import tool_schema
        legacy=tool_schema();self.assertNotIn('v3_spec',json.dumps(legacy))
        validate(specification(),tool_schema('v3_spec'));validate(response(),tool_schema('v3_plan'))
        self.assertTrue(tokens('原住民就业率') & tokens('就业率'))

    def test_shared_budget_includes_spec_inspection_and_review(self):
        client=FakeClient([specification()]);record={}
        self.assertEqual(answer(self.path,QUESTION,client,record,{'v3_max_api_calls':1}),'')
        self.assertEqual(record['usage']['api_calls'],1);self.assertFalse(record['validated'])

    def test_explicit_locale_and_date_parsing_do_not_guess(self):
        from src.tools.executor import expression, Result
        from decimal import Decimal
        from datetime import date
        self.assertEqual(expression({'parse_number':[{'value':'1.234,50'},{'value':'decimal_comma'}]}, {}, {}).value,Decimal('1234.50'))
        with self.assertRaises(PlanError): expression({'parse_number':[{'value':'1.23,50'},{'value':'decimal_comma'}]}, {}, {})
        self.assertEqual(expression({'parse_date':[{'value':'03/04/2020'},{'value':'%m/%d/%Y'}]}, {}, {}).value,date(2020,3,4))
        self.assertEqual(expression({'parse_date':[{'value':'03/04/2020'},{'value':'%d/%m/%Y'}]}, {}, {}).value,date(2020,4,3))

    def test_usage_keeps_reasoning_separate_from_completion(self):
        from src.usage import UsageLedger
        ledger=UsageLedger();call=ledger.reserve([],1)
        call.update(prompt_tokens=10,completion_tokens=8,total_tokens=18,reasoning_tokens=3)
        self.assertEqual(ledger.summary()['input_plus_reasoning_tokens'],13)
        self.assertEqual(ledger.summary()['total_tokens'],18)
        call['reasoning_tokens']=None
        self.assertIsNone(ledger.summary()['input_plus_reasoning_tokens'])

    def test_cli_routes_v3_and_writes_numeric_ids_in_order(self):
        import main
        queries=Path(self.temp.name)/'queries.json';output=Path(self.temp.name)/'submission.csv';logs=Path(self.temp.name)/'logs'
        queries.write_text(json.dumps([{'id':i,'file_name':self.path.name,'query':QUESTION} for i in [10,2,1]]),encoding='utf-8')
        cli_plan=computation();cli_plan['output']['decimals']=2
        client=FakeClient([cli_plan,{'verdict':'accept','issues':[]}]*3)
        args=['main.py','--mode','llm-v3','--xlsx_dir',self.temp.name,'--queries',str(queries),
              '--output',str(output),'--log-dir',str(logs),'--config',str(Path(self.temp.name)/'unused.yaml')]
        with patch('sys.argv',args),patch.object(main,'get_client',return_value=object()), \
             patch.object(main,'PlannerClient',return_value=client),patch('sys.stdout',new=io.StringIO()):
            main.main()
        with output.open(encoding='utf-8',newline='') as stream: rows=list(csv.DictReader(stream))
        self.assertEqual([r['id'] for r in rows],['1','2','10'])
        self.assertEqual([r['answer'] for r in rows],['60.00']*3)
        self.assertEqual(json.loads((logs/'summary.json').read_text())['pipeline_version'],'lmmv3-v2-dedup-v1')

    def test_dashboard_accepts_v3_without_launching_a_worker(self):
        import dashboard
        queries=[{'id':1,'file_name':self.path.name,'source_file':self.path.name,'query':QUESTION,'dataset':'synthetic'}]
        with patch.object(dashboard,'OUTPUTS',Path(self.temp.name)/'outputs'), \
             patch.object(dashboard,'datasets_for',return_value=[{'id':'synthetic','xlsx_dir':str(self.path.parent)}]), \
             patch.object(dashboard,'queries',return_value=queries),patch.object(dashboard,'reference',return_value={}), \
             patch.object(dashboard.threading,'Thread'),patch.dict(dashboard.JOB,running=False):
            run_id=dashboard.start_run('llm-v3',None,'synthetic')
            metadata=json.loads((dashboard.OUTPUTS/'runs'/run_id/'metadata.json').read_text(encoding='utf-8'))
            self.assertEqual(metadata['mode'],'llm-v3')

    def test_contextual_ambiguity_can_be_resolved_before_planning(self):
        ambiguous=specification();ambiguous['ambiguities']=['Need to inspect amount heading']
        revise={'action':'revise_spec','reason':'Heading confirms requested metric','spec':specification()}
        client=FakeClient([ambiguous,revise,response()]);record={}
        self.assertEqual(answer(self.path,QUESTION,client,record,{}),'60.00')
        self.assertEqual(len(record['specification_revisions']),1)

    def test_provider_numeric_spellings_are_normalized_without_guessing(self):
        for count in (1.0,'1','1.0'):
            with self.subTest(count=count):
                spec=specification();spec['output_count']=count;spec['fields'][0]['decimals']='2'
                client=FakeClient([spec,response()]);record={}
                self.assertEqual(answer(self.path,QUESTION,client,record,{}),'60.00')
                self.assertTrue(record['specification_normalizations'])
                self.assertEqual(client.purposes,['v3_spec','v3_plan'])
        from src.lmmv3.contracts import normalize_spec
        for value in (True,1.5,'one','1,0','',float('inf')):
            spec=specification();spec['output_count']=value
            with self.assertRaises(PlanError): check_spec(normalize_spec(spec)[0],QUESTION)
        spec=specification();spec['output_count']='null'
        self.assertIsNone(check_spec(normalize_spec(spec)[0],QUESTION)['output_count'])

    def test_spec_failure_logs_raw_type_and_retry_has_previous_response(self):
        spec=specification();spec['output_count']='several'
        class Client(FakeClient):
            def complete(self,messages,ledger,purpose):
                if self.purposes and purpose=='v3_spec':
                    feedback=json.loads(messages[1]['content'])['feedback']
                    assert 'previous_response' in feedback
                    assert "received str" in feedback['error']
                return super().complete(messages,ledger,purpose)
        record={}
        self.assertEqual(answer(self.path,QUESTION,Client([spec,spec,response(),{'verdict':'accept','issues':[]}]),record,{}),'60')
        self.assertTrue(record['specification_fallback'])
        self.assertEqual(len(record['model_responses']),4)
        self.assertEqual(json.loads(record['model_responses'][0]['raw'])['output_count'],'several')

    def test_dynamic_schema_rejects_unknown_ids_and_bounds_feedback_is_precise(self):
        active=schema('v3_plan',specification(),[{'table_id':'t1'}])
        bad=response();bad['coverage'][0]['requirement']='fabricated'
        validate(bad,active)
        book=Workbook(self.path)
        with self.assertRaisesRegex(PlanError,'unknown requirement'): coverage_check(specification(),bad['coverage'],book)
        book=Workbook(self.path);book.preview(book.index['t1'],QUESTION,45)
        bad=response();bad['coverage'][0]['range']='B1:B100'
        with self.assertRaisesRegex(PlanError,'bounds are A1:C4'): coverage_check(specification(),bad['coverage'],book)

    def test_context_ambiguity_resolves_with_grounded_evidence_then_review(self):
        spec=specification();spec['ambiguities']=['Which column holds the amount?']
        ready=response();ready['ambiguity_resolutions']=[{'index':0,'requirement':'metric','explanation':'B1 is the amount heading; B2:B4 contain the requested data.'}]
        class Client(FakeClient):
            def complete(self,messages,ledger,purpose):
                if purpose=='v3_plan':
                    assert 'coverage' in ledger.response_schema['anyOf'][0]['properties']
                if purpose=='v3_review': assert json.loads(messages[1]['content'])['ambiguity_resolutions']
                return super().complete(messages,ledger,purpose)
        record={};client=Client([spec,ready,{'verdict':'accept','issues':[]}])
        self.assertEqual(answer(self.path,QUESTION,client,record,{}),'60.00')
        self.assertTrue(record['readiness']['ambiguity_resolutions'])

    def test_unresolved_ambiguity_requires_reviewer_clearance(self):
        spec=specification();spec['ambiguities']=['Which period?']
        ready=response();ready['ambiguity_resolutions']=[];record={}
        self.assertEqual(answer(self.path,QUESTION,FakeClient([spec,ready,{'verdict':'revise','issues':['Read the missing period header before accepting']}]),record,{'v3_max_repairs':0}),'')
        self.assertFalse(record['validated'])
        self.assertEqual(record['readiness']['pending_review_ambiguities'],['Which period?'])

    def test_descriptive_kind_and_paraphrase_do_not_deadlock_before_execution(self):
        spec=specification();spec['requirements'][0]['kind']='group'
        spec['requirements'][0]['quote']='a paraphrased statement, not an exact quote'
        spec['ambiguities']=['Need to identify the amount column from workbook']
        client=FakeClient([spec,response(),{'verdict':'accept','issues':[]}]);record={}
        self.assertEqual(answer(self.path,QUESTION,client,record,{}),'60.00')
        self.assertTrue(record['specification_warnings'])
        self.assertEqual(client.purposes.count('v3_review'),1)

    def test_nested_metadata_and_unary_container_are_syntax_only(self):
        from src.lmmv3.contracts import normalize_response
        ready=response();ready['plan']['ambiguity_resolutions']=[]
        ready['plan']['steps'][-1]={'id':'total','op':'calculate','expr':{'abs':{'value':-7}}}
        normalized,changes=normalize_response(ready)
        self.assertNotIn('ambiguity_resolutions',normalized['plan'])
        self.assertEqual(normalized['plan']['steps'][-1]['expr'],{'abs':[{'value':-7}]})
        self.assertIn('hoist_ambiguity_resolutions',changes)
        self.assertNotIn('ambiguity_resolutions',ready)

    def test_output_format_is_grounded_in_question_not_fabricated_cells(self):
        spec=specification();spec['requirements'].append({'id':'format','kind':'output','description':'Two decimals','quote':'两位小数'})
        book=Workbook(self.path);book.preview(book.index['t1'],QUESTION,45)
        result=coverage_check(spec,response()['coverage'],book,QUESTION)
        self.assertEqual(result['question_grounded'],['format'])
        with self.assertRaisesRegex(PlanError,'Missing requirement evidence'):
            coverage_check(spec,response()['coverage'],book)

    def test_disjoint_evidence_and_auto_read_without_extra_generation(self):
        for notation in ['A1:B1;A3:B3',['$A$1:$B$1',"'金额'!A3:B3"]]:
            with self.subTest(notation=notation):
                ready=response();ready['coverage'][0]['range']=notation
                record={};client=FakeClient([specification(),ready])
                self.assertEqual(answer(self.path,QUESTION,client,record,{'preview_rows':1}),'60.00')
                self.assertEqual(client.purposes,['v3_spec','v3_plan'])
                self.assertTrue(record['readiness']['complete'])
                self.assertTrue(any(i['action']=='auto_read_evidence' for i in record['interactions']))

    def test_missing_advisory_envelope_does_not_block_valid_computation(self):
        plan=computation();plan.pop('answer_spec')
        client=FakeClient([specification(),plan,{'verdict':'accept','issues':[]}]);record={}
        self.assertEqual(answer(self.path,QUESTION,client,record,{}),'60.00')
        self.assertTrue(record['readiness']['warnings'])
        self.assertTrue(record['validated'])

    def test_bad_coverage_never_clamps_coordinates_but_review_can_clear_real_plan(self):
        ready=response();ready['coverage'][0]['range']='B900:B910'
        client=FakeClient([specification(),ready,{'verdict':'accept','issues':[]}]);record={}
        self.assertEqual(answer(self.path,QUESTION,client,record,{}),'60.00')
        self.assertTrue(record['readiness']['warnings'])
        self.assertFalse(record['readiness']['complete'])
        self.assertNotIn('金额!B900',record['evidence'])

    def test_repeated_invalid_computation_stops_without_retry_loop_or_fake_answer(self):
        bad=response();bad['plan']['steps'][1]['expr']={'col':'MISSING'}
        client=FakeClient([specification(),bad,bad]);record={}
        self.assertEqual(answer(self.path,QUESTION,client,record,{}),'')
        self.assertEqual(client.purposes,['v3_spec','v3_plan','v3_plan'])
        self.assertEqual(record['status'],'unanswered');self.assertFalse(record['validated'])
        self.assertNotIn('failure_event',record)

    def test_failed_reviewer_metadata_keeps_computed_candidate_with_explicit_warning(self):
        spec=specification();spec['ambiguities']=['Check source column']
        client=FakeClient([spec,response(),{'unexpected':'metadata'}, {'unexpected':'metadata'}]);record={}
        self.assertEqual(answer(self.path,QUESTION,client,record,{}),'60.00')
        self.assertFalse(record['validated']);self.assertTrue(record['finalized_with_warning'])
        self.assertEqual(record['verification_level'],'computed_review_unavailable')

    def test_no_review_budget_keeps_computation_but_never_promotes_rejected_candidate(self):
        spec=specification();spec['ambiguities']=['Check units']
        record={};client=FakeClient([spec,response()])
        self.assertEqual(answer(self.path,QUESTION,client,record,{'v3_max_api_calls':2}),'60.00')
        self.assertTrue(record['finalized_with_warning'])
        record={};client=FakeClient([spec,response(),{'verdict':'revise','issues':['Wrong metric']}])
        self.assertEqual(answer(self.path,QUESTION,client,record,{'v3_max_api_calls':3}),'')
        self.assertTrue(record['candidate_rejected']);self.assertFalse(record.get('finalized_with_warning',False))


if __name__=='__main__': unittest.main()
