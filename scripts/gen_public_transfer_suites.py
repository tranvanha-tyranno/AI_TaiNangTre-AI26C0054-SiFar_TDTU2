"""Build 10 fresh, seeded transfer benchmarks inspired by BOTH public datasets.

python scripts/gen_public_transfer_suites.py --seed 261016
Then export the spec using scripts/author_public_transfer_workbooks.mjs.
Reference expressions are for evaluation only and never consumed by src/.
"""
from __future__ import annotations

import argparse
import csv
from datetime import date, timedelta
import hashlib
import json
from pathlib import Path
import random
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.gen_extended_test_suites import Suite, cell, table, filt, agg, ranked, op


def field(name): return {"field": name}
def derive(data, **fields): return {"derive": {"data": data, "fields": fields}}
def absolute(expr): return {"abs": expr}
def items(*expressions): return {"list": list(expressions)}
def formatted(expr, decimals): return {"format": {"value": expr, "decimals": decimals}}
def pct(a, b): return op("mul", op("div", a, b), 100)
def parsed(data, *columns): return derive(data, **{c: {"parse": field(c)} for c in columns})
INT = " Answer only the integer, without thousands separators or units."
ONE = " Answer only the number with exactly 1 decimal place, no units or percent sign."
TWO = " Answer only the number with exactly 2 decimal places, no units or percent sign."
ORDER = " Break ties by the label in ascending Unicode order."


class TransferSuite(Suite):
    def __init__(self, name, description, sources, rng, seed):
        super().__init__(name, description)
        self.sources, self.rng, self.seed = sources, rng, seed

    def sheet(self, name, rows, start="A1", merges=(), formats=None, headers=None, widths=None):
        super().sheet(name, rows, start, merges, formats)
        self.sheets[-1].update(headers=headers or [f"A1:{colname(max(map(len, rows)))}1"],
                               widths=widths or {}, previews=[f"A1:{colname(max(map(len, rows)))}{min(len(rows),12)}"])

    def save(self, target):
        assert len(self.queries) == 30, (self.name, len(self.queries))
        assert len({q["query"] for q in self.queries}) == 30
        spec = super().save(target)
        folder = target / self.name
        with (folder / "sample_submission.csv").open("w", encoding="utf-8", newline="") as stream:
            w = csv.writer(stream, lineterminator="\n")
            w.writerow(["id", "answer"])
            w.writerows((q["id"], "") for q in self.queries)
        provenance = {"seed": self.seed, "kind": "new synthetic data; structural transfer, not copied records",
                      "source_workbooks": self.sources, "description": self.description,
                      "questions": 30, "answer_rounding": "Decimal ROUND_HALF_UP; explicit precision in query",
                      "generator": "scripts/gen_public_transfer_suites.py",
                      "categories": sorted({c["category"] for c in self.checks})}
        (folder / "provenance.json").write_text(json.dumps(provenance, ensure_ascii=False, indent=2), encoding="utf-8")
        (folder / "README.md").write_text(
            f"# {self.name}\n\n{self.description}\n\n"
            f"Đúng 30 câu; seed `{self.seed}`. Toàn bộ bản ghi, kỳ báo cáo và số liệu được sinh mới. "
            "Dữ liệu chỉ là tình huống tổng hợp để kiểm thử, không phải thống kê thực tế.\n\n"
            "Cấu trúc tham khảo: " + ", ".join(f"`{p}`" for p in self.sources) + ".\n\n"
            "Nhóm câu hỏi: " + ", ".join(provenance["categories"]) + ".\n\n"
            "Đầu vào mô hình chỉ gồm `xlsx_files/` và `queries.json`. `expected_submission.csv`, "
            "`answer_checks.json` và `provenance.json` dành riêng cho đánh giá; không đưa chúng vào prompt/retrieval. "
            "`sample_submission.csv` có đáp án rỗng để làm mẫu I/O. Đáp án được tính từ địa chỉ ô, "
            "sau đó kiểm lại trên chính XLSX đã export, độc lập với bộ giải.\n\n"
            "```powershell\n"
            f"python main.py --mode llm --xlsx_dir Test/{self.name}/xlsx_files --queries Test/{self.name}/queries.json "
            f"--output outputs/{self.name}.csv --log-dir outputs/{self.name}-logs\n"
            f"python scripts/evaluate.py --prediction outputs/{self.name}.csv --reference Test/{self.name}/expected_submission.csv "
            f"--report outputs/{self.name}-evaluation.json\n"
            "```\n\nChưa đo độ chính xác LLM trên bộ này. "
            "Dashboard tự phát hiện khi tải lại danh sách dataset.\n", encoding="utf-8")
        return spec


def colname(n):
    out = ""
    while n:
        n, r = divmod(n-1, 26)
        out = chr(65+r) + out
    return out


def admission(s):
    rows = [["No", "Name", None, None, "Date of", None, "Score", "Campus"],
            [None, "First", "Middle", "Last", "Birth", "Admission", None, None]]
    for i in range(1, 65):
        birth = date(1989, 1, 1) + timedelta(days=s.rng.randrange(6200))
        admission_date = date(2026, 1, 1) + timedelta(days=s.rng.randrange(700))
        rows.append([i, ["Bao", "Quynh", "Hieu", "Thao"][i%4], f"Minh{i:02d}",
                     ["Truong", "Ho", "Dinh", "Luong", "Le"][i%5], birth, admission_date,
                     s.rng.randint(40, 100), ["Hue", "Can Tho", "Da Lat"][i%3]])
    # Birthday boundary and leap-day examples are intentional, not random accidents.
    rows[2][4:6] = [date(2004, 2, 29), date(2026, 2, 28)]
    rows[3][4:6] = [date(2004, 2, 29), date(2026, 3, 1)]
    for idx, score in [(4,100),(5,100),(6,75),(7,60),(8,90)]: rows[idx][6]=score
    s.sheet("People", rows, merges=["A1:A2", "B1:D1", "E1:F1", "G1:G2", "H1:H2"],
            formats={"E3:F66": "yyyy-mm-dd"}, headers=["A1:H2"])
    t = table("People", dict(id="A", first="B", middle="C", last="D", birth="E", admission="F", score="G", campus="H"), 3, 66)
    t = derive(t, age={"age": [field("birth"), field("admission")]}, days={"days": [field("birth"), field("admission")]})
    for k in [2, 3, 5, 9]:
        s.add("age_rank", f"Rank people by exact age in days on their own Admission date, youngest first, ties by No ascending. What is the last name at position {k} (1-based)? Answer only the last name.", ranked(t, [["days", "asc"], ["id", "asc"]], ["last"], k-1))
    for i in [1, 2, 19, 64]:
        s.add("completed_age", f"For No {i}, what is the age in completed years on Admission? A Feb 29 birthday is reached on March 1 in a non-leap year."+INT, {"age": [cell("People", "E", i+2), cell("People", "F", i+2)]})
    for n in [5, 10, 15]:
        youngest = filt(t, id={"gt": 0})
        # The rank output is a list of scores, so max is computed on exactly n selected records.
        s.add("top_k_age_score", f"Among the {n} youngest people by exact age in days at Admission, what is the highest Score? Ties in age use No ascending."+INT, {"max": ranked(youngest, [["days", "asc"], ["id", "asc"]], ["score"], limit=n)})
    for campus in ["Hue", "Can Tho", "Da Lat"]:
        sub = filt(t, campus=campus)
        s.add("conditional_average", f"What is the mean Score for Campus {campus}?"+TWO, agg("mean", sub, "score"), 2)
        s.add("multi_condition", f"How many people in Campus {campus} have Score >= 75 and completed age at Admission < 30?"+INT, agg("count", filt(sub, score={"ge": 75}, age={"lt": 30})))
        s.add("conditional_rank", f"Which No in Campus {campus} has the highest Score, ties by No ascending?"+INT, ranked(sub, [["score", "desc"], ["id", "asc"]], ["id"]))
    for threshold in [60, 75, 90]:
        s.add("strict_boundary", f"How many people have Score strictly greater than {threshold}?"+INT, agg("count", filt(t, score={"gt": threshold})))
    s.add("date_filter", "How many admissions fall from 2026-01-01 through 2026-12-31 inclusive?"+INT, agg("count", filt(t, admission={"ge": "2026-01-01", "le": "2026-12-31"})))
    s.add("rank_score", "What Score did the second-oldest person receive, using exact age in days at Admission and No ascending to break ties?"+INT, ranked(t, [["days", "desc"], ["id", "asc"]], ["score"], 1))
    s.add("top_three", "List No of the 3 highest-scoring people, Score descending then No ascending. Separate with comma and one space.", ranked(t, [["score", "desc"], ["id", "asc"]], ["id"], limit=3))
    s.add("empty_result", "How many people have Score above 100?"+INT, agg("count", filt(t, score={"gt": 100})))
    s.add("exact_age_days", "How many days old was No 64 on their Admission date?"+INT, {"days": [cell("People", "E", 66), cell("People", "F", 66)]})
    s.add("distinct", "How many distinct last names appear in People?"+INT, agg("distinct_count", t, "last"))
    s.add("birthday_boundary", "What is No 2 completed age at Admission minus No 1 completed age at Admission? Use March 1 for Feb 29 birthdays in non-leap years."+INT, op("sub", {"age":[cell("People","E",4),cell("People","F",4)]}, {"age":[cell("People","E",3),cell("People","F",3)]}))


def trade(s):
    rows = [["Commodity", "Fiscal years", None, "Fiscal year-to-date", None, "September"],
            [None, 2025, 2026, 2025, 2026, 2026], [None, "Billion dollars", None, None, None, None]]
    labels = ["Cocoa", "Rice", "Coffee", "Fruit", "Dairy", "Seafood"]
    for label in labels:
        exports = [s.rng.randint(40, 250) for _ in range(5)]
        imports = [s.rng.randint(70, 230) for _ in range(5)]
        rows += [[label+" exports", *exports], [label+" imports", *imports],
                 [label+" trade balance", *[a-b for a, b in zip(exports, imports)]]]
    rows += [[None]*6, ["Commodity", "Calendar years", None, "Calendar year-to-date", None, "September"],
             [None, 2025, 2026, 2025, 2026, 2026], [None, "Billion dollars", None, None, None, None]]
    calendar_first = len(rows)+1
    for label in labels:
        rows.append([label+" exports", *[s.rng.randint(30, 220) for _ in range(5)]])
    s.sheet("Trade", rows, merges=["A1:A3", "B1:C1", "D1:E1", "B3:F3", "A23:A25", "B23:C23", "D23:E23", "B25:F25"], headers=["A1:F3", "A23:F25"], widths={"A": 30})
    for j, label in enumerate(labels):
        r = 4+j*3
        s.add("fiscal_vs_ytd", f"For {label} exports, is fiscal year-to-date 2026 greater than fiscal year-to-date 2025? Answer only Yes or No.", op("gt", cell("Trade", "E", r), cell("Trade", "D", r)))
        s.add("trade_balance", f"What is fiscal year-to-date 2026 exports minus imports for {label}, in billion dollars?"+INT, op("sub", cell("Trade", "E", r), cell("Trade", "E", r+1)))
        s.add("period_disambiguation", f"For {label} exports, what is calendar year-to-date 2026 minus fiscal year-to-date 2026?"+INT, op("sub", cell("Trade", "E", calendar_first+j), cell("Trade", "E", r)))
        s.add("growth", f"What is the percentage change of {label} exports from fiscal year 2025 to fiscal year 2026? Use (new-old)/old*100."+TWO, pct(op("sub", cell("Trade", "C", r), cell("Trade", "B", r)), cell("Trade", "B", r)), 2)
    for c, period in [("B", "fiscal year 2025"), ("C", "fiscal year 2026"), ("D", "fiscal year-to-date 2025"), ("E", "fiscal year-to-date 2026")]:
        s.add("noncontiguous_sum", f"Sum exports for all six commodities for {period}. Exclude imports and balances."+INT, op("add", op("add", op("add", cell("Trade", c, 4), cell("Trade", c, 7)), op("add", cell("Trade", c, 10), cell("Trade", c, 13))), op("add", cell("Trade", c, 16), cell("Trade", c, 19))))
    s.add("signed_balance", "How many of the six commodities have a negative fiscal year-to-date 2026 trade balance?"+INT, {"count": filt(table("Trade", {"label":"A", "balance":"E"},4,21), label={"not_blank": True}, balance={"lt":0})})
    # Count only the explicit balance rows, since imports/exports are always positive.
    s.add("absolute_difference", "What is the absolute difference between Coffee and Seafood exports in September 2026?"+INT, absolute(op("sub", cell("Trade", "F", 10), cell("Trade", "F", 19))))


def prices(s):
    rows = [["Indent Level", "Category", "Relative importance", "Unadjusted indexes", None, "Unadjusted percent change", None, "Seasonally adjusted change"],
            [None, None, "2026", "Sep 2026", "Oct 2026", "Oct 2025 to Oct 2026", "Sep 2026 to Oct 2026", "Sep 2026 to Oct 2026"]]
    definitions = [(0,"All items"),(1,"Food and beverages"),(2,"Food at home"),(2,"Dining out"),(1,"Housing"),
                   (2,"Shelter"),(3,"Rent"),(3,"Owner equivalent rent"),(2,"Utilities"),(3,"Electricity"),(3,"Gas"),
                   (2,"Household furnishings"),(1,"Transport"),(2,"Vehicles"),(2,"Transport services"),
                   (1,"Commodities"),(1,"Services"),(1,"Health"),(2,"Medicine"),(2,"Hospital services")]
    for level, label in definitions:
        sep=s.rng.randint(180000,340000)/1000
        monthly=s.rng.randint(-24,40)/10
        rows.append([level,label,100 if not level else s.rng.randint(100,42000)/1000,sep,round(sep*(1+monthly/100),3),s.rng.randint(-30,85)/10,monthly,s.rng.randint(-20,30)/10])
    s.sheet("Prices", rows, merges=["D1:E1","F1:G1"],headers=["A1:H2"],formats={"C3:E22":"0.000","F3:H22":"0.0"},widths={"B":30})
    t=table("Prices",dict(level="A",label="B",weight="C",sep="D",oct="E",annual="F",monthly="G",adjusted="H"),3,22)
    t=derive(t, magnitude=absolute(field("monthly")),gap=absolute(op("sub",field("monthly"),field("adjusted"))))
    for label in ["Housing","Food and beverages","Services","Commodities","Transport"]:
        r=next(i+3 for i,(_,x) in enumerate(definitions) if x==label)
        s.add("percent_points",f"For {label}, what is annual unadjusted change minus monthly unadjusted change, in percentage points?"+ONE,op("sub",cell("Prices","F",r),cell("Prices","G",r)),1)
        s.add("adjusted_disambiguation",f"For {label}, what is monthly unadjusted change minus seasonally adjusted change?"+ONE,op("sub",cell("Prices","G",r),cell("Prices","H",r)),1)
        s.add("index_difference",f"What is the October 2026 index minus September 2026 index for {label}? Answer exactly 3 decimal places.",op("sub",cell("Prices","E",r),cell("Prices","D",r)),3)
    for threshold in [0,1,2,3]:
        s.add("filter",f"Excluding All items, how many categories have monthly unadjusted change strictly greater than {threshold}?"+INT,agg("count",filt(t,level={"gt":0},monthly={"gt":threshold})))
    # Only indent level 2 descendants of Housing (rows 8, 11, 14), not grandchildren.
    housing=filt(table("Prices",dict(level="A",label="B",annual="F",monthly="G"),8,14),level=2)
    s.add("direct_children","Rank the categories directly under Housing by annual unadjusted change descending; exclude grandchildren."+ORDER+" Separate category names with comma and one space.",ranked(housing,[["annual","desc"],["label","asc"]],["label"],limit=3))
    s.add("hierarchy_average","What is mean annual unadjusted change of the three categories directly under Housing? Exclude grandchildren."+TWO,agg("mean",housing,"annual"),2)
    for metric,desc in [("magnitude","largest absolute monthly unadjusted change"),("gap","largest absolute gap between monthly unadjusted and seasonally adjusted change"),("weight","highest relative importance"),("annual","highest annual unadjusted change")]:
        s.add("ranking",f"Excluding All items, which category has the {desc}?"+ORDER+" Answer only the category name.",ranked(filt(t,level={"gt":0}),[[metric,"desc"],["label","asc"]],["label"]))
    for level in [1,2,3]:
        s.add("hierarchy_scope",f"How many categories have Indent Level exactly {level}?"+INT,agg("count",filt(t,level=level)))
    s.add("weight_gap","What is Housing relative importance minus Food and beverages relative importance? Answer exactly 3 decimal places.",op("sub",cell("Prices","C",7),cell("Prices","C",4)),3)
    s.add("comparison","Is the annual unadjusted change of Services greater than that of Commodities? Answer only Yes or No.",op("gt",cell("Prices","F",19),cell("Prices","F",18)))


def demographics(s):
    levels=["No diploma","High school","Some college, no degree","Associate degree","Bachelor degree","Advanced degree"]
    rows=[["Sex / employment status","2026",None,None,None,None,None], [None,*levels]]
    blocks={}
    for sex in ["Men","Women"]:
        rows.append([sex]); first=len(rows)+1
        population=[s.rng.randint(6000,40000) for _ in levels]
        force=[p*s.rng.randint(45,85)//100 for p in population]
        unemployed=[v*s.rng.randint(2,12)//100 for v in force]
        employed=[a-b for a,b in zip(force,unemployed)]
        rates=[round(100*u/f,1) for u,f in zip(unemployed,force)]
        for label,vals in [("Civilian noninstitutional population",population),("Civilian labor force",force),("Employed",employed),("Unemployed",unemployed),("Unemployment rate",rates)]: rows.append([label,*vals])
        blocks[sex]=first
        rows.append([])
    s.sheet("Demographics",rows,merges=["B1:G1"],headers=["A1:G2","A3:G3","A10:G10"],widths={"A":43,"D":28})
    for sex in ["Women","Men"]:
        r=blocks[sex]
        for j,level in enumerate(levels):
            col=colname(j+2)
            s.add("population_gap",f"For {sex} with {level}, what is population minus civilian labor force? All counts are in thousands; do not multiply by 1000."+INT,op("sub",cell("Demographics",col,r),cell("Demographics",col,r+1)))
        s.add("unweighted_average",f"What is the simple unweighted average Unemployment rate across the six education levels for {sex}?"+TWO,{"mean":items(*[cell("Demographics",colname(j+2),r+4) for j in range(6)])},2)
        s.add("weighted_rate",f"Using summed unemployed divided by summed civilian labor force across the six disjoint education groups for {sex}, what is the unemployment rate in percent?"+TWO,pct({"sum":items(*[cell("Demographics",colname(j+2),r+3) for j in range(6)])},{"sum":items(*[cell("Demographics",colname(j+2),r+1) for j in range(6)])}),2)
        for j in [0,2,5]:
            col=colname(j+2)
            s.add("participation_rate",f"What percentage of the population is in the civilian labor force for {sex} with {levels[j]}?"+TWO,pct(cell("Demographics",col,r+1),cell("Demographics",col,r)),2)
    for j,level in enumerate(levels):
        col=colname(j+2)
        s.add("repeated_label_scope",f"For {level}, what is men's employed count minus women's employed count, in thousands?"+INT,op("sub",cell("Demographics",col,blocks["Men"]+2),cell("Demographics",col,blocks["Women"]+2)))
    s.add("comparison","Is the unemployment rate for Women with Advanced degree higher than Men with Advanced degree? Answer only Yes or No.",op("gt",cell("Demographics","G",blocks["Women"]+4),cell("Demographics","G",blocks["Men"]+4)))
    s.add("nonoverlap_total","Sum Employed across all six education levels and both sexes, in thousands."+INT,{"sum":items(*[cell("Demographics",colname(j+2),blocks[sex]+2) for sex in blocks for j in range(6)])})


def industry(s):
    names=["Agriculture","Manufacturing","Construction","Financial activities","Transportation and utilities","Education","Hospitality","Information"]
    rows=[["Industry","Employed by age (thousands)",None,None,"At work by hours (thousands)",None,"Average hours"],
          [None,"16 to 19","20 to 24","25 and over","1 to 34 hours","35 hours or more","All at work"]]
    occupations=[["Industry","Protective service","Other service","Transportation and material moving","Production"]]
    for name in names:
        ages=[s.rng.randint(100,900),s.rng.randint(400,2000),s.rng.randint(1000,8000)]
        total=sum(ages); part=total*s.rng.randint(15,40)//100
        rows.append([name,*ages,part,total-part,s.rng.randint(320,450)/10])
        occupations.append([name,*[s.rng.randint(50,1500) for _ in range(4)]])
    rows.append(["TOTAL",*[sum(row[c] for row in rows[2:]) for c in range(1,6)],None])
    s.sheet("Employment",rows,merges=["A1:A2","B1:D1","E1:F1","G1:G2"],headers=["A1:G2"],widths={"A":40})
    s.sheet("Occupations",occupations,widths={"A":40,"D":35})
    t=table("Employment",dict(label="A",teen="B",young="C",adult="D",part="E",full="F",hours="G"),3,10)
    t=derive(t,total=op("add",op("add",field("teen"),field("young")),field("adult")),gap=op("sub",field("young"),field("teen")))
    for j,name in enumerate(names):
        r=j+3
        s.add("age_difference",f"In Employment, how many more employed people are 20 to 24 than 16 to 19 in {name}? Keep the table's thousands unit."+INT,op("sub",cell("Employment","C",r),cell("Employment","B",r)))
        s.add("hours_share",f"In {name}, what percentage of persons at work worked 1 to 34 hours? Use part-time count/(part-time+35 hours or more)*100."+TWO,pct(cell("Employment","E",r),op("add",cell("Employment","E",r),cell("Employment","F",r))),2)
    for j,name in enumerate(names[:4]):
        s.add("occupation_sum",f"In Occupations, what is Protective service plus Other service for {name}?"+INT,op("add",cell("Occupations","B",j+2),cell("Occupations","C",j+2)))
    for metric,label in [("teen","16 to 19 employed count"),("full","35 hours or more count"),("hours","Average hours"),("gap","20 to 24 minus 16 to 19 employed gap")]:
        s.add("exclude_total_rank",f"Excluding TOTAL, which industry has the highest {label}?"+ORDER+" Answer only the industry name.",ranked(t,[[metric,"desc"],["label","asc"]],["label"]))
    occ=table("Occupations",dict(label="A",transport="D"),2,4)
    s.add("sector_order","Rank Agriculture, Manufacturing and Construction by Transportation and material moving count descending."+ORDER+" Separate names with comma and one space.",ranked(occ,[["transport","desc"],["label","asc"]],["label"],limit=3))
    s.add("no_double_count","What is total 1 to 34 hours across the eight industries? Do not add TOTAL again."+INT,agg("sum",t,"part"))
    s.add("pair_output","Return TOTAL 1 to 34 hours and TOTAL 35 hours or more, in that order, as two integers separated by comma and one space.",items(cell("Employment","E",11),cell("Employment","F",11)))
    s.add("cross_sheet_ratio","For Manufacturing, what percentage of all employed people work in Production occupations? Use Production from Occupations divided by the three age counts in Employment."+TWO,pct(cell("Occupations","E",3),op("add",op("add",cell("Employment","B",4),cell("Employment","C",4)),cell("Employment","D",4))),2)
    s.add("two_industries","What is combined 20 to 24 employed count in Construction and Manufacturing?"+INT,op("add",cell("Employment","C",5),cell("Employment","C",4)))
    s.add("comparison","Does Financial activities have more employed persons aged 25 and over than Transportation and utilities? Answer only Yes or No.",op("gt",cell("Employment","D",6),cell("Employment","D",7)))


def time_use(s):
    rows=[["Activity","Average hours per day",None,None,None,"Percent doing activity",None,None,None],
          [None,"Married mothers",None,"Married fathers",None,"Married mothers",None,"Married fathers",None],
          [None,"Full time","Not employed","Full time","Not employed","Full time","Not employed","Full time","Not employed"]]
    groups=["Own household children under 6","Own household children 6-12","Own household children 13-17"]
    activities=["Housework","Childcare","Sleeping","Leisure","Travel"]
    refs={}
    for group in groups:
        rows.append([group]); rows.append(["Total, all activities",24,24,24,24,100,100,100,100])
        for activity in activities:
            refs[group,activity]=len(rows)+1
            rows.append([activity,*[s.rng.randint(700,900)/100 if activity=="Sleeping" else s.rng.randint(30,300)/100 for _ in range(4)],*[s.rng.randint(200,980)/10 for _ in range(4)]])
        rows.append([])
    s.sheet("Time use",rows,merges=["A1:A3","B1:E1","F1:I1","B2:C2","D2:E2","F2:G2","H2:I2"],headers=["A1:I3","A4:I4","A12:I12","A20:I20"],formats={"B4:E27":"0.00","F4:I27":"0.0"},widths={"A":45})
    for group in groups:
        for activity in activities[:4]:
            r=refs[group,activity]
            s.add("group_hours_gap",f"For {group}, what is not-employed fathers' average hours on {activity} minus full-time mothers' average hours? Use hours, not participation percent."+TWO,op("sub",cell("Time use","E",r),cell("Time use","B",r)),2)
        for activity in ["Housework","Childcare"]:
            r=refs[group,activity]
            s.add("percent_vs_hours",f"For {group}, what is not-employed mothers' participation percent minus full-time mothers' participation percent for {activity}, in percentage points?"+ONE,op("sub",cell("Time use","G",r),cell("Time use","F",r)),1)
        r=refs[group,"Housework"]
        a=cell("Time use","B",r); b=cell("Time use","E",r)
        # Winner is ranked from source cells as well; the text is never taken from a public answer.
        winner={"winner": {"left":a,"right":b,"left_label":"full-time employed mothers","right_label":"not-employed fathers"}}
        s.add("mixed_output",f"For {group}, who spends more average hours on Housework: full-time employed mothers or not-employed fathers? Answer `group, difference`, using one of those two group labels and exactly 2 decimals for the positive difference. If equal, choose full-time employed mothers.",items(winner,formatted(absolute(op("sub",a,b)),2)))
    for activity in activities:
        # A compact table expression preserves labels and their actual source values.
        mapped={"records": [{"label":g,"value":cell("Time use","B",refs[g,activity])} for g in groups]}
        s.add("cross_block_rank",f"Across the three child groups, which has the highest full-time mothers' average hours on {activity}?"+ORDER+" Answer the group name exactly.",ranked(mapped,[["value","desc"],["label","asc"]],["label"]))
    for group in groups:
        s.add("activity_sum",f"For {group}, sum full-time fathers' average hours on Housework and Childcare."+TWO,op("add",cell("Time use","D",refs[group,"Housework"]),cell("Time use","D",refs[group,"Childcare"])),2)
    s.add("unit_conversion","For Own household children under 6, how many minutes per day do full-time mothers spend on Travel? Multiply average hours by 60."+TWO,op("mul",cell("Time use","B",refs[groups[0],"Travel"]),60),2)


def safety(s):
    rows=[["Industry","NAICS code","Total recordable cases","Days away, restriction or transfer",None,None,"Other cases"],
          [None,None,None,"Reported combined total","Days away","Restriction or transfer",None]]
    for i in range(1,1101):
        code=str(400000+i) if i%5 else str(400+i%500)
        total=5.0 if i%41==0 else s.rng.randint(0,110)/10
        combined=s.rng.randint(0,60)/10
        # The reported combined rate can differ from rounded components.
        rows.append([f"Industry {i:04d}",code,total,combined,round(combined*.55,1),round(combined*.45+.1,1),s.rng.randint(0,40)/10])
    rows[551][2]=None; rows[1001][2]=".."
    s.sheet("Incidence",rows,merges=["A1:A2","B1:B2","C1:C2","D1:F1","G1:G2"],headers=["A1:G2"],formats={"C3:G1102":"0.0"},widths={"A":28,"D":28})
    s.sheets[-1]["previews"] += ["A1095:G1102"]
    t=parsed(table("Incidence",dict(label="A",code="B",total="C",combined="D",away="E",restriction="F",other="G"),3,1102),"total")
    t=derive(t,digits={"length":field("code")})
    for threshold in [0,2.5,5,7.5,10]:
        s.add("six_digit_scope",f"How many industries with a 6-digit NAICS code have Total recordable cases strictly above {threshold}? Ignore missing data."+INT,agg("count",filt(t,digits=6,total={"gt":threshold})))
    for i in [1,499,550,999,1099,1100]:
        s.add("reported_not_components",f"For Industry {i:04d}, what is the REPORTED combined total for days away, restriction or transfer? Do not add the two rounded components."+TWO,cell("Incidence","D",i+2),2)
    for low,high in [(1,300),(301,700),(701,1100)]:
        sub=filt(t,label={"ge":f"Industry {low:04d}","le":f"Industry {high:04d}"},total={"not_blank":True})
        s.add("full_scan_rank",f"Among Industry {low:04d} through Industry {high:04d} inclusive, which has the highest available Total recordable cases rate?"+ORDER+" Answer only the industry name.",ranked(sub,[["total","desc"],["label","asc"]],["label"]))
    quartiles=[["Industry / size","Mean","Q1","Median","Q3"]]
    industries=["Crop production","Machinery manufacturing","Food manufacturing","Warehousing"]
    blocks={}
    for name in industries:
        quartiles.append([name]); first=len(quartiles)+1
        for j,size in enumerate(["Total all sizes","1 - 10","11 - 49","50 - 249","250 - 999"]):
            q1=s.rng.randint(0,15)/10; median=q1+s.rng.randint(3,20)/10; q3=median+s.rng.randint(4,35)/10
            quartiles.append([size,"(-8-)" if name=="Food manufacturing" and j==4 else round(median+.3,1),round(q1,1),round(median,1),round(q3,1)])
        blocks[name]=first
    s.sheet("Quartiles",quartiles,widths={"A":32})
    for name in industries:
        first=blocks[name]
        sub=derive(table("Quartiles",dict(label="A",median="D",q3="E"),first+1,first+4),gap=op("sub",field("q3"),field("median")))
        s.add("quartile_rank",f"In {name}, excluding Total all sizes, which employment size has the largest Q3-minus-Median gap?"+ORDER+" Answer `size class, gap`, with 1 decimal for gap.",items(ranked(sub,[["gap","desc"],["label","asc"]],["label"]),formatted({"max":ranked(sub,[["gap","desc"],["label","asc"]],["gap"])},1)))
        s.add("quartile_threshold",f"In {name}, is Q3 for size 11 - 49 strictly greater than 5.0? Answer only Yes or No.",op("gt",cell("Quartiles","E",first+2),5))
        s.add("quartile_lookup",f"In {name}, what is the Median for size 250 - 999?"+ONE,cell("Quartiles","D",first+4),1)
    s.add("missing_marker","How many industries in Quartiles have a missing Mean for size 250 - 999? Treat (-8-) as missing."+INT,agg("count",filt(parsed(table("Quartiles",dict(label="A",mean="B"),2,25),"mean"),label="250 - 999",mean={"not_blank":False})))
    means={"records":[{"label":name,"value":{"parse":cell("Quartiles","B",blocks[name]+4)}} for name in industries]}
    s.add("ignore_suppressed","Which industry has the highest available Mean for size 250 - 999? Ignore (-8-) and break ties by industry name ascending. Answer only the industry name.",ranked(filt(means,value={"not_blank":True}),[["value","desc"],["label","asc"]],["label"]))
    s.add("exact_boundary","How many 6-digit NAICS industries have Total recordable cases exactly 5.0?"+INT,agg("count",filt(t,digits=6,total=5)))
    s.add("missing_count","How many rows in Incidence have missing Total recordable cases? Treat blanks and .. as missing."+INT,agg("count",filt(t,total={"not_blank":False})))


def productivity(s):
    rows=[["Period","Row type","Output per job index (2026=100)",None,None,None],
          [None,None,"Rubber, plastics & non-metallic minerals","Chemicals, Pharmaceuticals","Machinery & equipment","Food, beverages & tobacco"]]
    for year in range(2008,2028):
        rows.append([year,"Annual",*[100 if year==2026 else s.rng.randint(650,1350)/10 for _ in range(4)]])
        for q in range(1,5): rows.append([f"{year} Q{q}","Quarter",*[s.rng.randint(500,1700)/10 for _ in range(4)]])
    rows += [[None]*6,["Period","Row type","Level (GBP thousand)",None,None,None],[2026,"Level",55,180,80,65]]
    s.sheet("Productivity",rows,merges=["A1:A2","B1:B2","C1:F1","C104:F104"],headers=["A1:F2","A104:F104"],widths={"C":40,"D":32,"E":28,"F":32})
    t=table("Productivity",dict(year="A",type="B",rubber="C",chem="D",machine="E",food="F"),3,102)
    annual=filt(t,type="Annual")
    annual=derive(annual,gap=op("sub",field("rubber"),field("chem")),higher=op("gt",field("rubber"),field("chem")))
    for key,label in [("rubber","Rubber, plastics & non-metallic minerals"),("chem","Chemicals, Pharmaceuticals"),("machine","Machinery & equipment"),("food","Food, beverages & tobacco")]:
        sub=derive(annual,label=field("year"))
        s.add("annual_max",f"Using Annual index rows only, which year has the highest {label} output per job? Exclude quarters and Level rows; break ties by earliest year."+INT,ranked(sub,[[key,"desc"],["year","asc"]],["year"]))
        s.add("annual_mean",f"What is mean {label} index across all 20 Annual rows? Exclude quarters and Level."+TWO,agg("mean",annual,key),2)
        s.add("quarter_max",f"Using Quarter rows only, which period has the highest {label} index? Break ties by period ascending. Answer the period exactly.",ranked(filt(t,type="Quarter"),[[key,"desc"],["year","asc"]],["year"]))
        for year in [2017,2027]:
            col={"rubber":"C","chem":"D","machine":"E","food":"F"}[key]
            r=3+(year-2008)*5
            s.add("year_on_year",f"Using Annual index rows, what is {label} index in {year} minus {year-1}?"+ONE,op("sub",cell("Productivity",col,r),cell("Productivity",col,r-5)),1)
    s.add("annual_comparison_count","Across Annual index rows only, in how many years is Rubber, plastics & non-metallic minerals strictly above Chemicals, Pharmaceuticals?"+INT,agg("count",filt(annual,higher=True)))
    s.add("annual_gap_rank","Among Annual rows only, which year has the greatest Rubber minus Chemicals index gap? Break ties by earliest year."+INT,ranked(annual,[["gap","desc"],["year","asc"]],["year"]))
    for threshold in [90,100,110]:
        s.add("annual_threshold",f"How many Annual rows have Machinery & equipment index >= {threshold}? Exclude all Quarter and Level rows."+INT,agg("count",filt(annual,machine={"ge":threshold})))
    previous={"previous":{"data":annual,"order":"year","value":"rubber"}}
    changes=derive(previous,decrease=op("sub",field("previous"),field("rubber")))
    s.add("largest_drop","In Annual Rubber index rows, which year had the largest decrease versus the previous year? Answer `year, decrease` with decrease to exactly 1 decimal; break ties by earliest year.",items(ranked(changes,[["decrease","desc"],["year","asc"]],["year"]),formatted({"max":ranked(changes,[["decrease","desc"],["year","asc"]],["decrease"])},1)))
    s.add("period_count","How many Annual index rows are from 2015 through 2022 inclusive?"+INT,agg("count",filt(annual,year={"ge":2015,"le":2022})))
    for key,label in [("chem","Chemicals, Pharmaceuticals"),("food","Food, beverages & tobacco")]:
        sub=filt(annual,year={"ge":2020,"le":2027})
        s.add("recent_mean",f"What is mean Annual {label} index for 2020 through 2027 inclusive?"+TWO,agg("mean",sub,key),2)
    s.add("level_lookup","In the Level (GBP thousand) table, what is Chemicals, Pharmaceuticals for 2026? Use Level, not Annual or Quarter index."+INT,cell("Productivity","D",105))


def finance(s):
    segments=["Cloud","Devices","Services"]
    for segment in segments:
        prod=[s.rng.randint(400,2600) for _ in range(3)]; service=[s.rng.randint(900,4000) for _ in range(3)]
        rows=[["Dollars in millions","Fiscal2027","Fiscal2026","Fiscal2025"],
              ["Product revenue",*[f"${v:,}" for v in prod]],
              ["Service and other revenue",*[f"{v:,}" for v in service]],
              ["Total segment revenue",*[f"${a+b:,}" for a,b in zip(prod,service)]],
              ["Segment operating income",*[f"${s.rng.randint(100,900):,}" for _ in range(3)]],
              ["Restructuring expense",*[f"({s.rng.randint(10,90)})" for _ in range(3)]]]
        s.sheet(segment,rows,widths={"A":33})
    assets=[["Assets (millions)","July 31 2027",None,None,None,"July 31 2026",None,None,None],
            [None,"Level 1","Level 2","Level 3","Total fair value","Level 1","Level 2","Level 3","Total fair value"]]
    labels=["Cash equivalents","Municipal bonds","Corporate notes","Agency securities","Equity securities"]
    for label in labels:
        a=[s.rng.randint(0,900) for _ in range(3)]; b=[s.rng.randint(0,900) for _ in range(3)]
        assets.append([label,*a,sum(a),*b,sum(b)])
    assets.append(["TOTAL",*[sum(row[c] for row in assets[2:]) for c in range(1,9)]])
    s.sheet("Fair value",assets,merges=["B1:E1","F1:I1"],headers=["A1:I2"],widths={"A":30})
    funding=[["Funding (millions)","Year end",None,"Annual average",None],[None,2027,2026,2027,2026]]
    for label in ["Wholesale funding","Client cash management","Other borrowed funds","Preferred stock"]:
        funding.append([label,*[s.rng.randint(10000,50000) for _ in range(4)]])
    funding.append(["Net adjustment","(1,275)","—","(850)","—"])
    s.sheet("Funding",funding,merges=["B1:C1","D1:E1"],headers=["A1:E2"],widths={"A":33})
    for year,col in [(2027,"B"),(2026,"C"),(2025,"D")]:
        s.add("cross_sheet_sum",f"Across Cloud, Devices and Services, what is combined Total segment revenue in Fiscal{year}, in millions? Do not sum Product, Service and Total together."+INT,{"sum":items(*[{"parse":cell(seg,col,4)} for seg in segments])})
        records={"records":[{"label":seg,"value":{"parse":cell(seg,col,3)}} for seg in segments]}
        s.add("sheet_rank",f"Which sheet has the highest Service and other revenue in Fiscal{year}? Break ties by sheet name ascending. Answer `sheet name, integer value` without thousands separators.",ranked(records,[["value","desc"],["label","asc"]],["label","value"]))
        s.add("accounting_negative",f"What is combined signed Restructuring expense across the three segments in Fiscal{year}? Parentheses denote negative values."+INT,{"sum":items(*[{"parse":cell(seg,col,6)} for seg in segments])})
    for seg in segments:
        a={"parse":cell(seg,"B",4)}; b={"parse":cell(seg,"C",4)}
        s.add("revenue_growth",f"What is percentage revenue growth of {seg} from Fiscal2026 to Fiscal2027? Use Total segment revenue."+TWO,pct(op("sub",a,b),b),2)
        s.add("operating_margin",f"What is {seg} operating income as a percentage of Total segment revenue in Fiscal2027?"+TWO,pct({"parse":cell(seg,"B",5)},a),2)
    for j,label in enumerate(labels):
        s.add("fair_value_difference",f"What is July 31 2027 Total fair value minus July 31 2026 Total fair value for {label}, in millions?"+INT,op("sub",cell("Fair value","E",j+3),cell("Fair value","I",j+3)))
    s.add("avoid_total_double_count","By how much did TOTAL recurring fair value increase from July 31 2026 to July 31 2027, in millions? Use the TOTAL row once."+INT,op("sub",cell("Fair value","E",8),cell("Fair value","I",8)))
    for col,desc in [("B","year end 2027"),("C","year end 2026"),("D","annual average 2027"),("E","annual average 2026")]:
        s.add("cross_table_ratio",f"What is the ratio of Other borrowed funds to Preferred stock for {desc}? Answer exactly 3 decimal places.",op("div",cell("Funding",col,5),cell("Funding",col,6)),3)
    s.add("year_end_vs_average","For Other borrowed funds in 2027, what is year-end value minus annual average?"+INT,op("sub",cell("Funding","B",5),cell("Funding","D",5)))
    s.add("parentheses","What is signed Net adjustment at year end 2027? Parentheses indicate negatives."+INT,{"parse":cell("Funding","B",7)})
    s.add("fair_value_share","What percent of July 31 2027 TOTAL fair value is Level 2? Use the TOTAL row."+TWO,pct(cell("Fair value","C",8),cell("Fair value","E",8)),2)
    s.add("boolean","Is Cloud Fiscal2027 Total segment revenue greater than Devices Fiscal2027 Total segment revenue? Answer only Yes or No.",op("gt",{"parse":cell("Cloud","B",4)},{"parse":cell("Devices","B",4)}))
    s.add("exact_text","What literal missing marker appears in Funding for Net adjustment at year end 2026? Answer only that symbol.",cell("Funding","C",7))


def multilingual(s):
    care=[["Type de besoin","Statut du besoin (%)",None,None],[None,"unmet","partially met","fully met"]]
    for label in ["any","information","medication","counselling","other"]:
        a=s.rng.randint(50,350)/10; b=s.rng.randint(20,200)/10
        care.append([label,f"{a:.1f}",f"{b:.1f}",f"{100-a-b:.1f}"])
    care[-1][1]=".."
    s.sheet("Care",care,merges=["B1:D1"],headers=["A1:D2"],widths={"A":30})
    income=[["Trình độ học vấn","Thu nhập trung vị (đô la)",None],[None,"nam","nữ"]]
    levels=["chưa tốt nghiệp phổ thông","tốt nghiệp phổ thông","sau trung học","cao đẳng","đại học","sau đại học"]
    for label in levels:
        male=s.rng.randint(40000,90000); female=s.rng.randint(25000,80000)
        income.append([label,f"{male:,}",f"{female:,}"])
    s.sheet("Thu nhập",income,merges=["B1:C1"],headers=["A1:C2"],widths={"A":37})
    justice=[["연도","Total charges",None,"Total cases",None],[None,"number","percent change","number","percent change"]]
    for j in range(10):
        justice.append([f"{2017+j}/{2018+j}",f"{s.rng.randint(700000,1400000):,}","..",f"{s.rng.randint(260000,500000):,}",".."])
    s.sheet("Justice",justice,merges=["A1:A2","B1:C1","D1:E1"],headers=["A1:E2"])
    labour=[["指标 / 群体","已婚或同居","单身或曾结婚","差异（百分点）"]]
    for measure in ["劳动参与率","就业率","失业率"]:
        labour.append([measure])
        for label in ["原住民","非原住民"]:
            a=s.rng.randint(600,900)/10 if measure!="失业率" else s.rng.randint(30,150)/10
            b=s.rng.randint(500,850)/10 if measure!="失业率" else s.rng.randint(50,200)/10
            labour.append([label,f"{a:.1f}",f"{b:.1f}",f"{a-b:.1f}"])
    s.sheet("劳动",labour,widths={"A":30,"B":24,"C":24,"D":24})
    chronic=[["Groupe","arthrite","hypertension","asthme","trouble de l'humeur","anxiété"]]
    for label in ["Premières Nations hors réserve","Métis","Inuits","Population non autochtone"]:
        chronic.append([label,*[f"{s.rng.randint(50,300)/10:.1f}" for _ in range(5)]])
    s.sheet("Santé",chronic,widths={"A":38,"E":28})
    t=parsed(table("Care",dict(label="A",unmet="B",partial="C",full="D"),3,7),"unmet","partial","full")
    t=derive(filt(t,label={"not_blank":True},unmet={"not_blank":True}),combined=op("add",field("unmet"),field("partial")))
    for label in ["information","medication","counselling"]:
        s.add("vi_care",f"Trong Care, tổng unmet và partially met cho {label} là bao nhiêu? Chỉ trả số với đúng 1 chữ số thập phân.",agg("sum",filt(t,label=label),"combined"),1)
    s.add("vi_exclude_any","Trong Care, không tính dòng any và bỏ dòng thiếu unmet, tổng unmet + partially met cao nhất là bao nhiêu? Chỉ trả số với đúng 1 chữ số thập phân.",agg("max",filt(t,label={"ne":"any"}),"combined"),1)
    s.add("fr_missing","Dans Care, combien de types ont une valeur unmet manquante ? Le symbole .. signifie manquant. Répondez par un entier.",agg("count",filt(parsed(table("Care",dict(unmet="B"),3,7),"unmet"),unmet={"not_blank":False})))
    inc=parsed(table("Thu nhập",dict(label="A",male="B",female="C"),3,8),"male","female")
    inc=derive(inc,gap=absolute(op("sub",field("male"),field("female"))))
    for i,label in enumerate(levels):
        s.add("vi_income_ratio",f"Trong sheet Thu nhập, ở nhóm {label}, thu nhập nữ bằng bao nhiêu phần trăm thu nhập nam? Chỉ trả số với đúng 1 chữ số thập phân.",pct({"parse":cell("Thu nhập","C",i+3)},{"parse":cell("Thu nhập","B",i+3)}),1)
    s.add("vi_income_gap","Trong Thu nhập, chênh lệch tuyệt đối thu nhập nam và nữ lớn nhất giữa các trình độ là bao nhiêu đô la? Chỉ trả số nguyên không có dấu phân cách hàng nghìn.",agg("max",inc,"gap"))
    s.add("vi_income_rank","Trong Thu nhập, nhóm nào có thu nhập nữ cao nhất? Nếu bằng nhau chọn tên nhóm tăng dần Unicode. Chỉ trả tên nhóm.",ranked(inc,[["female","desc"],["label","asc"]],["label"]))
    for start,end in [(2017,2026),(2019,2024),(2020,2025)]:
        a={"parse":cell("Justice","D",start-2017+3)}; b={"parse":cell("Justice","D",end-2017+3)}
        s.add("ko_percent_change",f"Justice에서 {start}/{start+1}년 대비 {end}/{end+1}년 total cases의 감소율을 (이전-이후)/이전*100으로 계산하세요. 음수도 허용합니다. 소수점 둘째 자리까지 % 기호 없이 숫자만 답하십시오.",pct(op("sub",a,b),a),2)
    j=parsed(table("Justice",dict(label="A",charges="B",cases="D"),3,12),"charges","cases")
    s.add("ko_max","Justice에서 total cases가 가장 많은 연도를 표에 나온 그대로 답하십시오. 동률이면 가장 이른 연도입니다.",ranked(j,[["cases","desc"],["label","asc"]],["label"]))
    s.add("ko_total","Justice의 10개 연도 total charges 합계를 쉼표 없는 정수로 답하십시오.",agg("sum",j,"charges"))
    for measure,r in [("劳动参与率",3),("就业率",6),("失业率",9)]:
        s.add("zh_group_gap",f"在劳动表中，对于单身或曾结婚的人群，非原住民的{measure}减去原住民的{measure}是多少个百分点？只回答数字，保留一位小数。",op("sub",{"parse":cell("劳动","C",r+1)},{"parse":cell("劳动","C",r)}),1)
        s.add("zh_marital_gap",f"在劳动表中，原住民的{measure}，已婚或同居减去单身或曾结婚是多少个百分点？只回答数字，保留一位小数。",op("sub",{"parse":cell("劳动","B",r)},{"parse":cell("劳动","C",r)}),1)
    c=parsed(table("Santé",dict(label="A",mood="E",asthma="D"),2,4),"mood","asthma")
    s.add("fr_range","Dans Santé, parmi Premières Nations hors réserve, Métis et Inuits seulement, quel est l'écart entre le maximum et le minimum de trouble de l'humeur ? Répondez avec une décimale sans %.",op("sub",agg("max",c,"mood"),agg("min",c,"mood")),1)
    s.add("fr_rank","Dans Santé, parmi les trois groupes autochtones seulement, lequel a le taux d'asthme le plus élevé ? En cas d'égalité, ordre Unicode croissant. Répondez par le nom exact.",ranked(c,[["asthma","desc"],["label","asc"]],["label"]))
    s.add("fr_mean","Dans Santé, quelle est la moyenne simple du trouble de l'humeur parmi les trois groupes autochtones, sans Population non autochtone ? Répondez avec deux décimales.",agg("mean",c,"mood"),2)
    s.add("vi_missing_vs_zero","Trong Care, có bao nhiêu dòng có unmet bằng số 0? Không coi .. là 0. Chỉ trả số nguyên.",agg("count",filt(parsed(table("Care",dict(unmet="B"),3,7),"unmet"),unmet=0)))
    s.add("ko_ratio","Justice에서 2026/2027년 total charges를 total cases로 나눈 비율을 소수점 셋째 자리까지 숫자만 답하십시오.",op("div",{"parse":cell("Justice","B",12)},{"parse":cell("Justice","D",12)}),3)
    s.add("vi_income_total","Trong Thu nhập, tổng thu nhập nữ của hai nhóm đại học và sau đại học là bao nhiêu? Chỉ trả số nguyên không có dấu phân cách hàng nghìn.",op("add",{"parse":cell("Thu nhập","C",7)},{"parse":cell("Thu nhập","C",8)}))


DEFINITIONS=[
    ("13_admission_transfer",admission,"64 hồ sơ mới: tuổi thực theo ngày tại nhập học, tuổi tròn, ngày nhuận, lọc và top-k.",["Public_data_1/xlsx_files/p4_t01.xlsx"]),
    ("14_trade_periods",trade,"Xuất nhập khẩu của 6 mặt hàng mới; fiscal/calendar/FYTD trùng năm, số âm và tăng trưởng.",["Public_data_1/xlsx_files/p4_t02.xlsx"]),
    ("15_cpi_hierarchy",prices,"CPI mới với mức thụt cấp 0–3, con trực tiếp khác cháu, điều chỉnh mùa vụ khác chưa điều chỉnh.",["Public_data_1/xlsx_files/p4_t03.xlsx","Public_data_2/xlsx_files/p2_t02.xlsx"]),
    ("16_demographic_blocks",demographics,"Hai khối nam/nữ có nhãn lặp, 6 nhóm học vấn rời nhau; trung bình đơn khác tỷ lệ gộp.",["Public_data_1/xlsx_files/p4_t04.xlsx","Public_data_2/xlsx_files/p2_t04.xlsx","Public_data_2/xlsx_files/p2_t05.xlsx"]),
    ("17_industry_hours",industry,"Hai sheet ngành/nghề mới, tuổi, giờ làm, đơn vị nghìn, dòng TOTAL và phép tính chéo sheet.",[f"Public_data_1/xlsx_files/p4_t{i:02d}.xlsx" for i in range(5,11)]),
    ("18_time_use_blocks",time_use,"3 nhóm tuổi trẻ, mẹ/cha và tình trạng việc làm; giờ khác tỷ lệ tham gia, so sánh nhiều khối.",["Public_data_2/xlsx_files/p2_t01.xlsx"]),
    ("19_safety_naics",safety,"1.100 ngành mới, NAICS 3/6 chữ số, ô thiếu và bị ẩn, tỷ lệ tổng đã báo cáo, bảng quartile.",["Public_data_2/xlsx_files/p2_t06.xlsx","Public_data_2/xlsx_files/p2_t07.xlsx"]),
    ("20_productivity_periods",productivity,"20 năm xen 80 quý và bảng level riêng, chênh lệch liên năm, xếp hạng, phạm vi annual/quarter.",["Public_data_2/xlsx_files/p2_t03.xlsx","Public_data_2/xlsx_files/p2_t08.xlsx"]),
    ("21_finance_multisheet",finance,"5 sheet tài chính mới: số chuỗi $/dấu phẩy/ngoặc âm, doanh thu bộ phận, fair value, cuối kỳ khác bình quân.",["Public_data_2/xlsx_files/p2_t09.xlsx","Public_data_2/xlsx_files/p2_t10.xlsx"]),
    ("22_multilingual_social",multilingual,"Câu hỏi Việt/Pháp/Hàn/Trung; chăm sóc, thu nhập, tòa án, lao động và sức khỏe với số liệu mới.",[f"Public_data_2/xlsx_files/p2_t{i:02d}.xlsx" for i in range(11,16)]),
]


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--seed",type=int,default=261016)
    parser.add_argument("--test-root",type=Path,default=ROOT/"Test")
    parser.add_argument("--spec",type=Path,default=ROOT/"outputs/public-transfer/workbooks.json")
    args=parser.parse_args()
    inventory=[]
    import openpyxl
    for dataset in ["Public_data_1","Public_data_2"]:
        queries=json.loads((ROOT/dataset/"queries.json").read_text(encoding="utf-8"))
        for p in sorted((ROOT/dataset/"xlsx_files").glob("*.xlsx")):
            book=openpyxl.load_workbook(p,data_only=True)
            try:
                inventory.append({"file":p.relative_to(ROOT).as_posix(),"sha256":hashlib.sha256(p.read_bytes()).hexdigest(),
                                  "query_ids":[q["id"] for q in queries if q["file_name"]==p.name],
                                  "sheets":[{"name":w.title,"rows":w.max_row,"columns":w.max_column,"merged_ranges":list(map(str,w.merged_cells.ranges))} for w in book]})
            finally: book.close()
    specs=[]
    for name,fn,description,sources in DEFINITIONS:
        s=TransferSuite(name,description,sources,random.Random(args.seed+int(name[:2])),args.seed)
        fn(s)
        specs.append(s.save(args.test_root))
    args.spec.parent.mkdir(parents=True,exist_ok=True)
    args.spec.write_text(json.dumps(specs,ensure_ascii=False,indent=2),encoding="utf-8")
    (args.test_root/"public_transfer_manifest.json").write_text(json.dumps({"seed":args.seed,"suite_count":10,"question_count":300,
        "suites":[{"name":s["suite"],"workbook":"xlsx_files/data.xlsx","questions":30,"sheets":len(s["sheets"])} for s in specs],
        "source_inventory":inventory},ensure_ascii=False,indent=2),encoding="utf-8")
    print(f"Prepared 10 suites / 300 questions with seed {args.seed}; export spec: {args.spec}")


if __name__=="__main__": main()
