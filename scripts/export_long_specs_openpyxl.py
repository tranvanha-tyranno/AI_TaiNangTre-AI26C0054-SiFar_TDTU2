"""Export long-suite JSON specs to XLSX without the authoring-only Node dependency."""
from __future__ import annotations
import json, re, sys
from datetime import date, datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

def value(v):
    if isinstance(v,str):
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}",v):
            try: return date.fromisoformat(v)
            except ValueError: pass
        if re.fullmatch(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}",v):
            try: return datetime.fromisoformat(v)
            except ValueError: pass
    return v

def export(spec_path):
    spec=json.loads(Path(spec_path).read_text(encoding='utf-8'))[0]
    out=Path(spec['file']); out.parent.mkdir(parents=True,exist_ok=True)
    wb=openpyxl.Workbook(); wb.remove(wb.active)
    for source in spec['sheets']:
        ws=wb.create_sheet(source['name'])
        rows=source['rows']
        for r_idx,row in enumerate(rows,1):
            for c_idx,v in enumerate(row,1):
                cell=ws.cell(r_idx,c_idx,value(v))
                cell.font=Font(name='Arial',size=10,color='243247')
                cell.alignment=Alignment(vertical='center')
        for merged in source.get('merges',[]): ws.merge_cells(merged)
        for address in source.get('headers',[]):
            for row in ws[address]:
                for cell in row:
                    cell.fill=PatternFill('solid',fgColor='E8EEF7')
                    cell.font=Font(name='Arial',size=10,bold=True,color='243247')
                    cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
                    ws.row_dimensions[cell.row].height=40
        for rng,fmt in source.get('formats',{}).items():
            for row in ws[rng]:
                for cell in row: cell.number_format=fmt
        for col,width in source.get('widths',{}).items(): ws.column_dimensions[col].width=width
        header=source.get('headers',['A1:Z1'])[0]
        header_row=int(re.search(r'(\d+)$',header).group(1))
        ws.freeze_panes=f"A{header_row}"
        ws.sheet_view.showGridLines=False
    wb.save(out)
    print(f'Exported {spec["suite"]}: {out}')

if __name__=='__main__':
    for arg in sys.argv[1:]: export(arg)
