"""Authoring/evaluation utilities only. Never imported by the answering pipeline."""
from __future__ import annotations
from collections import Counter
import csv
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
import json
from pathlib import Path
import random
import sqlite3


def serial(value):
    if isinstance(value, datetime): return value.isoformat(sep=" ")
    if isinstance(value, date): return value.isoformat()
    return value


def quote(name): return '"' + name.replace('"', '""') + '"'


def colname(n):
    out=""
    while n:
        n,r=divmod(n-1,26); out=chr(65+r)+out
    return out


def sql_table(connection, spec, rows):
    columns=spec["columns"]
    declaration=", ".join(f'{quote(c["key"])} {c["type"]}' for c in columns)
    connection.execute(f'CREATE TABLE {quote(spec["table"])} ({declaration})')
    connection.executemany(f'INSERT INTO {quote(spec["table"])} VALUES ({",".join("?" for _ in columns)})',
                           ([serial(v) for v in row] for row in rows))


def number_text(value, digits=None):
    if value is None: raise ValueError("Unexpected missing answer; use explicit empty presentation")
    d=Decimal(str(value))
    if digits is not None:
        return f'{d.quantize(Decimal(1).scaleb(-digits),rounding=ROUND_HALF_UP):.{digits}f}'
    return str(int(d)) if d==int(d) else format(d.normalize(),"f")


def present(result, spec):
    kind=spec.get("kind","scalar")
    if kind=="empty":
        if result and any(v is not None for row in result for v in row): raise ValueError("Empty answer has unexpected source value")
        return ""
    if not result: raise ValueError("Question has no rows without an explicit missing-data contract")
    if kind=="ratio":
        a,b=result[0]
        if not b: raise ValueError("Zero answer denominator")
        return number_text(Decimal(str(a))/Decimal(str(b))*Decimal(str(spec.get("scale",1))),spec["digits"])
    formats=spec.get("formats",["text"])
    def item(value, style):
        if style=="text":
            if value is None: raise ValueError("Missing text answer")
            return str(value)
        if style=="yesno": return "Yes" if value else "No"
        if style=="int": return number_text(value)
        return number_text(value,int(style))
    return ", ".join(item(v,formats[min(i,len(formats)-1)]) for row in result for i,v in enumerate(row))


class LongSuite:
    def __init__(self, name, description, count, seed):
        self.name,self.description,self.count,self.seed=name,description,count,seed
        self.rng=random.Random(seed+int(name[:2]))
        self.db=sqlite3.connect(":memory:")
        self.sheets,self.tables,self.queries,self.checks,self.views=[],[],[],[],[]

    def table(self, name, sheet, columns, records, header_rows=None, merges=(), widths=None, repeats=0):
        """Columns: (SQL key, visible label, SQLite type). Every record is meaningful data."""
        records=list(records)
        headers=header_rows or [[c[1] for c in columns]]
        raw=[list(r) for r in headers]; excluded=[]
        for i,record in enumerate(records):
            if repeats and i and i%repeats==0:
                excluded.append(len(raw)+1); raw.append([c[1] for c in columns])
            raw.append(list(record))
        end=colname(len(columns)); first=len(headers)+1
        schema={"table":name,"sheet":sheet,"first_row":first,"last_row":len(raw),"record_count":len(records),
                "skip_rows":excluded,"columns":[{"key":c[0],"label":c[1],"type":c[2],"column":i+1} for i,c in enumerate(columns)]}
        sql_table(self.db,schema,records)
        self.tables.append(schema)
        formats={}
        for i,c in enumerate(columns):
            if c[2]=="INTEGER": formats[f"{colname(i+1)}{first}:{colname(i+1)}{len(raw)}"]="0"
        self.sheets.append({"name":sheet,"start":"A1","rows":[[serial(v) for v in row] for row in raw],
                            "merges":list(merges),"formats":formats,"headers":[f"A1:{end}{len(headers)}"],
                            "widths":widths or {},"previews":[f"A1:{min_col(end,'H')}10",f"A{max(first,len(raw)-6)}:{min_col(end,'H')}{len(raw)}"]})
        return schema

    def view(self, sql):
        self.db.execute(sql); self.views.append(sql)

    def add(self, language, category, question, sql, params=(), formats=None, ratio=None, empty=False):
        ident=int(self.name[:2])*1000+len(self.queries)+1
        presentation={"kind":"empty"} if empty else ({"kind":"ratio",**ratio} if ratio else {"kind":"scalar","formats":formats or ["int"]})
        answer=present(self.db.execute(sql,params).fetchall(),presentation)
        self.queries.append({"id":ident,"file_name":"data.xlsx","query":question})
        self.checks.append({"id":ident,"language":language,"category":category,"sql":sql,"params":list(params),
                            "presentation":presentation,"answer":answer})

    def save(self, test_root, spec_root):
        assert len(self.queries)==30,(self.name,len(self.queries))
        assert Counter(c["language"] for c in self.checks)=={"vi":15,"en":5,"zh":4,"fr":3,"ko":3}
        assert sum(c["answer"]=="" for c in self.checks)==1
        folder=(test_root/self.name).resolve(); (folder/"xlsx_files").mkdir(parents=True,exist_ok=True)
        self.rng.shuffle(self.queries)
        assert [q["id"] for q in self.queries]!=sorted(q["id"] for q in self.queries)
        (folder/"queries.json").write_text(json.dumps(self.queries,ensure_ascii=False,indent=2),encoding="utf-8")
        checks=sorted(self.checks,key=lambda c:c["id"])
        for filename,blank in [("expected_submission.csv",False),("sample_submission.csv",True)]:
            with (folder/filename).open("w",encoding="utf-8",newline="") as f:
                w=csv.writer(f,lineterminator="\n"); w.writerow(["id","answer"])
                w.writerows((c["id"],"" if blank else c["answer"]) for c in checks)
        audit={"schema_version":1,"views":self.views,"checks":checks}
        (folder/"answer_audit.json").write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding="utf-8")
        manifest={"suite":self.name,"seed":self.seed,"description":self.description,"questions":30,
                  "language_counts":dict(Counter(c["language"] for c in checks)),"tables":self.tables,
                  "agrees_with":"AGENTS.md sections 3, 11.1, 11.2; synthetic benchmark, not an official scoring replica",
                  "gold_order":"ascending numeric ID","unknown_answer":"","data_kind":"new fictional records"}
        (folder/"dataset_manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
        sizes="; ".join(f'`{t["sheet"]}`: {t["record_count"]:,} bản ghi / {t["last_row"]:,} dòng vật lý' for t in self.tables)
        (folder/"README.md").write_text(
            f"# {self.name}\n\n{self.description}\n\n{sizes}.\n\n"
            "30 câu: **15 Việt, 5 Anh, 4 Trung, 3 Pháp, 3 Hàn**. ID là số tự nhiên; câu hỏi cố ý xáo trộn. "
            "CSV chuẩn dùng UTF-8 và sắp ID theo giá trị số tăng dần. Một câu thiếu trường dữ liệu có đáp án chuỗi rỗng. "
            "Câu đó không tính là lỗi fixture; không tự đổi đáp án thành N/A.\n\n"
            "Dữ liệu hư cấu được sinh mới theo seed, có khóa và quan hệ nghiệp vụ nhất quán. "
            "Dòng cuối, phần giữa và các sheet phụ đều có câu yêu cầu truy xuất; không thể giải đầy đủ bằng preview đầu bảng. "
            "Các câu quy định phạm vi, đơn vị, xử lý thiếu dữ liệu, phá hòa và định dạng khi cần.\n\n"
            "Chỉ truyền `xlsx_files/` và `queries.json` vào solver. `answer_audit.json` chứa SQL/đáp án để kiểm chứng "
            "bên ngoài mô hình; `dataset_manifest.json` mô tả vùng dữ liệu cho bộ kiểm tra. Không đưa các file này hoặc "
            "`expected_submission.csv` vào prompt/retrieval. `sample_submission.csv` chỉ là mẫu rỗng.\n\n"
            f"```powershell\npython main.py --mode llm --xlsx_dir Test/{self.name}/xlsx_files --queries Test/{self.name}/queries.json "
            f"--output outputs/{self.name}.csv --log-dir outputs/{self.name}-logs\n"
            f"python scripts/evaluate.py --prediction outputs/{self.name}.csv --reference Test/{self.name}/expected_submission.csv "
            f"--report outputs/{self.name}-strict.json\n```\n\n"
            "Điểm trên là strict exact match, chưa tái tạo chuẩn hóa BTC. Chưa chạy LLM hoặc đo chi phí/thời gian của mô hình. "
            "Quy định tài nguyên bài thi là RAM ≤ 4 GB, disk ≤ 1 GB và ZIP ≤ 100 MB theo AGENTS.md; "
            "kích thước file nhỏ không chứng minh cả pipeline đạt các giới hạn đó.\n",encoding="utf-8")
        spec_root.mkdir(parents=True,exist_ok=True)
        spec_path=spec_root/(self.name+".json")
        spec_path.write_text(json.dumps([{"suite":self.name,"file":str(folder/"xlsx_files/data.xlsx"),"sheets":self.sheets}],ensure_ascii=False),encoding="utf-8")
        self.db.close()
        return {"name":self.name,"questions":30,"main_records":self.count,"sheets":len(self.tables),"spec":str(spec_path)}


def min_col(a,b): return a if len(a)==1 and a<b else b


def rebuild_database(folder):
    """Stream saved XLSX into SQLite; no generator and no solver required."""
    import openpyxl
    manifest=json.loads((folder/"dataset_manifest.json").read_text(encoding="utf-8"))
    book=openpyxl.load_workbook(folder/"xlsx_files/data.xlsx",data_only=True,read_only=True)
    db=sqlite3.connect(":memory:")
    try:
        for table in manifest["tables"]:
            skip=set(table["skip_rows"])
            worksheet=book[table["sheet"]]
            raw=worksheet.iter_rows(min_row=table["first_row"],max_row=table["last_row"],max_col=len(table["columns"]),values_only=True)
            sql_table(db,table,(row for r,row in enumerate(raw,table["first_row"]) if r not in skip))
        audit=json.loads((folder/"answer_audit.json").read_text(encoding="utf-8"))
        for sql in audit["views"]: db.execute(sql)
    except Exception:
        db.close(); raise
    finally: book.close()
    return db,manifest,audit
